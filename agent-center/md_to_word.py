import os
import sys
from docx import Document
from docx.shared import Pt, RGBColor
from markdown_it import MarkdownIt

def process_inline(paragraph, inline_token):
    """
    处理行内标记（加粗、斜体、代码等）
    """
    if not inline_token.children:
        paragraph.add_run(inline_token.content)
        return

    is_bold = False
    is_italic = False
    is_code = False

    for child in inline_token.children:
        if child.type == 'text':
            run = paragraph.add_run(child.content)
            run.bold = is_bold
            run.italic = is_italic
            if is_code:
                run.font.name = 'Courier New'
        elif child.type == 'strong_open':
            is_bold = True
        elif child.type == 'strong_close':
            is_bold = False
        elif child.type == 'em_open':
            is_italic = True
        elif child.type == 'em_close':
            is_italic = False
        elif child.type == 'code_inline':
            run = paragraph.add_run(child.content)
            run.font.name = 'Courier New'
        elif child.type == 'softbreak':
            paragraph.add_run('\n')
        elif child.type == 'link_open':
            pass
        elif child.type == 'link_close':
            pass

def convert_md_to_docx(md_filepath, docx_filepath):
    """
    将 Markdown 文件转换为 Word (docx) 文件 (无需 Pandoc)
    """
    if not os.path.exists(md_filepath):
        print(f"错误: 找不到文件 {md_filepath}")
        return

    print(f"正在读取 Markdown 文件: {md_filepath}")
    with open(md_filepath, 'r', encoding='utf-8') as f:
        md_text = f.read()

    # 使用 markdown-it-py 解析，启用表格支持
    md = MarkdownIt("commonmark").enable("table")
    tokens = md.parse(md_text)

    doc = Document()
    
    print("正在转换中...")
    i = 0
    while i < len(tokens):
        token = tokens[i]
        
        # 标题
        if token.type == 'heading_open':
            level = int(token.tag[1])
            i += 1
            p = doc.add_heading('', level=level)
            while tokens[i].type != 'heading_close':
                if tokens[i].type == 'inline':
                    process_inline(p, tokens[i])
                i += 1
        
        # 段落
        elif token.type == 'paragraph_open':
            i += 1
            p = doc.add_paragraph()
            while tokens[i].type != 'paragraph_close':
                if tokens[i].type == 'inline':
                    process_inline(p, tokens[i])
                i += 1
        
        # 列表 (无序和有序)
        elif token.type == 'bullet_list_open' or token.type == 'ordered_list_open':
            list_type = 'List Bullet' if token.type == 'bullet_list_open' else 'List Number'
            i += 1
            while tokens[i].type not in ['bullet_list_close', 'ordered_list_close']:
                if tokens[i].type == 'list_item_open':
                    i += 1
                    while tokens[i].type != 'list_item_close':
                        if tokens[i].type == 'paragraph_open':
                            i += 1
                            p = doc.add_paragraph(style=list_type)
                            while tokens[i].type != 'paragraph_close':
                                if tokens[i].type == 'inline':
                                    process_inline(p, tokens[i])
                                i += 1
                        else:
                            i += 1
                else:
                    i += 1
        
        # 表格
        elif token.type == 'table_open':
            table_data = []
            while tokens[i].type != 'table_close':
                if tokens[i].type == 'tr_open':
                    row = []
                    while tokens[i].type != 'tr_close':
                        if tokens[i].type in ['th_open', 'td_open']:
                            i += 1
                            cell_tokens = []
                            while tokens[i].type not in ['th_close', 'td_close']:
                                cell_tokens.append(tokens[i])
                                i += 1
                            row.append(cell_tokens)
                        i += 1
                    table_data.append(row)
                i += 1
            
            if table_data:
                table = doc.add_table(rows=len(table_data), cols=len(table_data[0]))
                table.style = 'Table Grid'
                for r_idx, row_data in enumerate(table_data):
                    for c_idx, cell_tokens in enumerate(row_data):
                        cell = table.cell(r_idx, c_idx)
                        if cell.paragraphs:
                            p = cell.paragraphs[0]
                        else:
                            p = cell.add_paragraph()
                        
                        for ct in cell_tokens:
                            if ct.type == 'inline':
                                process_inline(p, ct)
        
        # 分隔符
        elif token.type == 'hr':
            doc.add_page_break()
        
        i += 1

    doc.save(docx_filepath)
    print(f"\n✅ 转换成功！生成文件: {docx_filepath}")

if __name__ == "__main__":
    # Markdown 文件路径
    source_md = r"D:\choice_product\agent-center\1.1.2版本.md"
    # 输出的 Word 文件路径
    target_docx = r"D:\choice_product\agent-center\1.1.2版本.docx"
    
    convert_md_to_docx(source_md, target_docx)
