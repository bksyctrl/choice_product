import base64
import difflib
import io
import json
import math
import os
import py_compile
import re
import sys
import threading
import time
import traceback
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

import pymysql
import requests
from dotenv import load_dotenv
from flask import Flask, jsonify, redirect, render_template, request, send_file, session
from flask_cors import CORS
from openai import OpenAI
from openpyxl import Workbook
from PIL import Image, ImageFilter
from werkzeug.security import check_password_hash, generate_password_hash


load_dotenv()

TOOLS_DIR = Path(__file__).resolve().parent / "tools"
if str(TOOLS_DIR) not in sys.path:
    sys.path.insert(0, str(TOOLS_DIR))

from analyze_single_product import (  # noqa: E402
    MINIMAX_BASE_URL,
    MINIMAX_MODEL,
    PRODUCT_SOURCES as AI_PRODUCT_SOURCES,
    build_input_snapshot,
    call_minimax as call_product_vision_minimax,
    load_product as load_ai_product,
    parse_json_response as parse_product_json_response,
    resolve_minimax_api_key,
    run_analysis,
    run_material_prefilter,
)


DB_CONFIG = {
    "host": os.getenv("DB_HOST", "192.168.0.168"),
    "port": int(os.getenv("DB_PORT", "3306")),
    "user": os.getenv("DB_USER", "ITaimysql"),
    "password": os.getenv("DB_PASSWORD", "Ai12345678@"),
    "database": os.getenv("DB_NAME", "ecommerce_workflow"),
    "charset": "utf8mb4",
    "cursorclass": pymysql.cursors.DictCursor,
    "connect_timeout": int(os.getenv("DB_CONNECT_TIMEOUT", "5")),
    "read_timeout": int(os.getenv("DB_READ_TIMEOUT", "20")),
    "write_timeout": int(os.getenv("DB_WRITE_TIMEOUT", "20")),
}

STATUS_VALUES = {"PENDING", "REVIEWING", "READY", "PUBLISHED", "OTHER"}
IP_GRADES = ["E", "D", "C", "B", "A", "S"]
AI_DAEMON_ENABLED = os.getenv("AI_DAEMON_ENABLED", "false").strip().lower() == "true"
AI_DAEMON_SOURCE = os.getenv("AI_DAEMON_SOURCE", "fastmoss").strip()
AI_DAEMON_SOURCES = [
    item.strip()
    for item in os.getenv("AI_DAEMON_SOURCES", AI_DAEMON_SOURCE).split(",")
    if item.strip()
]
AI_DAEMON_ANALYSIS = os.getenv("AI_DAEMON_ANALYSIS", "both").strip().lower()
AI_DAEMON_INTERVAL_SECONDS = int(os.getenv("AI_DAEMON_INTERVAL_SECONDS", "300"))
AI_DAEMON_BATCH_SIZE = max(30, int(os.getenv("AI_DAEMON_BATCH_SIZE", "30")))
AI_DAEMON_CONCURRENCY = max(1, min(8, int(os.getenv("AI_DAEMON_CONCURRENCY", "8"))))
AI_DAEMON_WRITE = os.getenv("AI_DAEMON_WRITE", "true").strip().lower() != "false"
AI_DAEMON_LOCK_RETRIES = max(1, int(os.getenv("AI_DAEMON_LOCK_RETRIES", "3")))
AI_DAEMON_LOCK_WAIT_SECONDS = max(1, int(os.getenv("AI_DAEMON_LOCK_WAIT_SECONDS", "3")))
AI_DAEMON_LATEST_ONLY = os.getenv("AI_DAEMON_LATEST_ONLY", "true").strip().lower() != "false"
SCORE_DAEMON_ENABLED = os.getenv("SCORE_DAEMON_ENABLED", "true").strip().lower() != "false"
SCORE_DAEMON_INTERVAL_SECONDS = int(os.getenv("SCORE_DAEMON_INTERVAL_SECONDS", "180"))
SCORE_DAEMON_BATCH_SIZE = max(20, int(os.getenv("SCORE_DAEMON_BATCH_SIZE", "200")))
SCORE_RECALC_MAX_ATTEMPTS = max(1, int(os.getenv("SCORE_RECALC_MAX_ATTEMPTS", "3")))
SCORE_RECALC_LATEST_ONLY = os.getenv("SCORE_RECALC_LATEST_ONLY", "false").strip().lower() == "true"
SCORE_AUTO_CREATE_INDEXES = os.getenv("SCORE_AUTO_CREATE_INDEXES", "false").strip().lower() == "true"
ILLUSTRATION_DAEMON_ENABLED = os.getenv("ILLUSTRATION_DAEMON_ENABLED", "true").strip().lower() != "false"
ILLUSTRATION_DAEMON_INTERVAL_SECONDS = int(os.getenv("ILLUSTRATION_DAEMON_INTERVAL_SECONDS", "120"))
ILLUSTRATION_DAEMON_BATCH_SIZE = max(20, int(os.getenv("ILLUSTRATION_DAEMON_BATCH_SIZE", "120")))
ILLUSTRATION_DAEMON_CONCURRENCY = max(1, min(8, int(os.getenv("ILLUSTRATION_DAEMON_CONCURRENCY", "4"))))
ILLUSTRATION_DAEMON_SOURCES = [
    item.strip()
    for item in os.getenv("ILLUSTRATION_DAEMON_SOURCES", "pod_cross_category").split(",")
    if item.strip()
]
ILLUSTRATION_DAEMON_LATEST_ONLY = os.getenv("ILLUSTRATION_DAEMON_LATEST_ONLY", "true").strip().lower() != "false"
MINIMAX_IMAGE_GENERATION_URL = os.getenv("MINIMAX_IMAGE_GENERATION_URL", "https://api.minimax.io/v1/image_generation")
MINIMAX_IMAGE_MODEL = os.getenv("MINIMAX_IMAGE_MODEL", "image-01")
MINIMAX_IMAGE_SUBJECT_TYPE = os.getenv("MINIMAX_IMAGE_SUBJECT_TYPE", "character")
MINIMAX_ILLUSTRATION_MAX_ATTEMPTS = max(1, min(5, int(os.getenv("MINIMAX_ILLUSTRATION_MAX_ATTEMPTS", "3"))))
MINIMAX_ILLUSTRATION_PASS_SCORE = max(0, min(100, int(os.getenv("MINIMAX_ILLUSTRATION_PASS_SCORE", "85"))))
POD_ILLUSTRATION_STALE_MINUTES = max(1, int(os.getenv("POD_ILLUSTRATION_STALE_MINUTES", "10")))
DETAIL_ANALYSIS_TABLE = "cp_user_detail_analysis"
EXPERT_TEAM_SESSION_TABLE = "cp_expert_team_session"
EXPERT_TEAM_MESSAGE_TABLE = "cp_expert_team_message"
EXPERT_WORKFLOW_INSTANCE_TABLE = "cp_expert_workflow_instance"
EXPERT_WORKFLOW_STEP_TABLE = "cp_expert_workflow_step"
EXPERT_TEAM_ALLOWED_ROLES = {"admin", "manager"}
PERMISSION_EXPERT_TEAM = "expert_team"
PROJECT_ROOT = Path(__file__).resolve().parent
EXPERT_FILE_CONTEXT_MAX_FILES = 12
EXPERT_FILE_CONTEXT_MAX_MATCHES_PER_FILE = 5
EXPERT_FILE_CONTEXT_MAX_CHARS = 30000
EXPERT_FILE_ALLOWED_SUFFIXES = {".py", ".html", ".css", ".js", ".md", ".yml", ".yaml", ".xml", ".json", ".txt"}
EXPERT_FILE_EXCLUDED_DIRS = {
    ".git",
    ".idea",
    ".venv",
    "__pycache__",
    "target",
    "node_modules",
    "tmp_render",
    "ComfyUI_ImageToText-main",
}
CODEPATCH_ALLOWED_SUFFIXES = EXPERT_FILE_ALLOWED_SUFFIXES | {".java", ".properties", ".toml", ".sql"}
CODEPATCH_WRITE_ROOTS = (
    "app.py",
    "static",
    "tools",
    "agent-center/src/main",
    "agent-center/pom.xml",
    "agent-center/README.md",
    "agent-center/1.1.2版本.md",
    "requirements.txt",
    ".env.example",
    "product_selection_platform.md",
    "private_product_library_plan.md",
    "1.1版本-2026年5月30日.md",
    "1.1.2版本.md",
)
CODEPATCH_NEVER_WRITE_DIRS = EXPERT_FILE_EXCLUDED_DIRS | {
    "uploads",
    "backups",
}
CODEPATCH_DANGEROUS_PATTERNS = [
    r"\bos\.system\s*\(",
    r"\bsubprocess\.",
    r"\bshutil\.rmtree\s*\(",
    r"\bos\.remove\s*\(",
    r"\bos\.unlink\s*\(",
    r"\bRemove-Item\b",
    r"\brm\s+-rf\b",
    r"\bdel\s+/[fsq]\b",
    r"\bformat\s+[A-Za-z]:",
]
DB_READONLY_DIAGNOSTIC_PREFIXES = ("SHOW INDEX", "SHOW COLUMNS", "SHOW CREATE TABLE", "EXPLAIN SELECT")
DB_READONLY_PERMISSION_PHRASES = [
    "允许访问数据库",
    "同意访问数据库",
    "可以访问数据库",
    "授权访问数据库",
    "允许查数据库",
    "同意查数据库",
    "可以查库",
    "授权查库",
    "允许执行只读sql",
    "允许执行只读 SQL",
]
DETAIL_ANALYSIS_EXECUTOR = ThreadPoolExecutor(max_workers=max(1, min(4, int(os.getenv("DETAIL_ANALYSIS_CONCURRENCY", "2")))))
_ai_daemon_started = False
_score_daemon_started = False
_illustration_daemon_started = False
_ai_daemon_current_concurrency = AI_DAEMON_CONCURRENCY
SYSTEM_DAILY_STATS_LOG = PROJECT_ROOT / "logs" / "system_daily_stats.jsonl"
GENERATED_ILLUSTRATION_DIR = PROJECT_ROOT / "static" / "generated" / "illustrations"
GENERATED_ILLUSTRATION_URL_PREFIX = "/static/generated/illustrations"
SYSTEM_DAILY_STATS_KEYS = (
    "illustration_checked",
    "ip_analyzed",
    "material_analyzed",
    "selection_scored",
)
_system_daily_stats_lock = threading.Lock()
FASTMOSS_ANALYSIS_REUSE_SOURCES = ("fastmoss", "fastmoss_rank")


class ApiAuthError(Exception):
    pass


class ApiPermissionError(Exception):
    pass


def empty_daily_stats():
    return {key: 0 for key in SYSTEM_DAILY_STATS_KEYS}


def load_latest_daily_stats(target_date):
    if not SYSTEM_DAILY_STATS_LOG.exists():
        return empty_daily_stats()
    latest = None
    try:
        with SYSTEM_DAILY_STATS_LOG.open("r", encoding="utf-8") as handle:
            for line in handle:
                line = line.strip()
                if not line:
                    continue
                try:
                    item = json.loads(line)
                except json.JSONDecodeError:
                    continue
                if item.get("date") == target_date and isinstance(item.get("totals"), dict):
                    latest = item["totals"]
    except OSError:
        return empty_daily_stats()
    totals = empty_daily_stats()
    if latest:
        for key in SYSTEM_DAILY_STATS_KEYS:
            totals[key] = int(float(latest.get(key) or 0))
    return totals


def record_daily_system_stats(event, delta=None, detail=None):
    normalized_delta = empty_daily_stats()
    for key, value in (delta or {}).items():
        if key in normalized_delta:
            normalized_delta[key] = int(float(value or 0))
    if not any(normalized_delta.values()):
        return
    target_date = datetime.now().strftime("%Y-%m-%d")
    with _system_daily_stats_lock:
        totals = load_latest_daily_stats(target_date)
        for key, value in normalized_delta.items():
            totals[key] += value
        entry = {
            "timestamp": datetime.now().isoformat(timespec="seconds"),
            "date": target_date,
            "event": event,
            "delta": normalized_delta,
            "totals": totals,
            "detail": detail or {},
        }
        SYSTEM_DAILY_STATS_LOG.parent.mkdir(parents=True, exist_ok=True)
        with SYSTEM_DAILY_STATS_LOG.open("a", encoding="utf-8") as handle:
            handle.write(json.dumps(entry, ensure_ascii=False) + "\n")
    print(f"[SYSTEM_DAILY_STATS] {event} delta={normalized_delta} totals={totals}", flush=True)


def reset_stale_pod_illustration_tasks():
    try:
        with db() as conn, conn.cursor() as cursor:
            cursor.execute(
                """
                UPDATE pod_illustration_task
                SET status = 3,
                    error_msg = '图生图任务超时未完成，已自动标记失败，可重新提取'
                WHERE status IN (0, 1)
                  AND updated_at < DATE_SUB(NOW(), INTERVAL %s MINUTE)
                """,
                (POD_ILLUSTRATION_STALE_MINUTES,),
            )
            cursor.execute(
                """
                UPDATE pod_cross_category_product p
                JOIN pod_illustration_task t
                  ON t.product_id = p.product_id
                 AND t.date_record = p.date_record
                SET p.illustration_status = 3
                WHERE p.illustration_status = 1
                  AND p.illustration_extractable = 1
                  AND t.status = 3
                  AND t.updated_at >= DATE_SUB(NOW(), INTERVAL %s MINUTE)
                """,
                (POD_ILLUSTRATION_STALE_MINUTES + 1,),
            )
            changed = cursor.rowcount
            conn.commit()
            if changed:
                print(f"[POD_ILLUSTRATION] reset stale tasks changed={changed}", flush=True)
    except Exception as exc:
        print(f"[POD_ILLUSTRATION] reset stale tasks failed: {exc}", flush=True)

SOURCES = {
    "unified": {
        "label": "统一表",
        "table": "fastmoss_product_aggregate",
        "id": "product_id",
        "date": "date_record",
        "title": "title",
        "image": "image_base64",
        "price": "real_price",
        "rating": "rating",
        "commission": "commission_rate",
        "sold": "rank_sold_count",
        "total_sold": "sold_count",
        "sale_amount": "sale_amount",
        "author_count": "author_count",
        "base_price": "base_price",
        "transport_fee": "transport_fee",
        "video_ratio": None,
        "product_card_ratio": None,
        "distribution_30d": "content_ratio",
        "distribution_7d": "distribution_7d",
        "distribution_90d": "distribution_90d",
        "distribution_180d": "distribution_180d",
        "overview_30d": "sales_overview",
        "overview_7d": "overview_7d",
        "overview_90d": "overview_90d",
        "overview_180d": "overview_180d",
        "sku_analysis_7d": "sku_analysis_7d",
        "sku_analysis_28d": "sku_analysis_28d",
        "detail_url": "detail_url",
        "platform": "fastmoss",
        "selling_points": "selling_points",
        "attributes": "attributes",
        "score": "selection_score",
        "score_reason": "score_reason",
        "sales_mom": "sales_mom",
        "status": "illustration_status",
        "illustration_result_url": "illustration_result_url",
        "illustration_extractable": "illustration_extractable",
        "illustration_extract_reason": "illustration_extract_reason",
    },
    "fastmoss": {
        "label": "FastMoss",
        "table": "fastmoss_product_rank_aggregate",
        "id": "product_id",
        "date": "date_record",
        "title": "title",
        "image": "image_base64",
        "price": "real_price",
        "rating": "rating",
        "commission": "commission_rate",
        "sold": "rank_sold_count",
        "total_sold": "sold_count",
        "sale_amount": "sale_amount",
        "author_count": "author_count",
        "base_price": "base_price",
        "transport_fee": "transport_fee",
        "video_ratio": None,
        "product_card_ratio": None,
        "distribution_30d": "content_ratio",
        "distribution_7d": "distribution_7d",
        "distribution_90d": "distribution_90d",
        "distribution_180d": "distribution_180d",
        "overview_30d": "sales_overview",
        "overview_7d": "overview_7d",
        "overview_90d": "overview_90d",
        "overview_180d": "overview_180d",
        "sku_analysis_7d": "sku_analysis_7d",
        "sku_analysis_28d": "sku_analysis_28d",
        "detail_url": "detail_url",
        "platform": "fastmoss",
        "selling_points": None,
        "attributes": None,
        "score": "selection_score",
        "score_reason": "score_reason",
        "sales_mom": "sales_mom",
        "status": "illustration_status",
        "illustration_result_url": "illustration_result_url",
        "illustration_extractable": "illustration_extractable",
        "illustration_extract_reason": "illustration_extract_reason",
    },
    "kalodata": {
        "label": "Kalodata",
        "table": "kalodata_youwei_product",
        "id": "商品ID",
        "date": "date_record",
        "title": "商品标题",
        "image": "商品主图",
        "price": "价格",
        "rating": "商品评分",
        "commission": "佣金比例",
        "sold": "总销量",
        "total_sold": "all_solds",
        "sale_amount": "总成交额",
        "author_count": "关联达人数",
        "base_price": None,
        "transport_fee": None,
        "video_ratio": "近28天视频占比",
        "product_card_ratio": "近28天商品卡占比",
        "distribution_30d": None,
        "distribution_7d": "distribution_7d",
        "distribution_90d": "distribution_90d",
        "distribution_180d": "distribution_180d",
        "overview_30d": None,
        "overview_7d": "overview_7d",
        "overview_90d": "overview_90d",
        "overview_180d": "overview_180d",
        "sku_analysis_7d": "sku_analysis_7d",
        "sku_analysis_28d": "sku_analysis_28d",
        "detail_url": "商品链接",
        "platform": "kalodata",
        "selling_points": "卖点",
        "attributes": "属性信息",
        "score": "selection_score",
        "score_reason": "score_reason",
        "sales_mom": "sales_mom",
        "status": "illustration_status",
        "illustration_result_url": "illustration_result_url",
        "illustration_extractable": "illustration_extractable",
        "illustration_extract_reason": "illustration_extract_reason",
    },
    "pod_cross_category": {
        "label": "POD素材榜",
        "table": "pod_cross_category_product",
        "id": "product_id",
        "date": "date_record",
        "title": "title",
        "image": "image_base64",
        "price": "real_price",
        "rating": "rating",
        "commission": "commission_rate",
        "sold": "rank_sold_count",
        "total_sold": "sold_count",
        "sale_amount": "sale_amount",
        "author_count": "author_count",
        "base_price": "base_price",
        "transport_fee": "transport_fee",
        "video_ratio": None,
        "product_card_ratio": None,
        "distribution_30d": None,
        "distribution_7d": None,
        "distribution_90d": None,
        "distribution_180d": None,
        "overview_30d": None,
        "overview_7d": None,
        "overview_90d": None,
        "overview_180d": None,
        "sku_analysis_7d": "sku_analysis_7d",
        "sku_analysis_28d": "sku_analysis_28d",
        "detail_url": "detail_url",
        "platform": "fastmoss",
        "selling_points": None,
        "attributes": None,
        "score": "selection_score",
        "score_reason": "score_reason",
        "sales_mom": "sales_mom",
        "status": "illustration_status",
        "illustration_result_url": "illustration_result_url",
        "illustration_extractable": "illustration_extractable",
        "illustration_extract_reason": "illustration_extract_reason",
        "category_l1": "category_l1",
        "category_l2": "category_l2",
        "category_l3": "category_l3",
        "aweme_count": "aweme_count",
    },
}

SORT_FIELDS = {
    "date_record": "date_record",
    "sold": "sold_count_view",
    "sale_amount": "sale_amount_view",
    "rating": "rating_view",
    "price": "price_view",
    "author_count": "author_count_view",
    "launch_time": "launch_time",
    "product_card_sales": "product_card_sales_view",
    "product_card_ratio": "product_card_ratio_view",
    "sales_growth": "sales_growth_view",
    "ip_grade": "ip_grade",
}


def create_app():
    app = Flask(__name__, static_folder="static", template_folder="static")
    app.secret_key = os.getenv("SECRET_KEY", "choice-product-dev-secret")
    CORS(app, supports_credentials=True)
    ensure_pod_cross_category_tables()
    ensure_selection_score_columns()
    ensure_detail_analysis_table()
    ensure_expert_team_tables()

    @app.errorhandler(ApiAuthError)
    def handle_auth_error(_error):
        return api_error("未登录", 401)

    @app.errorhandler(ApiPermissionError)
    def handle_permission_error(_error):
        return api_error("没有专家团队使用权限", 403)

    @app.errorhandler(PermissionError)
    def handle_legacy_permission_error(_error):
        return api_error("没有专家团队使用权限", 403)

    @app.get("/")
    def index():
        return render_template("index.html")

    @app.get("/login")
    def login_page():
        return render_template("login.html")

    @app.post("/api/auth/register")
    def register():
        payload = request.get_json(silent=True) or {}
        username = stringify(payload.get("username")).strip()
        password = stringify(payload.get("password"))
        display_name = stringify(payload.get("display_name")).strip() or username
        if not username or not password:
            return api_error("用户名和密码不能为空", 400)
        if len(username) > 64:
            return api_error("用户名过长", 400)
        if len(password) < 6:
            return api_error("密码至少 6 位", 400)

        password_hash = generate_password_hash(password)
        try:
            with db() as conn, conn.cursor() as cursor:
                cursor.execute(
                    """
                    INSERT INTO app_user (username, password_hash, display_name, role, status)
                    VALUES (%s, %s, %s, 'user', 'ACTIVE')
                    """,
                    (username, password_hash, display_name),
                )
                conn.commit()
                user_id = cursor.lastrowid
                cursor.execute(
                    "SELECT id, username, display_name, role, status FROM app_user WHERE id = %s",
                    (user_id,),
                )
                user = cursor.fetchone()
        except pymysql.err.IntegrityError:
            return api_error("用户名已存在", 409)
        session["uid"] = user["id"]
        return api_ok({"user": public_user(user)})

    @app.post("/api/auth/login")
    def login():
        payload = request.get_json(silent=True) or {}
        username = stringify(payload.get("username")).strip()
        password = stringify(payload.get("password"))
        if not username or not password:
            return api_error("用户名和密码不能为空", 400)
        with db() as conn, conn.cursor() as cursor:
            cursor.execute(
                "SELECT id, username, password_hash, display_name, role, status FROM app_user WHERE username = %s",
                (username,),
            )
            user = cursor.fetchone()
        if not user or user.get("status") != "ACTIVE" or not check_password_hash(user.get("password_hash") or "", password):
            return api_error("用户名或密码错误", 401)
        session["uid"] = user["id"]
        return api_ok({"user": public_user(user)})

    @app.get("/api/auth/me")
    def auth_me():
        user = current_user()
        if not user:
            return api_error("未登录", 401)
        return api_ok({"user": public_user(user)})

    @app.post("/api/auth/logout")
    def logout():
        session.clear()
        return api_ok({"logged_out": True})

    @app.get("/api/products")
    def products():
        require_current_user()
        source = normalize_source(request.args.get("source"))
        if source == "pod_cross_category":
            reset_stale_pod_illustration_tasks()
        period = normalize_period(request.args.get("period"))
        sales_period = build_sales_period(
            period,
            request.args.get("sales_start"),
            request.args.get("sales_end"),
        )
        page = clamp_int(request.args.get("page"), 1, 1, 100000)
        page_size = clamp_int(request.args.get("page_size"), 30, 10, 100)
        meta = SOURCES[source]
        date_start_arg = request.args.get("date_start")
        date_end_arg = request.args.get("date_end")
        where, params = build_filters(meta, request.args, latest_by_default=True)
        has_date_filter = bool(date_start_arg or date_end_arg)
        is_single_day_filter = bool(date_start_arg and date_end_arg and date_end_arg == date_start_arg)
        use_latest_join = has_date_filter and not is_single_day_filter
        sort_by = request.args.get("sort_by")
        include_sales_growth_join = sort_by == "sales_growth"
        use_sales_aggregate = (not is_single_day_filter) or include_sales_growth_join
        order_sql = build_order(meta, sort_by, request.args.get("sort_order"), sales_period, use_sales_aggregate, request.args)
        offset = (page - 1) * page_size

        select_sql = build_product_select(source, meta)
        latest_join = build_latest_product_join(meta, where) if use_latest_join else ""
        sales_aggregate_join = build_sales_aggregate_join(meta, where) if use_sales_aggregate else ""
        main_where_sql = "" if latest_join else f"WHERE {' AND '.join(where)}"
        previous_sales_aggregate_join = build_previous_sales_aggregate_join(
            meta,
            date_start_arg,
            date_end_arg,
        ) if use_sales_aggregate and include_sales_growth_join else ""
        sales_delta_join = build_sales_delta_join(
            meta,
            date_start_arg,
            date_end_arg,
        ) if use_sales_aggregate and include_sales_growth_join else ""
        if use_sales_aggregate and include_sales_growth_join:
            runtime_sold_select = (
                "sales_aggregate.aggregate_sold_count AS runtime_sold_count, "
                "previous_sales_aggregate.previous_sold_count AS previous_runtime_sold_count, "
                f"{sales_delta_current_expr(meta)} AS current_period_sold_count, "
                f"{sales_delta_previous_expr()} AS previous_period_sold_count"
            )
        elif use_sales_aggregate:
            runtime_sold_select = (
                "sales_aggregate.aggregate_sold_count AS runtime_sold_count, "
                "NULL AS previous_runtime_sold_count, "
                "NULL AS current_period_sold_count, "
                "NULL AS previous_period_sold_count"
            )
        else:
            runtime_sold_select = (
                "NULL AS runtime_sold_count, NULL AS previous_runtime_sold_count, "
                "NULL AS current_period_sold_count, NULL AS previous_period_sold_count"
            )
        sql = f"""
            SELECT {select_sql}, {runtime_sold_select}
            FROM `{meta["table"]}`
            {latest_join}
            {sales_aggregate_join}
            {previous_sales_aggregate_join}
            {sales_delta_join}
            {main_where_sql}
            {order_sql}
            LIMIT %s OFFSET %s
        """
        count_sql = f"""
            SELECT COUNT(*) AS total FROM (
                SELECT `{meta["id"]}` AS product_id
                FROM `{meta["table"]}`
                WHERE {" AND ".join(where)}
                GROUP BY `{meta["id"]}`
            ) AS latest_count
        """

        with db() as conn, conn.cursor() as cursor:
            cursor.execute(count_sql, params)
            total = int(cursor.fetchone()["total"])
            query_params = []
            if latest_join:
                query_params.extend(params)
            if sales_aggregate_join:
                query_params.extend(params)
            if main_where_sql:
                query_params.extend(params)
            cursor.execute(sql, [*query_params, page_size, offset])
            raw_rows = cursor.fetchall()
            attach_runtime_metrics(
                cursor,
                meta,
                raw_rows,
                sales_period,
                request.args.get("date_start"),
                request.args.get("date_end"),
            )
            rows = [normalize_product_row(source, row, sales_period) for row in raw_rows]

        return api_ok({
            "source": source,
            "period": sales_period["period"],
            "period_label": sales_period["label"],
            "page": page,
            "page_size": page_size,
            "total": total,
            "items": rows,
        })

    @app.get("/api/products/stats")
    def stats():
        require_current_user()
        source = normalize_source(request.args.get("source"))
        if source == "pod_cross_category":
            reset_stale_pod_illustration_tasks()
        meta = SOURCES[source]
        where, params = build_filters(meta, request.args)
        with db() as conn, conn.cursor() as cursor:
            cursor.execute(f"SELECT COUNT(*) AS total FROM `{meta['table']}` WHERE {' AND '.join(where)}", params)
            total = int(cursor.fetchone()["total"])
            if source == "pod_cross_category":
                cursor.execute(
                    f"""
                    SELECT illustration_status AS status, COUNT(*) AS count
                    FROM `{meta['table']}`
                    WHERE {" AND ".join(where)}
                    GROUP BY illustration_status
                    """,
                    params,
                )
                pod_statuses = {str(row["status"] if row["status"] is not None else 0): int(row["count"]) for row in cursor.fetchall()}
                return api_ok({
                    "source": source,
                    "total": total,
                    "pod_statuses": pod_statuses,
                    "statuses": {},
                    "ip_grades": {},
                    "materials": {},
                })
            cursor.execute(
                f"""
                SELECT audit_status, COUNT(*) AS count
                FROM `{meta['table']}`
                WHERE {" AND ".join(where)}
                GROUP BY audit_status
                """,
                params,
            )
            statuses = {row["audit_status"] or "PENDING": int(row["count"]) for row in cursor.fetchall()}
            cursor.execute(
                f"""
                SELECT ip_grade, COUNT(*) AS count
                FROM `{meta['table']}`
                WHERE {" AND ".join(where)} AND ip_grade IS NOT NULL AND ip_grade <> ''
                GROUP BY ip_grade
                """,
                params,
            )
            grades = {row["ip_grade"]: int(row["count"]) for row in cursor.fetchall()}
            try:
                material_counts = count_material_types(cursor, meta["table"], where, params)
            except pymysql.err.OperationalError as exc:
                material_counts = {}
                print(f"[STATS] skip material count table={meta['table']} error={exc}", flush=True)
        return api_ok({
            "source": source,
            "total": total,
            "statuses": statuses,
            "ip_grades": grades,
            "materials": material_counts,
        })

    @app.get("/api/products/latest-date")
    def latest_date():
        require_current_user()
        source = normalize_source(request.args.get("source"))
        meta = SOURCES[source]
        with db() as conn, conn.cursor() as cursor:
            cursor.execute(f"SELECT MAX(`{meta['date']}`) AS latest_date FROM `{meta['table']}`")
            row = cursor.fetchone()
        return api_ok({
            "source": source,
            "latest_date": stringify(row.get("latest_date") if row else None),
        })

    @app.post("/api/products/recalculate-scores")
    def recalculate_scores_api():
        user = require_current_user()
        if stringify(user.get("role")).strip().lower() not in {"admin", "manager"}:
            return api_error("没有评分回算权限", 403)
        payload = request.get_json(silent=True) or {}
        source = normalize_source(payload.get("source") or request.args.get("source")) if (payload.get("source") or request.args.get("source")) else None
        date_record = stringify(payload.get("date_record") or request.args.get("date_record")).strip() or None
        limit = parse_int(payload.get("limit") or request.args.get("limit"), None)
        force = str(payload.get("force") or request.args.get("force") or "").strip().lower() in {"1", "true", "yes"}
        latest_only_arg = payload.get("latest_only", request.args.get("latest_only"))
        latest_only = None if latest_only_arg in {None, ""} else str(latest_only_arg).strip().lower() in {"1", "true", "yes"}
        summary = recalculate_selection_scores(source=source, date_record=date_record, limit=limit, force=force, latest_only=latest_only)
        return api_ok({"summary": summary})

    @app.post("/api/products/pod_cross_category/<product_id>/extract")
    def create_pod_illustration_task(product_id):
        user = require_current_user()
        payload = request.get_json(silent=True) or {}
        style_prompt = stringify(payload.get("style_prompt")).strip()[:255]
        force_regenerate = str(payload.get("force") or request.args.get("force") or "").strip().lower() in {"1", "true", "yes"}
        date_record = stringify(payload.get("date_record") or request.args.get("date_record")).strip()
        print(
            f"[POD_EXTRACT] request product_id={product_id} date_record={date_record or '-'} "
            f"user={user.get('username') or user.get('id')}",
            flush=True,
        )
        meta = SOURCES["pod_cross_category"]
        where = [f"`{meta['id']}` = %s"]
        params = [product_id]
        if date_record:
            where.append(f"`{meta['date']}` = %s")
            params.append(date_record)
        with db() as conn, conn.cursor() as cursor:
            cursor.execute(
                f"""
                SELECT id, `{meta['id']}` AS product_id, `{meta['date']}` AS date_record,
                       {sql_alias(meta.get("title"), "title")},
                       {sql_alias(meta.get("image"), "image_value")},
                       illustration_extractable,
                       illustration_result_url
                FROM `{meta['table']}`
                WHERE {" AND ".join(where)}
                ORDER BY `{meta['date']}` DESC
                LIMIT 1
                """,
                params,
            )
            product = cursor.fetchone()
            if not product:
                print(f"[POD_EXTRACT] product not found product_id={product_id}", flush=True)
                return api_error("POD商品不存在", 404)
            if product.get("illustration_extractable") not in {1, "1", True}:
                print(
                    f"[POD_EXTRACT] rejected not_extractable product_id={product_id} "
                    f"extractable={product.get('illustration_extractable')}",
                    flush=True,
                )
                return api_error("该商品未通过插画可提取性判断，不能生成插画", 400)
            if product.get("illustration_result_url") and not force_regenerate:
                print(
                    f"[POD_EXTRACT] already generated product_id={product_id} "
                    f"url={product.get('illustration_result_url')}",
                    flush=True,
                )
                return api_ok({
                    "task_id": None,
                    "status": 2,
                    "result_image_url": stringify(product.get("illustration_result_url")),
                    "message": "该商品已有插画结果",
                })
            task_id = f"pod_{product_id}_{int(datetime.now().timestamp() * 1000)}"
            task_date = stringify(product.get("date_record"))
            image_url = f"/api/products/pod_cross_category/{product_id}/image?date_record={task_date}"
            print(f"[POD_EXTRACT] create task task_id={task_id} product_id={product_id} date_record={task_date}", flush=True)
            cursor.execute(
                """
                INSERT INTO pod_illustration_task
                  (task_id, product_id, date_record, user_id, username, style_prompt, original_image_url, status)
                VALUES (%s, %s, %s, %s, %s, %s, %s, 1)
                """,
                (task_id, product_id, task_date, user.get("id"), user.get("username"), style_prompt, image_url),
            )
            cursor.execute(
                f"""
                UPDATE `{meta['table']}`
                SET illustration_status = 1
                WHERE `{meta['id']}` = %s AND `{meta['date']}` = %s
                """,
                (product_id, task_date),
            )
            conn.commit()
        try:
            print(f"[POD_EXTRACT] calling minimax task_id={task_id}", flush=True)
            result_url = call_minimax_illustration_generation(
                product.get("image_value"),
                product.get("title") or "",
                style_prompt,
                task_id,
            )
            print('result_',result_url)
            with db() as conn, conn.cursor() as cursor:
                cursor.execute(
                    """
                    UPDATE pod_illustration_task
                    SET result_image_url = %s, status = 2, error_msg = NULL
                    WHERE task_id = %s
                    """,
                    (result_url, task_id),
                )
                cursor.execute(
                    """
                    UPDATE pod_cross_category_product
                    SET illustration_status = 2,
                        illustration_result_url = %s
                    WHERE product_id = %s AND date_record = %s
                    """,
                    (result_url, product_id, task_date),
                )
                conn.commit()
            print(f"[POD_EXTRACT] success task_id={task_id} result_url={result_url}", flush=True)
            return api_ok({
                "task_id": task_id,
                "status": 2,
                "result_image_url": result_url,
                "message": "插画已生成",
            })
        except Exception as exc:
            error_msg = stringify(exc)[:1000]
            print(f"[POD_EXTRACT] failed task_id={task_id} error={error_msg}", flush=True)
            with db() as conn, conn.cursor() as cursor:
                cursor.execute(
                    """
                    UPDATE pod_illustration_task
                    SET status = 3, error_msg = %s
                    WHERE task_id = %s
                    """,
                    (error_msg, task_id),
                )
                cursor.execute(
                    """
                    UPDATE pod_cross_category_product
                    SET illustration_status = 3
                    WHERE product_id = %s AND date_record = %s
                    """,
                    (product_id, task_date),
                )
                conn.commit()
            return api_error(f"插画生成失败：{error_msg}", 502)

    @app.get("/api/products/pod_cross_category/extract-tasks/<task_id>")
    def get_pod_illustration_task(task_id):
        require_current_user()
        with db() as conn, conn.cursor() as cursor:
            cursor.execute(
                """
                SELECT task_id, product_id, date_record, style_prompt, original_image_url, result_image_url, status, error_msg, created_at, updated_at
                FROM pod_illustration_task
                WHERE task_id = %s
                LIMIT 1
                """,
                (task_id,),
            )
            task = cursor.fetchone()
        if not task:
            return api_error("任务不存在", 404)
        return api_ok({key: serialize_value(value) for key, value in task.items()})

    @app.post("/api/products/<source>/<product_id>/illustration-check")
    def check_product_illustration_extractable(source, product_id):
        require_current_user()
        source = normalize_source(source)
        if source != "pod_cross_category":
            return api_error("插画可提取性判断仅用于 POD 跨品类商品", 400)
        payload = request.get_json(silent=True) or {}
        date_record = stringify(payload.get("date_record") or request.args.get("date_record")).strip()
        meta = SOURCES[source]
        where = [f"`{meta['id']}` = %s"]
        params = [product_id]
        if date_record:
            where.append(f"`{meta['date']}` = %s")
            params.append(date_record)

        with db() as conn, conn.cursor() as cursor:
            cursor.execute(
                f"""
                SELECT `{meta['id']}` AS product_id, `{meta['date']}` AS date_record,
                       {sql_alias(meta.get("title"), "title")},
                       {sql_alias(meta.get("image"), "image_value")}
                FROM `{meta['table']}`
                WHERE {" AND ".join(where)}
                ORDER BY `{meta['date']}` DESC
                LIMIT 1
                """,
                params,
            )
            product = cursor.fetchone()
            if not product:
                return api_error("商品不存在", 404)
            if not product.get("image_value"):
                return api_error("商品主图为空，无法判断插画可提取性", 422)

            try:
                result = analyze_illustration_extractability(product.get("image_value"), product.get("title") or "")
            except Exception as exc:
                print(f"[MINIMAX_V2] manual check failed product_id={product_id}: {exc}", flush=True)
                return api_error(f"插画判断服务失败，未写入判断结果：{exc}", 502)
            extractable = 1 if result.get("can_extract_illustration") else 0
            reason = stringify(result.get("reason")).strip()[:500]
            cursor.execute(
                f"""
                UPDATE `{meta['table']}`
                SET illustration_extractable = %s,
                    illustration_extract_reason = %s
                WHERE `{meta['id']}` = %s AND `{meta['date']}` = %s
                """,
                (extractable, reason, product_id, product.get("date_record")),
            )
            conn.commit()

        return api_ok({
            "product_id": product_id,
            "date_record": product.get("date_record"),
            "illustration_extractable": extractable,
            "illustration_extract_reason": reason,
        })

    @app.post("/api/translate-title")
    def translate_title():
        require_current_user()
        payload = request.get_json(silent=True) or {}
        text = stringify(payload.get("text")).strip()
        if not text:
            return api_error("text 不能为空", 400)
        if len(text) > 800:
            return api_error("商品名过长，无法翻译", 400)
        try:
            translation = translate_text_to_chinese(text)
        except RuntimeError as exc:
            translation = fallback_translate_title(text)
            return api_ok({"source": text, "translation": translation, "provider": "local_fallback", "warning": str(exc)})
        except requests.RequestException as exc:
            translation = fallback_translate_title(text)
            return api_ok({"source": text, "translation": translation, "provider": "local_fallback", "warning": str(exc)})
        return api_ok({"source": text, "translation": strip_llm_think_blocks(translation), "provider": "minimax"})

    @app.post("/api/translate-text")
    def translate_text():
        require_current_user()
        payload = request.get_json(silent=True) or {}
        text = stringify(payload.get("text")).strip()
        if not text:
            return api_error("text 不能为空", 400)
        if len(text) > 5000:
            return api_error("文本过长，无法翻译", 400)
        try:
            translation = translate_general_text_to_chinese(text)
        except RuntimeError as exc:
            translation = fallback_translate_title(text)
            return api_ok({"source": text, "translation": translation, "provider": "local_fallback", "warning": str(exc)})
        except requests.RequestException as exc:
            translation = fallback_translate_title(text)
            return api_ok({"source": text, "translation": translation, "provider": "local_fallback", "warning": str(exc)})
        return api_ok({"source": text, "translation": strip_llm_think_blocks(translation), "provider": "minimax"})

    @app.get("/api/products/<source>/<product_id>")
    def product_detail(source, product_id):
        require_current_user()
        source = normalize_source(source)
        meta = SOURCES[source]
        date_record = request.args.get("date_record")
        where = [f"`{meta['id']}` = %s"]
        params = [product_id]
        if date_record:
            where.append(f"`{meta['date']}` = %s")
            params.append(date_record)
        with db() as conn, conn.cursor() as cursor:
            cursor.execute(
                f"""
                SELECT *
                FROM `{meta['table']}`
                WHERE {" AND ".join(where)}
                ORDER BY `{meta['date']}` DESC
                LIMIT 1
                """,
                params,
            )
            row = cursor.fetchone()
            if not row:
                return api_error("商品不存在", 404)
            detail = normalize_detail(source, row)
            enrich_detail_period_fields(cursor, source, meta, row, detail)
        return api_ok(detail)

    @app.post("/api/products/<source>/<product_id>/analysis-detail")
    def product_analysis_detail(source, product_id):
        user = require_current_user()
        source = normalize_source(source)
        payload = request.get_json(silent=True) or {}
        analysis_type = stringify(payload.get("analysis_type")).strip().lower()
        date_record = payload.get("date_record") or request.args.get("date_record")
        if analysis_type not in {"ip", "material"}:
            return api_error("analysis_type 必须是 ip 或 material", 400)
        if not date_record:
            return api_error("date_record 不能为空", 400)

        ai_source = "fastmoss" if source == "unified" else source
        if ai_source not in AI_PRODUCT_SOURCES:
            return api_error(f"暂不支持该数据源的详细分析：{source}", 400)

        with db() as conn, conn.cursor() as cursor:
            record = fetch_detail_analysis_record(cursor, user["id"], source, product_id, date_record, analysis_type)
            if record and record.get("status") == "SUCCESS":
                return api_ok({"cached": True, "record": normalize_detail_analysis_record(record), "result": parse_json(record.get("result_json"))})
            if record and record.get("status") == "RUNNING":
                return api_ok({"cached": False, "record": normalize_detail_analysis_record(record), "status": "RUNNING"})
            record_id = upsert_detail_analysis_record(
                cursor,
                user,
                source,
                AI_PRODUCT_SOURCES[ai_source]["table"],
                product_id,
                date_record,
                analysis_type,
                payload.get("title") or "",
            )
            conn.commit()
        DETAIL_ANALYSIS_EXECUTOR.submit(run_detail_analysis_task, record_id, user["id"], ai_source, product_id, date_record, analysis_type)
        return api_ok({"cached": False, "status": "RUNNING", "record": {"id": record_id, "status": "RUNNING"}})

    @app.get("/api/detail-analyses")
    def detail_analyses():
        user = require_current_user()
        analysis_type = stringify(request.args.get("analysis_type")).strip().lower() or "ip"
        status = stringify(request.args.get("status")).strip().upper()
        if analysis_type not in {"ip", "material"}:
            return api_error("analysis_type 蹇呴』鏄?ip 鎴?material", 400)
        where = ["user_id = %s", "analysis_type = %s"]
        params = [user["id"], analysis_type]
        if status in {"RUNNING", "SUCCESS", "FAILED"}:
            where.append("status = %s")
            params.append(status)
        with db() as conn, conn.cursor() as cursor:
            cursor.execute(
                f"""
                SELECT *
                FROM {DETAIL_ANALYSIS_TABLE}
                WHERE {" AND ".join(where)}
                ORDER BY updated_at DESC, id DESC
                LIMIT 200
                """,
                params,
            )
            rows = [normalize_detail_analysis_record(row) for row in cursor.fetchall()]
        return api_ok({"items": rows, "analysis_type": analysis_type})

    @app.get("/api/expert-team/roles")
    def expert_team_roles():
        require_expert_team_permission()
        return api_ok({"roles": expert_team_roles_payload()})

    @app.get("/api/expert-team/sessions")
    def expert_team_sessions():
        user = require_expert_team_permission()
        with db() as conn, conn.cursor() as cursor:
            cursor.execute(
                f"""
                SELECT id, title, project_code, created_at, updated_at
                FROM {EXPERT_TEAM_SESSION_TABLE}
                WHERE user_id = %s
                ORDER BY updated_at DESC, id DESC
                LIMIT 30
                """,
                (user["id"],),
            )
            sessions = [normalize_expert_session(row) for row in cursor.fetchall()]
        return api_ok({"sessions": sessions})

    @app.get("/api/expert-team/sessions/<int:session_id>/messages")
    def expert_team_messages(session_id):
        user = require_expert_team_permission()
        with db() as conn, conn.cursor() as cursor:
            cursor.execute(
                f"SELECT id FROM {EXPERT_TEAM_SESSION_TABLE} WHERE id = %s AND user_id = %s",
                (session_id, user["id"]),
            )
            if not cursor.fetchone():
                return api_error("专家团队会话不存在", 404)
            cursor.execute(
                f"""
                SELECT id, role, content, team_role, created_at
                FROM {EXPERT_TEAM_MESSAGE_TABLE}
                WHERE session_id = %s AND user_id = %s
                ORDER BY id ASC
                """,
                (session_id, user["id"]),
            )
            messages = [normalize_expert_message(row) for row in cursor.fetchall()]
        return api_ok({"messages": messages})

    @app.delete("/api/expert-team/sessions/<int:session_id>")
    def delete_expert_team_session(session_id):
        user = require_expert_team_permission()
        with db() as conn, conn.cursor() as cursor:
            cursor.execute(
                f"SELECT id FROM {EXPERT_TEAM_SESSION_TABLE} WHERE id = %s AND user_id = %s",
                (session_id, user["id"]),
            )
            if not cursor.fetchone():
                return api_error("专家团队会话不存在", 404)
            cursor.execute(
                f"DELETE FROM {EXPERT_TEAM_MESSAGE_TABLE} WHERE session_id = %s AND user_id = %s",
                (session_id, user["id"]),
            )
            cursor.execute(
                f"DELETE FROM {EXPERT_TEAM_SESSION_TABLE} WHERE id = %s AND user_id = %s",
                (session_id, user["id"]),
            )
            conn.commit()
        return api_ok({"deleted": True, "session_id": session_id})

    @app.get("/api/expert-team/workflows")
    def expert_team_workflows():
        user = require_expert_team_permission()
        session_id = parse_int(request.args.get("session_id"))
        where = ["user_id = %s"]
        params = [user["id"]]
        if session_id:
            where.append("session_id = %s")
            params.append(session_id)
        with db() as conn, conn.cursor() as cursor:
            cursor.execute(
                f"""
                SELECT *
                FROM {EXPERT_WORKFLOW_INSTANCE_TABLE}
                WHERE {" AND ".join(where)}
                ORDER BY updated_at DESC, id DESC
                LIMIT 100
                """,
                params,
            )
            rows = [normalize_expert_workflow_instance(row) for row in cursor.fetchall()]
        return api_ok({"items": rows})

    @app.get("/api/expert-team/workflows/<int:workflow_id>")
    def expert_team_workflow_detail(workflow_id):
        user = require_expert_team_permission()
        with db() as conn, conn.cursor() as cursor:
            cursor.execute(
                f"""
                SELECT *
                FROM {EXPERT_WORKFLOW_INSTANCE_TABLE}
                WHERE id = %s AND user_id = %s
                LIMIT 1
                """,
                (workflow_id, user["id"]),
            )
            workflow = cursor.fetchone()
            if not workflow:
                return api_error("工作流不存在", 404)
            cursor.execute(
                f"""
                SELECT *
                FROM {EXPERT_WORKFLOW_STEP_TABLE}
                WHERE workflow_id = %s
                ORDER BY step_order ASC, id ASC
                """,
                (workflow_id,),
            )
            steps = [normalize_expert_workflow_step(row) for row in cursor.fetchall()]
        return api_ok({"workflow": normalize_expert_workflow_instance(workflow), "steps": steps})

    @app.post("/api/expert-team/chat")
    def expert_team_chat():
        user = require_expert_team_permission()
        payload = request.get_json(silent=True) or {}
        message = stringify(payload.get("message")).strip()
        images = payload.get("images") if isinstance(payload.get("images"), list) else []
        image_count = len(images)
        base_files = normalize_expert_base_files(payload.get("base_files"))
        if not message and image_count:
            message = f"请分析我发送的{image_count}张图片"
        if not message:
            return api_error("请输入要和专家团队讨论的问题", 400)
        session_id = parse_int(payload.get("session_id"))
        project_code = stringify(payload.get("project_code")).strip() or "choice_product"
        project_context = stringify(payload.get("project_context")).strip()

        with db() as conn, conn.cursor() as cursor:
            if session_id:
                cursor.execute(
                    f"SELECT * FROM {EXPERT_TEAM_SESSION_TABLE} WHERE id = %s AND user_id = %s",
                    (session_id, user["id"]),
                )
                if not cursor.fetchone():
                    return api_error("专家团队会话不存在", 404)
            else:
                cursor.execute(
                    f"""
                    INSERT INTO {EXPERT_TEAM_SESSION_TABLE}
                      (user_id, username, title, project_code, project_context)
                    VALUES (%s, %s, %s, %s, %s)
                    """,
                    (
                        user["id"],
                        user.get("username") or "",
                        make_expert_session_title(message),
                        project_code,
                        project_context,
                    ),
                )
                conn.commit()
                session_id = cursor.lastrowid

            cursor.execute(
                f"""
                INSERT INTO {EXPERT_TEAM_MESSAGE_TABLE}
                  (session_id, user_id, role, team_role, content)
                VALUES (%s, %s, 'user', 'user', %s)
                """,
                (session_id, user["id"], message),
            )
            conn.commit()

            cursor.execute(
                f"""
                SELECT role, team_role, content
                FROM {EXPERT_TEAM_MESSAGE_TABLE}
                WHERE session_id = %s AND user_id = %s
                ORDER BY id DESC
                LIMIT 12
                """,
                (session_id, user["id"]),
            )
            history = list(reversed(cursor.fetchall()))

        import requests
        try:
            java_payload = {
                "sessionId": str(session_id),
                "userId": user["id"],
                "message": message,
                "context": {
                    "project_code": project_code,
                    "project_context": project_context,
                    "base_files": base_files,
                    "images": images
                }
            }
            java_res = requests.post("http://localhost:5010/api/agent/chat", json=java_payload, timeout=120)
            java_json = java_res.json()
            if java_json.get("code") == 200 and java_json.get("data"):
                answer = java_json["data"].get("reply") or java_json["data"].get("message") or "专家团队已处理，但未返回文本回复。"
                status = "SUCCESS"
                readonly_context = ""
                ceo_decision = {"note": "Handled by Java Agent (port 5010)"}
                expert_execution = []
                learning_result = {}
                validation_result = {"passed": True}
            else:
                raise Exception(f"Java API error: {java_json.get('message', 'Unknown error')}")
        except Exception as exc:
            answer = sanitize_expert_team_answer(fallback_expert_team_answer(message, str(exc)))
            validation_result = {"passed": False, "fallback": True, "reason": str(exc), "retry_count": 0}
            status = "FALLBACK"
            readonly_context = ""
            ceo_decision = {}
            expert_execution = []
            learning_result = {}

        with db() as conn, conn.cursor() as cursor:
            cursor.execute(
                f"""
                INSERT INTO {EXPERT_TEAM_MESSAGE_TABLE}
                  (session_id, user_id, role, team_role, content, meta_json)
                VALUES (%s, %s, 'assistant', 'chief_planner', %s, %s)
                """,
                (session_id, user["id"], answer, json.dumps({"status": status, "readonly_context": readonly_context, "image_count": image_count, "base_files": base_files, "ceo_decision": ceo_decision, "expert_execution": expert_execution, "learning_result": learning_result, "response_validation": validation_result}, ensure_ascii=False)),
            )
            cursor.execute(
                f"UPDATE {EXPERT_TEAM_SESSION_TABLE} SET updated_at = NOW() WHERE id = %s AND user_id = %s",
                (session_id, user["id"]),
            )
            conn.commit()
            cursor.execute(
                f"SELECT * FROM {EXPERT_TEAM_SESSION_TABLE} WHERE id = %s AND user_id = %s",
                (session_id, user["id"]),
            )
            session_row = cursor.fetchone()
        return api_ok({"session": normalize_expert_session(session_row), "message": answer, "status": status})

    @app.patch("/api/products/<source>/<product_id>/status")
    def update_status(source, product_id):
        require_current_user()
        source = normalize_source(source)
        payload = request.get_json(silent=True) or {}
        audit_status = payload.get("audit_status")
        date_record = payload.get("date_record")
        if audit_status not in STATUS_VALUES:
            return api_error("audit_status 非法", 400)
        if not date_record:
            return api_error("date_record 不能为空", 400)

        meta = SOURCES[source]
        with db() as conn, conn.cursor() as cursor:
            cursor.execute(
                f"""
                UPDATE `{meta['table']}`
                SET audit_status = %s
                WHERE `{meta['id']}` = %s AND `{meta['date']}` = %s
                """,
                (audit_status, product_id, date_record),
            )
            conn.commit()
            affected = cursor.rowcount
        if affected == 0:
            return api_error("未找到需要修改的商品", 404)
        return api_ok({"updated": affected, "audit_status": audit_status})

    @app.get("/api/products/<source>/ready-links")
    def ready_links(source):
        require_current_user()
        source = normalize_source(source)
        links = fetch_ready_links(source)
        return api_ok({"source": source, "count": len(links), "links": links})

    @app.get("/api/products/<source>/export-ready-links")
    def export_ready_links(source):
        require_current_user()
        source = normalize_source(source)
        links = fetch_ready_links(source)
        wb = Workbook()
        ws = wb.active
        ws.title = "ready_links"
        ws.append(["link"])
        for link in links:
            ws.append([link])
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return send_file(
            output,
            as_attachment=True,
            download_name=f"{source}_ready_links.xlsx",
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    @app.get("/api/products/<source>/<product_id>/image")
    def product_image(source, product_id):
        require_current_user()
        source = normalize_source(source)
        meta = SOURCES[source]
        date_record = request.args.get("date_record")
        where = [f"`{meta['id']}` = %s"]
        params = [product_id]
        if date_record:
            where.append(f"`{meta['date']}` = %s")
            params.append(date_record)
        with db() as conn, conn.cursor() as cursor:
            cursor.execute(
                f"""
                SELECT `{meta['image']}` AS image_value
                FROM `{meta['table']}`
                WHERE {" AND ".join(where)}
                ORDER BY `{meta['date']}` DESC
                LIMIT 1
                """,
                params,
            )
            row = cursor.fetchone()
        if not row or not row.get("image_value"):
            return api_error("图片不存在", 404)
        return serve_image_value(row["image_value"])

    @app.after_request
    def add_cache_headers(response):
        if request.path in {"/", "/login"} or response.content_type.startswith("text/html"):
            response.headers["Cache-Control"] = "no-cache, no-store, must-revalidate"
            response.headers["Pragma"] = "no-cache"
            response.headers["Expires"] = "0"
        return response
    return app


def db():
    return pymysql.connect(**DB_CONFIG)


def safe_rollback(conn):
    try:
        conn.rollback()
    except Exception:
        pass


def safe_close(conn):
    try:
        conn.close()
    except Exception:
        pass


def translate_text_to_chinese(text):
    api_key = resolve_minimax_api_key()
    if not api_key:
        raise RuntimeError("缺少 MINIMAX_API_KEY，无法翻译商品名")
    api_key = stringify(api_key).strip().removeprefix("Bearer ").strip()
    url = f"{MINIMAX_BASE_URL.rstrip('/')}/chat/completions"
    payload = {
        "model": MINIMAX_MODEL,
        "messages": [
            {
                "role": "system",
                "content": "你是电商商品标题翻译助手。只输出简体中文译文，不要解释。",
            },
            {
                "role": "user",
                "content": f"请把下面的英文商品名翻译成简体中文，保留品牌名、型号和关键规格：\n{text}",
            },
        ],
        "temperature": 0,
        "stream": False,
    }
    
    data = call_minimax_api_with_retry(url, api_key, payload, timeout=30)
    try:
        translation = strip_llm_think_blocks(data["choices"][0]["message"]["content"])
    except (KeyError, IndexError, TypeError) as exc:
        raise RuntimeError("翻译服务返回格式异常") from exc
    if not translation:
        raise RuntimeError("翻译服务返回为空")
    return translation


def translate_general_text_to_chinese(text):
    api_key = resolve_minimax_api_key()
    if not api_key:
        raise RuntimeError("缺少 MINIMAX_API_KEY，无法翻译文本")
    api_key = stringify(api_key).strip().removeprefix("Bearer ").strip()
    url = f"{MINIMAX_BASE_URL.rstrip('/')}/chat/completions"
    payload = {
        "model": MINIMAX_MODEL,
        "messages": [
            {
                "role": "system",
                "content": "你是电商商品信息翻译助手。只输出简体中文译文，不要解释。保留品牌名、型号、规格、数字 and 专有名词。",
            },
            {
                "role": "user",
                "content": f"请把下面的商品卖点或商品描述翻译成简体中文，保持分隔符和关键信息清晰：\n{text}",
            },
        ],
        "temperature": 0,
        "stream": False,
    }
    
    data = call_minimax_api_with_retry(url, api_key, payload, timeout=30)
    try:
        translation = strip_llm_think_blocks(data["choices"][0]["message"]["content"])
    except (KeyError, IndexError, TypeError) as exc:
        raise RuntimeError("翻译服务返回格式异常") from exc
    if not translation:
        raise RuntimeError("翻译服务返回为空")
    return translation


def strip_llm_think_blocks(text):
    cleaned = stringify(text)
    cleaned = re.sub(r"<think>[\s\S]*?</think>", "", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"^\s*<think>[\s\S]*$", "", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"^\s*思考[:：][\s\S]*?(?=\n\s*(结论|翻译|译文)[:：]|\Z)", "", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"^\s*(结论|翻译|译文)[:：]\s*", "", cleaned.strip(), flags=re.IGNORECASE)
    return cleaned.strip()


def normalize_vision_image_url(image_value):
    value = stringify(image_value).strip()
    if not value:
        return ""
    if value.startswith("data:image/") or value.startswith("http://") or value.startswith("https://"):
        return value
    return "data:image/jpeg;base64," + value


def decode_image_value(image_value):
    value = stringify(image_value).strip()
    if not value:
        return None
    if value.startswith("data:image/") and "," in value:
        value = value.split(",", 1)[1]
    try:
        return Image.open(io.BytesIO(base64.b64decode(value, validate=False))).convert("RGB")
    except Exception:
        return None


def local_illustration_extractability_hint(image_value, title=""):
    title_text = stringify(title).lower()
    positive_terms = (
        "tattoo", "sticker", "skull", "skeleton", "graphic", "print", "printed",
        "pattern", "illustration", "embroidered", "embroidery", "boho",
        "bohemian", "dog", "dogs", "animal", "floral", "flower", "halloween",
        "cartoon", "tribal", "camo", "geometric", "all-over print",
    )
    negative_terms = (
        "plain", "solid", "minimalist", "basic", "ribbed", "knit", "zip",
        "zipper", "drawstring", "pocket", "pockets", "fleece", "pullover",
        "oversized", "puffy", "quilted",
    )
    has_positive_title = any(term in title_text for term in positive_terms)
    has_negative_title = any(term in title_text for term in negative_terms)
    if has_negative_title and not has_positive_title:
        return None
    image = decode_image_value(image_value)
    if image is None:
        return None
    gray = image.resize((240, 240)).convert("L")
    pixels = list(gray.getdata())
    dark_ratio = sum(1 for value in pixels if value < 85) / len(pixels)
    edge = gray.filter(ImageFilter.FIND_EDGES)
    edge_pixels = list(edge.getdata())
    edge_ratio = sum(1 for value in edge_pixels if value > 45) / len(edge_pixels)
    if has_positive_title and dark_ratio >= 0.012 and edge_ratio >= 0.025:
        return {
            "can_extract_illustration": True,
            "reason": (
                "本地兜底判断：标题和主图特征显示商品表面存在可迁移的印花/图案元素，"
                f"深色图案占比 {dark_ratio:.1%}，边缘密度 {edge_ratio:.1%}。"
            ),
        }
    return None


def parse_llm_json_object(raw_text):
    if not raw_text:
        return {}
    text = strip_llm_think_blocks(raw_text)
    if text.startswith("```"):
        text = text.strip("`")
        if text.lower().startswith("json"):
            text = text[4:].strip()
    start = text.find("{")
    end = text.rfind("}")
    if start >= 0 and end >= start:
        text = text[start:end + 1]
    try:
        return json.loads(text)
    except Exception:
        return {}


def call_minimax_api_with_retry(url, api_key, payload, timeout=90, max_retries=3):
    headers = {"Authorization": f"Bearer {api_key}", "Content-Type": "application/json"}
    last_exc = None
    session = requests.Session()
    session.trust_env = False
    for attempt in range(max_retries):
        try:
            response = session.post(url, headers=headers, json=payload, timeout=timeout)
            if response.status_code == 200:
                try:
                    return response.json()
                except ValueError as exc:
                    raise RuntimeError(f"MiniMax API {url} returned non-JSON response") from exc
            
            print(f"[MINIMAX_API] {url} attempt {attempt+1} error {response.status_code}: {response.text}", flush=True)
            if response.status_code in [429, 500, 502, 503, 504]:
                time.sleep(2 * (attempt + 1))
                continue
            response.raise_for_status()
        except requests.exceptions.RequestException as e:
            last_exc = e
            print(f"[MINIMAX_API] {url} attempt {attempt+1} network error: {e}", flush=True)
            if attempt < max_retries - 1:
                time.sleep(2 * (attempt + 1))
                continue
    raise last_exc or RuntimeError(f"MiniMax API {url} failed after {max_retries} attempts")


def extract_minimax_message_content(data, context="MiniMax"):
    if not isinstance(data, dict):
        raise RuntimeError(f"{context} returned invalid response type: {type(data).__name__}")
    choices = data.get("choices")
    if not isinstance(choices, list) or not choices:
        raise RuntimeError(f"{context} returned no choices: {stringify(data)[:300]}")
    first_choice = choices[0]
    if not isinstance(first_choice, dict):
        raise RuntimeError(f"{context} returned invalid choice: {stringify(first_choice)[:300]}")
    message = first_choice.get("message")
    if not isinstance(message, dict):
        raise RuntimeError(f"{context} returned no message: {stringify(first_choice)[:300]}")
    content = message.get("content")
    if content is None:
        raise RuntimeError(f"{context} returned empty content: {stringify(message)[:300]}")
    content = stringify(content).strip()
    if not content:
        raise RuntimeError(f"{context} returned blank content")
    return content


def analyze_illustration_extractability(image_value, title=""):
    api_key = resolve_minimax_api_key()
    if not api_key: raise RuntimeError("缺少 MINIMAX_API_KEY")
    prompt = {
        "system": (
            "你是商品视觉素材提取判断助手。必须结合商品主图和标题判断，"
            "只输出 JSON，不要输出 Markdown。你只判断图案是否适合提取迁移，"
            "不判断版权风险。必须按“能否从商品表面剥离出可复用装饰图案”判断，"
            "不能把商品本体、服装版型、包型结构、拉链、帽绳、口袋、缝线、普通面料纹理当成可提取图案。"
        ),
        "user_text": json.dumps(
            {
                "task": "判断商品主图里是否存在适合迁移到手机壳设计的可提取图案。",
                "product": {"title": title or ""},
                "criteria": [
                    "返回 true 的前提：商品表面必须有可从商品本体中剥离出来的清晰非品牌装饰图案，例如插画、重复印花、动物/花卉/几何图案、刺绣图案、骷髅/节日图案、明确的大面积装饰印花。",
                    "即使图片里有人物、模特或真实商品，只要商品表面的装饰图案本身足够清楚，也可以返回 true。",
                    "返回 false：纯色卫衣/纯色衣服/纯色包、普通商品结构、服装版型、包型结构、拉链、帽绳、口袋、缝线、褶皱、阴影、普通针织/罗纹/绗缝/磨毛/纯面料纹理都不算可提取图案。",
                    "返回 false：只有品牌 Logo、商标文字、广告文字、商品标签，或图案太小/太糊/被严重遮挡。",
                    "如果主图看起来只是一个空白或纯色商品，即使标题包含 hoodie、sweatshirt、shirt、bag 等品类词，也必须返回 false。",
                ],
                "required_json": {
                    "can_extract_illustration": "true/false",
                    "reason": "一句话说明判断依据；如果返回 true，必须点名主图里可剥离的具体图案元素；如果是纯色/结构/面料，返回 false 并说明没有可提取装饰图案",
                },
            },
            ensure_ascii=False,
        ),
    }
    try:
        content = call_product_vision_minimax(prompt, image_value)
    except Exception as exc:
        local_hint = local_illustration_extractability_hint(image_value, title)
        if local_hint:
            return local_hint
        raise RuntimeError(f"插画视觉判断接口调用失败: {exc}") from exc
    parsed = parse_product_json_response(content)
    if "can_extract_illustration" not in parsed:
        raise RuntimeError(f"插画判断服务返回格式异常: {content[:300]}")
    reason = stringify(parsed.get("reason"))
    if not parsed.get("can_extract_illustration") and any(
        marker in reason for marker in ("未提供商品主图", "无法看到图片", "无法查看图片", "未显示商品主图")
    ):
        local_hint = local_illustration_extractability_hint(image_value, title)
        if local_hint:
            return local_hint
    
    return {
        "can_extract_illustration": bool(parsed.get("can_extract_illustration")),
        "reason": reason,
    }

def generate_artwork_description(api_key, image_url):
    system_prompt = '''角色设定 (Role):
你现在是一位顶级的数字资产提取专家和高级纹理艺术家。你的任务是敏锐地观察用户上传的产品参考图，并为图像生成模型（如 Midjourney, DALL-E, 或 Stable Diffusion）撰写极其精确的英文提示词（Prompt），目的是将附着在 3D 物品上的平面图案完美剥离出来。

工作流 (Workflow):
视觉解构： 忽略所有 3D 结构（包的形状、衣服的褶皱、人物、背景、光影）。只盯住“印刷图案”本身。
细节提取： 准确识别图案的风格、核心元素（如花朵、几何体、Logo）、排列方式（单图居中还是无缝平铺铺满）、精确的颜色组成以及底色。
排除干扰： 敏锐识别出必须去除的元素（如水印、品牌 Logo、缝线、拉链、反光）。
输出提示词： 根据以上分析，输出一段用于生成平面资产的英文提示词。

提示词撰写规则 (Rules for Prompt Writing):
首句定调： 必须以 "A detailed, high-resolution photo of a flat, seamless textile pattern swatch..." 开头。
详尽描述图案： 用专业的视觉词汇描述提取出的图案细节。
强化否定指令： 明确指出“不要什么”。
格式要求： 只输出最终的英文提示词本身，不需要任何解释。'''.strip()

    prompt = {
        "system": system_prompt,
        "user_text": "请观察这张图片并直接输出英文提示词（Prompt）。",
    }

    print(f"[MINIMAX_VISION] Requesting artwork description for image...", flush=True)
    try:
        raw_answer = call_product_vision_minimax(prompt, image_url)
        return stringify(raw_answer).strip()
    except Exception as exc:
        print(f"[MINIMAX_VISION] Artwork description error: {exc}", flush=True)
        raise RuntimeError("无法获取原图的插画描述") from exc

def build_illustration_generation_prompt(artwork_description="", retry_feedback=""):
    feedback_text = stringify(retry_feedback).strip()
    feedback_part = f"\nPrevious failed because: {feedback_text[:260]}. Fix this exactly in this attempt." if feedback_text else ""
    
    if artwork_description:
        prompt = artwork_description
    else:
        # 这是一个兜底的极简提示词，防止由于某种原因视觉模型没返回内容
        prompt = "A detailed, high-resolution photo of a flat, seamless textile pattern swatch, precisely extracted from the printed design on the product in the reference image. The output must be a single, flat, repeating tile asset filling the entire square canvas. The texture is that of a scanned fabric sample, completely smooth and devoid of any three-dimensional product contours, creases, hardware, or the product shape itself. Do not generate a new product, model, or scene."
    
    # 强制加上基础铁律，防止视觉模型生成的提示词遗漏了“扁平化”和“去产品特征”的要求
    base_rules = "\nCRITICAL INSTRUCTIONS: The output MUST be a flat 2D asset. It must be completely devoid of 3D contours, creases, hardware, handles, straps, models, shadows, and watermarks. Do NOT generate the product shape (e.g. do not generate a bag or a shirt)."
    
    if base_rules not in prompt:
        prompt += base_rules
        
    prompt += feedback_part
    return prompt[:1450]


def image_bytes_to_data_url(image_bytes, mime_type="image/jpeg"):
    return f"data:{mime_type};base64,{base64.b64encode(image_bytes).decode('ascii')}"


def save_generated_illustration_bytes(image_bytes, task_id):
    safe_task_id = re.sub(r"[^A-Za-z0-9_.-]+", "_", stringify(task_id))[:80] or f"illustration_{int(time.time())}"
    filename = f"{safe_task_id}.jpeg"
    GENERATED_ILLUSTRATION_DIR.mkdir(parents=True, exist_ok=True)
    output_path = GENERATED_ILLUSTRATION_DIR / filename
    with output_path.open("wb") as handle:
        handle.write(image_bytes)
    return f"{GENERATED_ILLUSTRATION_URL_PREFIX}/{filename}"


def call_minimax_illustration_generation_once(api_key, image_file, prompt, task_id="illustration", attempt=1):
    image_file = normalize_vision_image_url(image_file)
    if not image_file:
        raise RuntimeError("商品主图为空，无法生成插画")
    payload = {
        "model": MINIMAX_IMAGE_MODEL,
        "prompt": prompt,
        "aspect_ratio": "1:1",
        "subject_reference": [
            {
                "type": MINIMAX_IMAGE_SUBJECT_TYPE,
                "image_file": image_file,
            }
        ],
        "response_format": "base64",
        "n": 1,
    }
    print(
        f"[MINIMAX_IMAGE] request task_id={task_id} attempt={attempt} model={MINIMAX_IMAGE_MODEL} "
        f"prompt_len={len(prompt)} reference={'url' if image_file.startswith(('http://', 'https://')) else 'base64'}",
        flush=True,
    )
    
    data = call_minimax_api_with_retry(MINIMAX_IMAGE_GENERATION_URL, api_key, payload, timeout=180)
    response_data = data.get("data") or {}
    print(
        f"[MINIMAX_IMAGE] response task_id={task_id} attempt={attempt} "
        f"data_keys={list(response_data.keys())}",
        flush=True,
    )
    base64_images = response_data.get("image_base64") or response_data.get("image_base64s") or []
    url_images = (
        response_data.get("image_urls")
        or response_data.get("image_url")
        or response_data.get("images")
        or response_data.get("urls")
        or []
    )
    if isinstance(base64_images, str):
        base64_images = [base64_images]
    if isinstance(url_images, str):
        url_images = [url_images]

    image_bytes = b""
    if base64_images:
        image_bytes = base64.b64decode(base64_images[0], validate=False)
    elif url_images:
        first_url = url_images[0]
        if isinstance(first_url, dict):
            first_url = first_url.get("url") or first_url.get("image_url") or first_url.get("image")
        first_url = stringify(first_url).strip()
        if not first_url:
            raise RuntimeError("MiniMax 图生图返回了空图片链接")
        session = requests.Session()
        session.trust_env = False
        image_response = session.get(first_url, timeout=120)
        image_response.raise_for_status()
        image_bytes = image_response.content
    else:
        error_hint = stringify(data.get("base_resp") or response_data.get("base_resp") or data)[:500]
        raise RuntimeError(f"MiniMax 图生图未返回图片数据：{error_hint}")
    return image_bytes


def evaluate_generated_illustration(api_key, source_image_file, generated_bytes, task_id="illustration"):
    generated_image_url = image_bytes_to_data_url(generated_bytes)
    user_text = "Compare the artwork in Image 2 with the original in Image 1. Return JSON with 'match_score' (0-100) and 'reason'."
    
    payload = {
        "model": MINIMAX_MODEL,
        "messages": [
            {
                "role": "user",
                "content": [
                    {"type": "text", "text": user_text},
                    {"type": "image_url", "image_url": {"url": source_image_file}},
                    {"type": "image_url", "image_url": {"url": generated_image_url}}
                ]
            }
        ]
    }
    
    try:
        data = call_minimax_api_with_retry("https://api.minimax.io/v1/text/chatcompletion_v2", api_key, payload)
        content = data["choices"][0]["message"]["content"]
        parsed = parse_llm_json_object(content)
    except (KeyError, IndexError, TypeError, Exception) as e:
        print(f"[MINIMAX_V2] Eval Error: {e}", flush=True)
        # 质检失败时，默认给 0 分
        parsed = {"match_score": 0, "reason": f"质检接口调用失败: {e}"}
        
    score = float(parsed.get("match_score", 0))
    return {
        "pass": score >= MINIMAX_ILLUSTRATION_PASS_SCORE,
        "is_same": score >= 80,
        "score": score,
        "feedback": parsed.get("reason", ""),
    }

def call_minimax_illustration_generation(image_value, title="", style_prompt="", task_id="illustration"):
    api_key = resolve_minimax_api_key()
    if not api_key:
        raise RuntimeError("缺少 MINIMAX_API_KEY，无法调用 MiniMax 图生图")
    api_key = stringify(api_key).strip().removeprefix("Bearer ").strip()
    image_file = normalize_vision_image_url(image_value)
    if not image_file:
        raise RuntimeError("商品主图为空，无法生成插画")

    # 阶段一：用大模型查看原图，生成详细描述
    print(f"[MINIMAX_IMAGE_QC] task_id={task_id} generating artwork description...", flush=True)
    try:
        artwork_description = generate_artwork_description(api_key, image_file)
        print(f"image_description task_id={task_id}: {artwork_description}", flush=True)
    except Exception as e:
        print(f"[MINIMAX_IMAGE_QC] task_id={task_id} failed to get description: {e}", flush=True)
        artwork_description = ""

    best_bytes = b""
    best_score = -1.0
    best_feedback = ""
    retry_feedback = ""
    for attempt in range(1, MINIMAX_ILLUSTRATION_MAX_ATTEMPTS + 1):
        prompt = build_illustration_generation_prompt(artwork_description, retry_feedback)
        image_bytes = call_minimax_illustration_generation_once(api_key, image_file, prompt, task_id, attempt)
        evaluation = evaluate_generated_illustration(api_key, image_file, image_bytes, task_id)
        print(
            f"[MINIMAX_IMAGE_QC] task_id={task_id} attempt={attempt} "
            f"pass={evaluation['pass']} score={evaluation['score']} feedback={evaluation['feedback']}",
            flush=True,
        )
        if evaluation["score"] > best_score:
            best_score = evaluation["score"]
            best_bytes = image_bytes
            best_feedback = evaluation["feedback"]
        if evaluation["pass"]:
            return save_generated_illustration_bytes(image_bytes, task_id)
        retry_feedback = evaluation["feedback"]
    if best_bytes:
        print(
            f"[MINIMAX_IMAGE_QC] task_id={task_id} all attempts failed, best_score={best_score} "
            f"feedback={best_feedback}",
            flush=True,
        )
        raise RuntimeError(f"插画生成质检未通过，最佳得分 {best_score:.0f}：{best_feedback}")
    raise RuntimeError("MiniMax 图生图未生成可保存的图片")


TITLE_TRANSLATION_PHRASES = [
    ("wireless charging", "无线充电"),
    ("magsafe compatible", "兼容 MagSafe"),
    ("shockproof", "防摔"),
    ("drop resistant", "防摔"),
    ("anti drop", "防摔"),
    ("screen protector", "屏幕保护膜"),
    ("camera lens protector", "摄像头镜头保护膜"),
    ("camera protection", "摄像头保护"),
    ("full coverage", "全覆盖"),
    ("protective film", "保护膜"),
    ("tempered glass", "钢化玻璃"),
    ("phone case", "手机壳"),
    ("case cover", "保护壳"),
    ("protective cover", "保护壳"),
    ("airbag", "气囊"),
    ("magnetic", "磁吸"),
    ("transparent", "透明"),
    ("clear", "透明"),
    ("glitter", "闪粉"),
    ("sparkly", "闪亮"),
    ("aesthetic", "高颜值"),
    ("solid color", "纯色"),
    ("jelly", "果冻质感"),
    ("matte", "磨砂"),
    ("soft", "柔软"),
    ("hard", "硬壳"),
    ("slim", "轻薄"),
    ("durable", "耐用"),
    ("compatible with", "适用于"),
    ("suitable for", "适用于"),
    ("for iphone", "适用于 iPhone"),
    ("for samsung", "适用于 Samsung"),
    ("women", "女士"),
    ("men", "男士"),
    ("cover", "保护壳"),
    ("case", "壳"),
    ("protector", "保护膜"),
    ("charging", "充电"),
    ("privacy", "防窥"),
    ("transparent", "透明"),
    ("black", "黑色"),
    ("white", "白色"),
    ("pink", "粉色"),
    ("blue", "蓝色"),
    ("green", "绿色"),
    ("purple", "紫色"),
    ("red", "红色"),
    ("gray", "灰色"),
    ("grey", "灰色"),
    ("gold", "金色"),
    ("silver", "银色"),
    ("brown", "棕色"),
]


def fallback_translate_title(text):
    translated = stringify(text)
    for english, chinese in TITLE_TRANSLATION_PHRASES:
        translated = re.sub(rf"\b{re.escape(english)}\b", chinese, translated, flags=re.IGNORECASE)
    translated = re.sub(r"\s+", " ", translated).strip()
    translated = re.sub(r"\s+([,;:)\]])", r"\1", translated)
    translated = re.sub(r"([(\[])\s+", r"\1", translated)
    translated = re.sub(r"\s*-\s*", " - ", translated)
    return translated or stringify(text)


def api_ok(data):
    return jsonify({"ok": True, "data": data})


def api_error(message, status):
    return jsonify({"ok": False, "error": message}), status


def current_user():
    uid = session.get("uid")
    if not uid:
        return None
    with db() as conn, conn.cursor() as cursor:
        cursor.execute(
            "SELECT id, username, display_name, role, status FROM app_user WHERE id = %s",
            (uid,),
        )
        user = cursor.fetchone()
    if not user or user.get("status") != "ACTIVE":
        return None
    return user


def require_current_user():
    user = current_user()
    if not user:
        raise ApiAuthError()
    return user


def require_expert_team_permission():
    user = require_current_user()
    if not has_user_permission(user, PERMISSION_EXPERT_TEAM):
        return_error = api_error("没有专家团队使用权限", 403)
        raise PermissionError(return_error)
    return user


def has_user_permission(user, permission):
    if permission == PERMISSION_EXPERT_TEAM:
        return stringify((user or {}).get("role")).strip().lower() in EXPERT_TEAM_ALLOWED_ROLES
    return False


def parse_int(value, default=None):
    try:
        if value is None or value == "":
            return default
        return int(value)
    except (TypeError, ValueError):
        return default


def ensure_pod_cross_category_tables():
    with db() as conn, conn.cursor() as cursor:
        cursor.execute(
            """
            CREATE TABLE IF NOT EXISTS pod_cross_category_product (
              id BIGINT NOT NULL AUTO_INCREMENT,
              product_id BIGINT NOT NULL,
              date_record VARCHAR(10) NOT NULL DEFAULT '',
              title VARCHAR(500) DEFAULT NULL,
              rank_sold_count INT DEFAULT NULL,
              rating DECIMAL(3,2) DEFAULT NULL,
              region VARCHAR(10) DEFAULT NULL,
              currency VARCHAR(10) DEFAULT NULL,
              base_price DECIMAL(10,2) DEFAULT NULL,
              real_price VARCHAR(50) DEFAULT NULL,
              original_price VARCHAR(50) DEFAULT NULL,
              launch_time DATETIME DEFAULT NULL,
              author_count INT DEFAULT NULL,
              aweme_count INT DEFAULT NULL,
              live_count INT DEFAULT NULL,
              sold_count BIGINT DEFAULT NULL,
              sale_amount DECIMAL(15,2) DEFAULT NULL,
              review_count INT DEFAULT NULL,
              commission_rate VARCHAR(20) DEFAULT NULL,
              category_l1 VARCHAR(100) DEFAULT NULL,
              category_l2 VARCHAR(100) DEFAULT NULL,
              category_l3 VARCHAR(100) DEFAULT NULL,
              stock_count INT DEFAULT NULL,
              viral_index INT DEFAULT NULL,
              popularity_index INT DEFAULT NULL,
              country_rank VARCHAR(50) DEFAULT NULL,
              category_rank VARCHAR(50) DEFAULT NULL,
              region_name VARCHAR(50) DEFAULT NULL,
              transport_fee VARCHAR(50) DEFAULT NULL,
              detail_url TEXT DEFAULT NULL,
              cover_list LONGTEXT DEFAULT NULL,
              channel_ratio TEXT DEFAULT NULL,
              channel_gmv_ratio TEXT DEFAULT NULL,
              content_ratio TEXT DEFAULT NULL,
              `投放_ratio` TEXT DEFAULT NULL,
              sales_overview LONGTEXT DEFAULT NULL,
              fan_distribution LONGTEXT DEFAULT NULL,
              author_type_distribution LONGTEXT DEFAULT NULL,
              ad_analysis LONGTEXT DEFAULT NULL,
              comments TEXT DEFAULT NULL,
              image_base64 LONGTEXT DEFAULT NULL,
              overview_7d LONGTEXT DEFAULT NULL,
              distribution_7d LONGTEXT DEFAULT NULL,
              overview_90d LONGTEXT DEFAULT NULL,
              distribution_90d TEXT DEFAULT NULL,
              overview_180d LONGTEXT DEFAULT NULL,
              distribution_180d TEXT DEFAULT NULL,
              ad_analysis_7d LONGTEXT DEFAULT NULL,
              ad_analysis_90d LONGTEXT DEFAULT NULL,
              ai_learned INT DEFAULT 0,
              sku_analysis_7d LONGTEXT DEFAULT NULL,
              sku_analysis_28d LONGTEXT DEFAULT NULL,
              selection_score DECIMAL(6,2) DEFAULT NULL,
              score_reason TEXT DEFAULT NULL,
              sales_mom DECIMAL(8,2) DEFAULT NULL,
              illustration_status TINYINT DEFAULT 0,
              illustration_result_url VARCHAR(512) DEFAULT NULL,
              score DECIMAL(8,2) DEFAULT NULL,
              status TINYINT DEFAULT 0,
              audit_status VARCHAR(20) DEFAULT 'PENDING',
              ip_grade VARCHAR(10) DEFAULT NULL,
              ip_reason TEXT DEFAULT NULL,
              ip_tags LONGTEXT DEFAULT NULL,
              material_analysis LONGTEXT DEFAULT NULL,
              ai_analysis_error LONGTEXT DEFAULT NULL,
              PRIMARY KEY (id),
              UNIQUE KEY uk_pod_product_date (product_id, date_record),
              KEY idx_pod_date_score (date_record, score),
              KEY idx_pod_selection_score_date (date_record, selection_score),
              KEY idx_pod_sales_mom_date (date_record, sales_mom),
              KEY idx_pod_illustration_status (illustration_status),
              KEY idx_pod_date_category_score (date_record, category_l1, category_l2, category_l3, score),
              KEY idx_pod_status (status),
              KEY idx_pod_audit_status (audit_status)
            ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_general_ci COMMENT='POD跨品类选品商品表'
            """
        )
        cursor.execute(
            """
            CREATE TABLE IF NOT EXISTS pod_illustration_task (
              task_id VARCHAR(64) NOT NULL,
              product_id BIGINT NOT NULL,
              date_record VARCHAR(10) DEFAULT '',
              user_id BIGINT DEFAULT NULL,
              username VARCHAR(100) DEFAULT NULL,
              style_prompt VARCHAR(255) DEFAULT NULL,
              original_image_url VARCHAR(512) DEFAULT NULL,
              result_image_url VARCHAR(512) DEFAULT NULL,
              status TINYINT NOT NULL DEFAULT 0,
              error_msg TEXT DEFAULT NULL,
              created_at DATETIME DEFAULT CURRENT_TIMESTAMP,
              updated_at DATETIME DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP,
              PRIMARY KEY (task_id),
              KEY idx_pod_task_product (product_id, date_record),
              KEY idx_pod_task_status (status, updated_at)
            ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_general_ci COMMENT='POD插画提取任务表'
            """
        )
        cursor.execute("SHOW INDEX FROM pod_cross_category_product")
        existing_indexes = {row["Key_name"] for row in cursor.fetchall()}
        if "idx_pod_illustration_queue" not in existing_indexes:
            cursor.execute(
                """
                CREATE INDEX idx_pod_illustration_queue
                ON pod_cross_category_product(date_record, illustration_extractable, id)
                """
            )
        conn.commit()


def ensure_selection_score_columns():
    field_definitions = {
        "selection_score": "DECIMAL(6,2) DEFAULT NULL COMMENT '选品综合分'",
        "score_reason": "TEXT DEFAULT NULL COMMENT '选品打分原因'",
        "sales_mom": "DECIMAL(8,2) DEFAULT NULL COMMENT '销售环比百分比'",
        "score_attempts": "TINYINT NOT NULL DEFAULT 0 COMMENT '选品分回算次数'",
        "score_updated_at": "DATETIME DEFAULT NULL COMMENT '选品分最后回算时间'",
        "illustration_status": "TINYINT DEFAULT 0 COMMENT '插画提取状态：0待处理 1处理中 2成功 3失败 4废弃'",
        "illustration_result_url": "VARCHAR(512) DEFAULT NULL COMMENT '插画提取结果图链接'",
        "illustration_extractable": "TINYINT DEFAULT NULL COMMENT '是否适合提取为手机壳插画：1可提取 0不可提取 NULL未判断'",
        "illustration_extract_reason": "VARCHAR(500) DEFAULT NULL COMMENT '插画可提取性判断原因'",
    }
    index_definitions = {
        "idx_selection_score_date": "(date_record, selection_score)",
        "idx_score_queue": "(selection_score, score_attempts, date_record)",
        "idx_sales_mom_date": "(date_record, sales_mom)",
        "idx_illustration_status": "(illustration_status)",
    }
    tables = sorted({meta["table"] for meta in SOURCES.values()})
    with db() as conn, conn.cursor() as cursor:
        for table in tables:
            cursor.execute(
                """
                SELECT COLUMN_NAME
                FROM information_schema.COLUMNS
                WHERE TABLE_SCHEMA = DATABASE() AND TABLE_NAME = %s
                """,
                (table,),
            )
            existing_columns = {row["COLUMN_NAME"] for row in cursor.fetchall()}
            for field, definition in field_definitions.items():
                if field not in existing_columns:
                    cursor.execute(f"ALTER TABLE `{table}` ADD COLUMN `{field}` {definition}")
                    existing_columns.add(field)

            if SCORE_AUTO_CREATE_INDEXES:
                cursor.execute(f"SHOW INDEX FROM `{table}`")
                existing_indexes = {row["Key_name"] for row in cursor.fetchall()}
                for index_name, index_columns in index_definitions.items():
                    if index_name not in existing_indexes:
                        cursor.execute(f"CREATE INDEX `{index_name}` ON `{table}` {index_columns}")
        conn.commit()


def ensure_detail_analysis_table():
    with db() as conn, conn.cursor() as cursor:
        cursor.execute(
            f"""
            CREATE TABLE IF NOT EXISTS {DETAIL_ANALYSIS_TABLE} (
              id BIGINT PRIMARY KEY AUTO_INCREMENT,
              user_id BIGINT NOT NULL,
              username VARCHAR(128) NOT NULL DEFAULT '',
              source VARCHAR(32) NOT NULL,
              product_table VARCHAR(128) NOT NULL,
              product_id VARCHAR(128) NOT NULL,
              date_record DATE NOT NULL,
              analysis_type VARCHAR(32) NOT NULL,
              product_title TEXT NULL,
              status VARCHAR(32) NOT NULL DEFAULT 'RUNNING',
              result_json LONGTEXT NULL,
              error_message TEXT NULL,
              created_at DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP,
              updated_at DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP,
              completed_at DATETIME NULL,
              UNIQUE KEY uq_user_detail_analysis (user_id, source, product_id, date_record, analysis_type),
              KEY idx_user_type_status (user_id, analysis_type, status),
              KEY idx_product_lookup (source, product_id, date_record)
            ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4
            """
        )
        conn.commit()


def ensure_expert_team_tables():
    with db() as conn, conn.cursor() as cursor:
        cursor.execute(
            f"""
            CREATE TABLE IF NOT EXISTS {EXPERT_TEAM_SESSION_TABLE} (
              id BIGINT PRIMARY KEY AUTO_INCREMENT,
              user_id BIGINT NOT NULL,
              username VARCHAR(128) NOT NULL DEFAULT '',
              title VARCHAR(255) NOT NULL DEFAULT '',
              project_code VARCHAR(64) NOT NULL DEFAULT 'general',
              project_context LONGTEXT NULL,
              created_at DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP,
              updated_at DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP,
              KEY idx_user_updated (user_id, updated_at),
              KEY idx_project (project_code)
            ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4
            """
        )
        cursor.execute(
            f"""
            CREATE TABLE IF NOT EXISTS {EXPERT_TEAM_MESSAGE_TABLE} (
              id BIGINT PRIMARY KEY AUTO_INCREMENT,
              session_id BIGINT NOT NULL,
              user_id BIGINT NOT NULL,
              role VARCHAR(20) NOT NULL,
              team_role VARCHAR(64) NOT NULL DEFAULT '',
              content LONGTEXT NOT NULL,
              meta_json LONGTEXT NULL,
              created_at DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP,
              KEY idx_session_id (session_id, id),
              KEY idx_user_created (user_id, created_at)
            ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4
            """
        )
        cursor.execute(
            f"""
            CREATE TABLE IF NOT EXISTS {EXPERT_WORKFLOW_INSTANCE_TABLE} (
              id BIGINT PRIMARY KEY AUTO_INCREMENT,
              session_id BIGINT NOT NULL,
              user_id BIGINT NOT NULL,
              workflow_code VARCHAR(128) NOT NULL,
              workflow_name VARCHAR(255) NOT NULL DEFAULT '',
              intent VARCHAR(128) NOT NULL DEFAULT '',
              mode VARCHAR(64) NOT NULL DEFAULT '',
              status VARCHAR(32) NOT NULL DEFAULT 'RUNNING',
              input_json LONGTEXT NULL,
              output_json LONGTEXT NULL,
              error_message TEXT NULL,
              created_at DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP,
              updated_at DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP,
              completed_at DATETIME NULL,
              KEY idx_session_workflow (session_id, id),
              KEY idx_user_status (user_id, status, updated_at),
              KEY idx_workflow_code (workflow_code)
            ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4
            """
        )
        cursor.execute(
            f"""
            CREATE TABLE IF NOT EXISTS {EXPERT_WORKFLOW_STEP_TABLE} (
              id BIGINT PRIMARY KEY AUTO_INCREMENT,
              workflow_id BIGINT NOT NULL,
              step_order INT NOT NULL,
              step_code VARCHAR(128) NOT NULL,
              step_name VARCHAR(255) NOT NULL DEFAULT '',
              executor VARCHAR(128) NOT NULL DEFAULT '',
              status VARCHAR(32) NOT NULL DEFAULT 'PENDING',
              input_json LONGTEXT NULL,
              output_json LONGTEXT NULL,
              error_message TEXT NULL,
              started_at DATETIME NULL,
              completed_at DATETIME NULL,
              created_at DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP,
              updated_at DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP,
              KEY idx_workflow_step (workflow_id, step_order),
              KEY idx_status (status)
            ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4
            """
        )
        conn.commit()


EXPERT_TEAM_ROLES = [
    {
        "code": "chief_planner",
        "name": "首席项目规划专家",
        "level": "leader",
        "focus": "理解用户目标，拆解版本路线，决定先做什么、后做什么。",
        "limits": "不直接替代具体数据/IP/材质专家做专业细节判断。",
    },
    {
        "code": "expert_manager",
        "name": "专家团队管理者",
        "level": "leader",
        "focus": "判断当前问题该调用哪些专家，识别专家能力边界和冲突。",
        "limits": "不单独给商品最终上架结论。",
    },
    {
        "code": "data_method_lead",
        "name": "数据方法论负责人",
        "level": "leader",
        "focus": "定义指标口径、分析框架、异常判断和数据证据标准。",
        "limits": "不判断 IP 侵权和图案来源。",
    },
    {
        "code": "ux_lead",
        "name": "用户体验负责人",
        "level": "leader",
        "focus": "降低用户阅读成本，设计更清楚的前端展示和交互路径。",
        "limits": "不决定数据库字段和业务风控规则。",
    },
    {
        "code": "data_analyst",
        "name": "数据分析专家",
        "level": "executor",
        "focus": "销量、环比、商品卡占比、视频/直播结构、达人带货、异常波动。",
        "limits": "不能单独给最终选品结论，不负责 IP/材质判断。",
    },
    {
        "code": "ip_compliance",
        "name": "IP合规专家",
        "level": "executor",
        "focus": "图片来源、作品/角色/品牌命中、侵权风险、IP等级和复核建议。",
        "limits": "不负责销量潜力判断。",
    },
    {
        "code": "material_visual",
        "name": "材质/视觉专家",
        "level": "executor",
        "focus": "材质、图案复杂度、Logo/插画、工厂可生产性、素材提取价值。",
        "limits": "不负责销售趋势和达人分析。",
    },
    {
        "code": "operations",
        "name": "运营专家",
        "level": "executor",
        "focus": "上架优先级、TikTok运营路径、达人/商品卡策略、监控预警动作。",
        "limits": "不负责代码实现和底层表结构。",
    },
    {
        "code": "product_manager",
        "name": "商品产品经理专家",
        "level": "executor",
        "focus": "商品定位、用户痛点、卖点清晰度、差异化、目标人群。",
        "limits": "这里是商品侧产品经理，不是系统代码维护角色。",
    },
    {
        "code": "selection_reviewer",
        "name": "智能选品总评专家",
        "level": "executor",
        "focus": "综合数据、IP、材质、运营、产品判断，输出是否值得选和复核点。",
        "limits": "必须引用其他专家证据，不能凭空下结论。",
    },
    {
        "code": "system_architect",
        "name": "技术架构专家",
        "level": "executor",
        "focus": "系统结构、接口、数据库、任务队列、可扩展性和技术债。",
        "limits": "不直接判断商品商业价值。",
    },
    {
        "code": "project_file_scout",
        "name": "项目文件侦察员",
        "level": "readonly_tool",
        "focus": "只读查看项目文件、定位前端页面、后端接口、样式、工具脚本和文档中的相关片段。",
        "limits": "只能读取项目白名单文件，不能写文件、不能执行命令、不能读取项目外路径。",
    },
    {
        "code": "codex_executor",
        "name": "技术执行Agent / Codex执行桥",
        "level": "write_tool",
        "focus": "在用户确认执行后，对项目白名单文件执行受控代码修改、验证并回传结果。",
        "limits": "只能执行内置白名单配方，不能运行任意命令，不能修改 D:\\choice_product 外文件，不能绕过用户确认。",
    },
]


EXPERT_EXECUTION_HANDOFF_RULES = """
专家团队执行交接规则：
1. 专家团队本身不直接写代码、不写库、不改线上数据，但必须判断“谁具备执行能力”。
2. 当用户表达“可以执行、开始执行、落地、修改、让A做、交给技术人员、你手下人员去做”等意图时，必须输出《执行交接单》。
3. 只要本轮只是专家团队输出方案或交接单、没有真实写入文件/数据库，就必须在回答最开头明确写：
   - 执行状态：未修改，仅生成执行方案/交接单。
   - 用户如何看到修改后内容：用户只需确认执行；由技术执行人员A/Codex按交接单完成代码修改、重启后端并刷新页面后，用户审核结果。
   - 如果已经由真实执行器完成修改，才允许写“已修改”，并必须列出修改文件、修改内容、验证方式。
4. 《执行交接单》必须让执行者A可以直接听懂，不能只说方向，必须包含：
   - 执行者A是谁：技术执行人员A/Codex/业务技术团队/业务运营团队/数据分析执行者等。
   - 为什么A有能力：A需要具备哪些技能、能访问哪些资源、能执行哪些动作。
   - 不应该交给谁：哪些专家只负责判断，不负责执行。
   - 执行目标：这次要完成什么，完成后用户能看到什么。
   - 执行范围：涉及页面、接口、数据库表、字段、权限、只读工具或业务流程。
   - 操作步骤：按 1、2、3 写清楚，尽量具体到文件、接口、字段、按钮、校验点。
   - 输入资料：A需要从用户、数据库、接口或截图拿到什么。
   - 验收标准：用户如何判断做完了，至少列出可测试的结果。
   - 风险边界：哪些不能做，哪些需要用户确认后再做。
5. 如果当前信息不足，仍然要先给出“可执行的第一步交接单”，并说明A需要补读哪些数据。
6. 如果任务是商品分析，执行者通常是业务专家；如果任务是系统改造，执行者通常是技术执行人员A/Codex/业务技术团队。
7. 禁止对用户说“你需要找技术执行人员”“如果你自己是技术执行人员就自己改”“你打开文件修改”。正确说法是：专家团队已判断执行对象为A/Codex，用户只需要确认是否执行，执行完成后审核结果。
""".strip()


def expert_team_roles_payload():
    return EXPERT_TEAM_ROLES


def normalize_expert_session(row):
    if not row:
        return None
    return {
        "id": row.get("id"),
        "title": row.get("title") or "专家团队会话",
        "project_code": row.get("project_code") or "general",
        "created_at": stringify(row.get("created_at")),
        "updated_at": stringify(row.get("updated_at")),
    }


def normalize_expert_message(row):
    role = row.get("role")
    content = row.get("content") or ""
    if role == "assistant":
        content = sanitize_expert_message_for_display(content)
    return {
        "id": row.get("id"),
        "role": role,
        "team_role": row.get("team_role") or "",
        "content": content,
        "created_at": stringify(row.get("created_at")),
    }


def normalize_expert_workflow_instance(row):
    if not row:
        return None
    return {
        "id": row.get("id"),
        "session_id": row.get("session_id"),
        "workflow_code": row.get("workflow_code") or "",
        "workflow_name": row.get("workflow_name") or "",
        "intent": row.get("intent") or "",
        "mode": row.get("mode") or "",
        "status": row.get("status") or "",
        "input": parse_json(row.get("input_json")),
        "output": parse_json(row.get("output_json")),
        "error_message": row.get("error_message") or "",
        "created_at": stringify(row.get("created_at")),
        "updated_at": stringify(row.get("updated_at")),
        "completed_at": stringify(row.get("completed_at")),
    }


def normalize_expert_workflow_step(row):
    if not row:
        return None
    return {
        "id": row.get("id"),
        "workflow_id": row.get("workflow_id"),
        "step_order": row.get("step_order"),
        "step_code": row.get("step_code") or "",
        "step_name": row.get("step_name") or "",
        "executor": row.get("executor") or "",
        "status": row.get("status") or "",
        "input": parse_json(row.get("input_json")),
        "output": parse_json(row.get("output_json")),
        "error_message": row.get("error_message") or "",
        "started_at": stringify(row.get("started_at")),
        "completed_at": stringify(row.get("completed_at")),
    }


def make_expert_session_title(message):
    text = re.sub(r"\s+", " ", stringify(message)).strip()
    return text[:40] or "专家团队会话"


def build_expert_context_query(message, history):
    current = stringify(message).strip()
    is_confirmation = current.lower() in {"a", "ok", "yes", "y"} or current in {"确认", "确认执行", "可以", "可以执行", "执行", "开始执行"}
    parts = []
    for item in (history or [])[-8:]:
        if item.get("role") == "user" or (is_confirmation and item.get("role") == "assistant"):
            content = stringify(item.get("content")).strip()
            if content == current:
                continue
            if content:
                parts.append(content)
    if current:
        if is_confirmation:
            parts.append(f"用户确认执行：{current}")
        else:
            parts.append(current)
    return "\n".join(list(dict.fromkeys(parts))) or current


def build_expert_ceo_decision(message, readonly_context="", image_count=0):
    text = stringify(message)
    lower = text.lower()
    readonly = stringify(readonly_context)
    wants_execution = any(word in text for word in ["执行", "开始做", "开始改", "修改", "落地", "部署", "交给", "让A", "让 Codex", "直接做", "新增", "添加", "加一个", "我要的是", "我想要", "不要包括", "只需要", "改成", "去掉", "删除", "隐藏"])
    code_related = any(word in lower for word in ["代码", "接口", "api", "页面", "前端", "后端", "字段", "数据库", "路由", "按钮", "登录", "注册", "权限", "导航栏", "排序", "筛选", "分页", "样式", "排版", "饼图", "条形图", "缓存", "刷新", "部署", "版本号", "静态资源", "工作流", "执行层", "执行能力", "productvideo", "workflow", "workflowengine", "html", "css", "js", "flask", "java", "spring", "agent", "handler", "app.py", "index.html", "styles.css", "renderproductdetail", "product_detail", "think", "<think>", "translation"])
    if any(word in text for word in ["站内详情", "站内详细", "商品详情", "商品详细", "详情弹窗", "属性信息", "卖点展示", "卖点", "翻译成中文", "展示出来", "中文结论"]):
        code_related = True
    if any(word in text for word in ["执行交接单", "技术执行人员", "Codex", "app.py", "index.html", "styles.css"]):
        code_related = True
    if "用户确认执行" in text:
        wants_execution = True
    product_related = any(word in text for word in ["商品", "选品", "销量", "销售", "IP", "材质", "标题", "卖点", "达人", "上架"])
    planning_related = any(word in text for word in ["规划", "方案", "团队", "专家", "架构", "流程", "中台", "模式"])
    has_file_context = "文件：" in readonly or "项目文件侦察结果" in readonly or "会话基底文件读取结果" in readonly
    has_data_context = "数据源：" in readonly or "只读工具结果" in readonly

    if wants_execution and code_related and is_explicit_code_write_request(text):
        intent = "system_change_handoff"
        mode = "PLAN_AND_HANDOFF"
        handlers = ["project_file_scout", "system_architect", "ux_lead", "codex_executor", "execution_handoff"]
        handoff_target = "技术执行Agent / Codex执行桥"
    elif code_related:
        intent = "project_diagnosis"
        mode = "READONLY_ANALYSIS"
        handlers = ["project_file_scout", "system_architect", "ux_lead"]
        handoff_target = "技术执行人员A / Codex"
    elif product_related:
        intent = "product_analysis"
        mode = "EXPERT_REVIEW"
        handlers = ["data_analysis", "ip_compliance", "material_visual", "operations", "product_manager", "selection_reviewer"]
        handoff_target = "业务专家团队"
    elif planning_related:
        intent = "team_planning"
        mode = "PLAN"
        handlers = ["chief_planner", "team_manager", "data_method_lead", "ux_lead", "system_architect"]
        handoff_target = "专家领导层"
    else:
        intent = "general_consulting"
        mode = "ANSWER"
        handlers = ["chief_planner", "team_manager"]
        handoff_target = "专家领导层"

    missing_params = []
    if image_count and "视觉模型" in readonly:
        missing_params.append("当前专家团队还没有接入视觉模型，只能接收图片附件元信息，不能直接识别图片内容。")

    return {
        "intent": intent,
        "mode": mode,
        "confidence": 0.86 if (has_file_context or has_data_context or wants_execution or product_related or code_related) else 0.68,
        "need_user_input": False,
        "user_role": "用户只负责提出要求、查看结果、审核结果；不负责粘贴代码、判断能力或整理项目材料。",
        "handlers": handlers,
        "handoff_target": handoff_target,
        "missing_params": missing_params,
        "evidence": {
            "has_file_context": has_file_context,
            "has_data_context": has_data_context,
            "image_count": image_count,
        },
        "dispatcher_rule": "先由CEO完成意图和模式判断，再分发专家；专家不得把可由只读工具完成的检索任务反问给用户。",
    }


def is_explicit_code_write_request(message):
    text = stringify(message)
    lower = text.lower()
    write_terms = [
        "修改代码", "改代码", "写入", "落地修改", "直接修改", "帮我改", "把代码", "补上",
        "新增接口", "新增路由", "新增按钮", "删除按钮", "改成", "替换成", "接入", "实现",
        "修复这个bug", "修复接口", "修复排序", "优化代码",
    ]
    if any(term in text for term in write_terms):
        return True
    if any(term in lower for term in ["patch", "codepatch", "apply patch", "commit"]):
        return True
    diagnostic_terms = ["为什么", "原因", "分析", "查看", "检查", "explain", "索引", "慢", "方案", "可行"]
    if any(term in text.lower() for term in diagnostic_terms):
        return False
    return False


def format_expert_ceo_decision(decision):
    return json.dumps(decision or {}, ensure_ascii=False, indent=2)


def dispatch_expert_handlers(ceo_decision, message, readonly_context="", base_files=None, session_id=None, user_id=None):
    decision = dict(ceo_decision or {})
    decision["_session_id"] = session_id
    decision["_user_id"] = user_id
    decision["_message"] = stringify(message)
    handlers = decision.get("handlers") or ["chief_planner"]
    registry = {
        "project_file_scout": handle_project_file_scout,
        "system_architect": handle_system_architect,
        "ux_lead": handle_ux_lead,
        "codex_executor": handle_codex_executor,
        "execution_handoff": handle_execution_handoff,
        "data_analysis": handle_business_expert,
        "ip_compliance": handle_business_expert,
        "material_visual": handle_business_expert,
        "operations": handle_business_expert,
        "product_manager": handle_business_expert,
        "selection_reviewer": handle_business_expert,
        "chief_planner": handle_planning_expert,
        "team_manager": handle_planning_expert,
        "data_method_lead": handle_planning_expert,
    }
    results = []
    workflow_result = run_expert_workflow_engine(
        "workflow_engine",
        decision,
        message,
        readonly_context,
        base_files or [],
    )
    decision["_current_workflow_id"] = workflow_result.get("workflow_id") if workflow_result.get("matched") else None
    decision["_current_workflow_result"] = workflow_result if workflow_result.get("matched") else None
    results.append({
        "handler": "workflow_engine",
        "status": workflow_result.get("verification", {}).get("workflow_status") or ("COMPLETED" if workflow_result.get("ok") else "NO_EXECUTABLE_WORKFLOW"),
        "summary": workflow_result.get("summary") or "WorkflowEngine 已处理本轮任务。",
        "data": workflow_result,
    })
    for handler_code in handlers:
        handler = registry.get(handler_code, handle_planning_expert)
        try:
            result = handler(handler_code, decision, message, readonly_context, base_files or [])
        except Exception as exc:
            result = {
                "handler": handler_code,
                "status": "FAILED",
                "summary": f"Handler 执行失败：{exc}",
                "data": {},
            }
        result = enhance_expert_handler_result(result, handler_code, decision, message, readonly_context, base_files or [])
        results.append(result)
    return {
        "stage": "Action Execution",
        "session_id": session_id,
        "intent": decision.get("intent"),
        "mode": decision.get("mode"),
        "dispatcher": "IntentHandlerFactory(local)",
        "handler_count": len(results),
        "results": results,
        "response_delivery": build_agent_response_summary(decision, results),
    }


def handle_project_file_scout(handler_code, decision, message, readonly_context, base_files):
    files = []
    for match in re.finditer(r"文件：([^\n]+)", stringify(readonly_context)):
        file_name = match.group(1).strip()
        if file_name not in files:
            files.append(file_name)
    return {
        "handler": handler_code,
        "status": "COMPLETED",
        "summary": f"已读取并检索项目白名单文件 {len(files)} 个。",
        "data": {
            "base_files": base_files,
            "matched_files": files[:12],
            "can_write": False,
        },
    }


def handle_system_architect(handler_code, decision, message, readonly_context, base_files):
    text = stringify(message)
    details = []
    if "product_detail" in readonly_context or "/api/products" in readonly_context:
        details.append("已定位商品详情后端接口：/api/products/<source>/<product_id> / product_detail。")
    if "renderProductDetail" in readonly_context or "openDetail" in readonly_context:
        details.append("已定位商品详情前端渲染入口：renderProductDetail/openDetail 相关代码。")
    if any(word in text for word in ["attributes", "selling_points", "属性", "卖点"]):
        details.append("任务涉及字段透传与详情弹窗展示，优先检查后端 detail 返回对象和前端渲染区块。")
    return {
        "handler": handler_code,
        "status": "COMPLETED",
        "summary": "技术架构专家已完成接口/字段/页面链路判断。",
        "data": {
            "findings": details or ["已完成系统结构初步判断，当前任务可进入方案或执行交接。"],
            "executor": decision.get("handoff_target"),
        },
    }


def handle_ux_lead(handler_code, decision, message, readonly_context, base_files):
    return {
        "handler": handler_code,
        "status": "COMPLETED",
        "summary": "用户体验负责人已完成展示策略判断。",
        "data": {
            "principles": [
                "用户只看结果和审核结果，不承担找文件、找接口、找执行者的任务。",
                "详情页新增字段应有值才显示，避免空区块增加阅读成本。",
                "新增区块应放在商品基础信息之后、深度分析之前，保持扫描顺序。",
            ]
        },
    }


EXPERT_CLOSURE_PROFILES = {
    "project_file_scout": {
        "capability": "只读检索项目文件与上下文",
        "evidence_required": "必须返回已读取文件、命中片段或未命中说明",
        "cannot_claim": "不能声称已修改代码或已执行数据库写入",
        "next_action": "把定位结果交给技术架构专家或执行适配器",
    },
    "system_architect": {
        "capability": "判断接口、字段、数据库、任务队列、前后端链路",
        "evidence_required": "必须引用文件、接口、字段、配置或查询结果",
        "cannot_claim": "未读代码/未查配置时不能断言系统缺失某能力",
        "next_action": "输出可执行落点，交给 Codex 执行适配器或标记缺口",
    },
    "ux_lead": {
        "capability": "检查用户阅读成本、页面提示、最终回复是否可读",
        "evidence_required": "必须说明用户能看到什么、点击哪里、如何验收",
        "cannot_claim": "不能替代技术执行结果说已完成",
        "next_action": "把用户可见要求交给前端/回复校验器",
    },
    "data_analysis": {
        "capability": "分析销量、环比、队列效率、覆盖率、异常波动",
        "evidence_required": "必须优先使用数据库统计、配置值或明确口径",
        "cannot_claim": "没查库时不能给出总量、覆盖率、瓶颈结论",
        "next_action": "输出需要查询的指标或基于已查数据给出结论",
    },
    "ip_compliance": {
        "capability": "判断 IP 来源、侵权风险、等级和复核建议",
        "evidence_required": "必须基于标题、属性、卖点、图片线索、规则库或经验库",
        "cannot_claim": "无法识别来源时不能直接打高风险 A/S，也不能编造来源",
        "next_action": "给出等级、来源、侵权判断和是否需要人工复核",
    },
    "material_visual": {
        "capability": "判断材质、工厂可生产性、视觉复杂度、图案/Logo/插画",
        "evidence_required": "必须基于标题、属性、卖点、图片线索或材质规则",
        "cannot_claim": "不能替代 IP 合规判断侵权等级",
        "next_action": "给出一两句材质总结或详细分析请求",
    },
    "operations": {
        "capability": "制定上架、达人、商品卡、监控和预警动作",
        "evidence_required": "必须引用商品状态、销售指标或用户运营目标",
        "cannot_claim": "不能在没有数据依据时断言必爆或必亏",
        "next_action": "输出可执行运营动作和优先级",
    },
    "product_manager": {
        "capability": "判断商品定位、痛点、卖点、差异化和目标人群",
        "evidence_required": "必须引用标题、属性、卖点、价格或竞品线索",
        "cannot_claim": "不能替代系统产品经理维护代码路线",
        "next_action": "输出商品侧改进建议或卖点重组",
    },
    "selection_reviewer": {
        "capability": "综合各专家证据，给出选品总评",
        "evidence_required": "必须引用数据/IP/材质/运营/产品中至少两个维度",
        "cannot_claim": "不能脱离前序专家证据直接给最终结论",
        "next_action": "输出是否值得选、风险点和复核点",
    },
    "codex_executor": {
        "capability": "按白名单适配器写文件、验证、回传结果",
        "evidence_required": "必须返回修改文件、验证结果或未命中适配器原因",
        "cannot_claim": "没有 changed_files/verification 时不能说已完成代码修改",
        "next_action": "执行适配器、补适配器或返回待补能力项",
    },
    "execution_handoff": {
        "capability": "把执行结果包装成用户可审核结果",
        "evidence_required": "必须基于真实 handler 状态",
        "cannot_claim": "不能暴露内部调试状态给普通用户",
        "next_action": "交给回复体验校验层",
    },
}


def enhance_expert_handler_result(result, handler_code, decision, message, readonly_context, base_files):
    enhanced = dict(result or {})
    data = dict(enhanced.get("data") or {})
    profile = EXPERT_CLOSURE_PROFILES.get(handler_code, {
        "capability": "规划、分发或辅助判断",
        "evidence_required": "必须说明依据和边界",
        "cannot_claim": "不能把未验证结论当事实",
        "next_action": "交给具备能力的专家或执行器",
    })
    evidence = build_handler_evidence(handler_code, data, readonly_context, base_files)
    confidence = estimate_handler_confidence(handler_code, enhanced, evidence, readonly_context)
    data.setdefault("closure", {
        "capability": profile["capability"],
        "evidence_required": profile["evidence_required"],
        "cannot_claim": profile["cannot_claim"],
        "evidence": evidence,
        "confidence": confidence,
        "next_action": profile["next_action"],
        "verified": bool(evidence) or handler_code in {"ux_lead", "execution_handoff"},
    })
    enhanced["data"] = data
    enhanced["quality"] = {
        "closed_loop": data["closure"]["verified"],
        "confidence": confidence,
        "has_evidence": bool(evidence),
    }
    return enhanced


def build_handler_evidence(handler_code, data, readonly_context, base_files):
    evidence = []
    if base_files:
        evidence.append({"type": "base_files", "items": [safe_relpath(Path(item)) if isinstance(item, str) else stringify(item) for item in base_files[:6]]})
    matched_files = data.get("matched_files") or []
    if matched_files:
        evidence.append({"type": "matched_files", "items": matched_files[:8]})
    findings = data.get("findings") or data.get("principles") or []
    if findings:
        evidence.append({"type": "handler_findings", "items": findings[:6]})
    if readonly_context:
        context_files = []
        for match in re.finditer(r"文件：([^\n]+)", stringify(readonly_context)):
            file_name = match.group(1).strip()
            if file_name not in context_files:
                context_files.append(file_name)
        if context_files:
            evidence.append({"type": "readonly_context_files", "items": context_files[:8]})
    if handler_code == "codex_executor":
        changed_files = data.get("changed_files") or []
        verification = data.get("verification") or {}
        if changed_files:
            evidence.append({"type": "changed_files", "items": changed_files})
        if verification:
            evidence.append({"type": "verification", "items": verification})
    return evidence


def estimate_handler_confidence(handler_code, result, evidence, readonly_context):
    status = result.get("status")
    if status == "FAILED":
        return 0.0
    if handler_code == "codex_executor":
        data = result.get("data") or {}
        verification = data.get("verification") or {}
        if result.get("status") == "COMPLETED" and verification:
            return 0.9
        if result.get("status") in {"WAITING_ADAPTER", "WAITING_USER_CONFIRMATION", "NO_EXECUTABLE_WORKFLOW", "UNSUPPORTED_AUTOMATION"}:
            return 0.45
    if evidence:
        return 0.75
    if readonly_context:
        return 0.55
    return 0.35


def handle_codex_executor(handler_code, decision, message, readonly_context, base_files):
    current_workflow_id = (decision or {}).get("_current_workflow_id")
    existing_workflow = fetch_expert_workflow_instance_by_id(current_workflow_id, (decision or {}).get("_session_id"), (decision or {}).get("_user_id"))
    if existing_workflow:
        status = existing_workflow.get("status")
        output = parse_json(existing_workflow.get("output_json"))
        return {
            "handler": handler_code,
            "status": "COMPLETED" if status == "SUCCESS" else status,
            "summary": f"codex_executor 已接入当前工作流 `{existing_workflow.get('workflow_code')}`，当前状态：{status}。",
            "data": {
                "workflow_id": existing_workflow.get("id"),
                "workflow_code": existing_workflow.get("workflow_code"),
                "workflow_name": existing_workflow.get("workflow_name"),
                "workflow_status": status,
                "output": output,
                "changed_files": ((output or {}).get("changed_files") or []),
                "pending_confirmation": ((output or {}).get("pending_confirmation") or {}),
                "verification": ((output or {}).get("verification") or {"workflow_status": status}),
            },
        }
    workflow_result = run_expert_workflow_engine(
        handler_code,
        decision,
        message,
        readonly_context,
        base_files,
    )
    if workflow_result.get("matched"):
        workflow_status = (workflow_result.get("verification") or {}).get("workflow_status")
        handler_status = "COMPLETED" if workflow_result.get("ok") else (workflow_status or "FAILED")
        return {
            "handler": handler_code,
            "status": handler_status,
            "summary": workflow_result.get("summary") or "业务执行层已运行匹配工作流。",
            "data": workflow_result,
        }
    return {
        "handler": handler_code,
        "status": "NO_EXECUTABLE_WORKFLOW",
        "summary": "WorkflowEngine 已接入，但当前任务没有匹配到工作流定义，未启动工作流实例。",
        "data": {
            "can_write": True,
            "write_scope": "D:\\choice_product 内白名单文本文件",
            "supported_workflows": list_expert_execution_workflows(),
            "safety": "WorkflowEngine 只执行已注册步骤和能力适配器；不执行任意命令、不删除文件、不修改数据库。",
        },
    }


def list_expert_execution_workflows():
    return [
        {
            "workflow_code": item["workflow_code"],
            "name": item["name"],
            "intent_codes": item["intent_codes"],
            "owner": item["owner"],
            "step_count": len(item["steps"]),
        }
        for item in get_expert_workflow_definitions()
    ]


def get_expert_workflow_definitions():
    return [
        {
            "workflow_code": "system_change_execution",
            "name": "系统改造执行工作流",
            "intent_codes": ["system_change_handoff"],
            "match_keywords": [],
            "owner": "technical_execution_team",
            "steps": [
                {"code": "understand_task", "name": "理解业务目标", "executor": "system_architect", "kind": "sync"},
                {"code": "inspect_context", "name": "读取项目上下文", "executor": "project_file_scout", "kind": "sync"},
                {"code": "route_capability", "name": "判断执行能力与边界", "executor": "codex_executor", "kind": "sync"},
                {"code": "execute_adapter", "name": "调用安全执行适配器", "executor": "codex_executor", "kind": "adapter"},
                {"code": "verify_result", "name": "验证与回传", "executor": "codex_executor", "kind": "verify"},
            ],
        },
        {
            "workflow_code": "web_ui_feature_development",
            "name": "网站页面与交互开发工作流",
            "intent_codes": ["system_change_handoff"],
            "match_keywords": ["页面", "前端", "按钮", "弹窗", "导航", "展示", "布局", "交互", "粘贴图片", "上传图片", "详情", "列表", "卡片"],
            "owner": "frontend_feature_team",
            "steps": [
                {"code": "ui_requirement_parse", "name": "解析页面目标与用户路径", "executor": "ux_lead", "kind": "sync"},
                {"code": "locate_frontend_entry", "name": "定位前端入口与状态流", "executor": "project_file_scout", "kind": "sync"},
                {"code": "design_component_state", "name": "设计组件结构和交互状态", "executor": "ux_lead", "kind": "sync"},
                {"code": "frontend_adapter", "name": "调用前端安全执行适配器", "executor": "codex_executor", "kind": "adapter"},
                {"code": "ui_smoke_verify", "name": "页面冒烟验证", "executor": "codex_executor", "kind": "verify"},
            ],
        },
        {
            "workflow_code": "web_api_integration_development",
            "name": "网站接口联调开发工作流",
            "intent_codes": ["system_change_handoff", "project_diagnosis"],
            "match_keywords": ["接口", "api", "404", "500", "json", "fetch", "请求", "响应", "路由", "flask", "后端", "network"],
            "owner": "backend_integration_team",
            "steps": [
                {"code": "api_requirement_parse", "name": "解析接口目标", "executor": "system_architect", "kind": "sync"},
                {"code": "locate_route_handler", "name": "定位路由和调用方", "executor": "project_file_scout", "kind": "sync"},
                {"code": "contract_check", "name": "校验请求/响应契约", "executor": "system_architect", "kind": "sync"},
                {"code": "backend_adapter", "name": "调用后端安全执行适配器", "executor": "codex_executor", "kind": "adapter"},
                {"code": "api_smoke_verify", "name": "接口冒烟验证", "executor": "codex_executor", "kind": "verify"},
            ],
        },
        {
            "workflow_code": "web_data_field_development",
            "name": "网站数据字段与列表开发工作流",
            "intent_codes": ["system_change_handoff", "project_diagnosis"],
            "match_keywords": ["字段", "数据库", "表", "排序", "筛选", "分页", "销量", "环比", "采集日期", "总销量", "sku", "date_record"],
            "owner": "data_ui_integration_team",
            "steps": [
                {"code": "field_requirement_parse", "name": "解析字段口径", "executor": "data_method_lead", "kind": "sync"},
                {"code": "inspect_data_source", "name": "定位数据源和字段映射", "executor": "project_file_scout", "kind": "sync"},
                {"code": "query_contract_check", "name": "校验查询、排序和分页契约", "executor": "system_architect", "kind": "sync"},
                {"code": "data_ui_adapter", "name": "调用数据展示安全执行适配器", "executor": "codex_executor", "kind": "adapter"},
                {"code": "data_ui_verify", "name": "列表/详情数据验证", "executor": "codex_executor", "kind": "verify"},
            ],
        },
        {
            "workflow_code": "web_responsive_style_development",
            "name": "网站样式与响应式开发工作流",
            "intent_codes": ["system_change_handoff"],
            "match_keywords": ["样式", "css", "排版", "居中", "响应式", "移动端", "宽度", "高度", "颜色", "图表", "饼图", "条形图", "太小", "挤压"],
            "owner": "frontend_visual_team",
            "steps": [
                {"code": "visual_problem_parse", "name": "解析视觉问题", "executor": "ux_lead", "kind": "sync"},
                {"code": "locate_style_scope", "name": "定位样式作用域", "executor": "project_file_scout", "kind": "sync"},
                {"code": "responsive_rule_plan", "name": "制定响应式规则", "executor": "ux_lead", "kind": "sync"},
                {"code": "style_adapter", "name": "调用样式安全执行适配器", "executor": "codex_executor", "kind": "adapter"},
                {"code": "visual_verify", "name": "视觉验收", "executor": "codex_executor", "kind": "verify"},
            ],
        },
        {
            "workflow_code": "web_auth_permission_development",
            "name": "网站登录权限开发工作流",
            "intent_codes": ["system_change_handoff", "project_diagnosis"],
            "match_keywords": ["登录", "注册", "权限", "role", "admin", "manager", "退出", "session", "导航栏"],
            "owner": "auth_security_team",
            "steps": [
                {"code": "auth_requirement_parse", "name": "解析认证与权限目标", "executor": "system_architect", "kind": "sync"},
                {"code": "locate_auth_flow", "name": "定位登录、注册、权限链路", "executor": "project_file_scout", "kind": "sync"},
                {"code": "permission_boundary_check", "name": "校验权限边界", "executor": "system_architect", "kind": "sync"},
                {"code": "auth_adapter", "name": "调用认证安全执行适配器", "executor": "codex_executor", "kind": "adapter"},
                {"code": "auth_verify", "name": "登录权限验证", "executor": "codex_executor", "kind": "verify"},
            ],
        },
        {
            "workflow_code": "web_release_cache_development",
            "name": "网站发布与缓存治理工作流",
            "intent_codes": ["system_change_handoff", "project_diagnosis"],
            "match_keywords": ["缓存", "清缓存", "刷新", "新版", "版本号", "cache-control", "部署", "cloudflare", "静态资源", "css版本", "js版本", "上线"],
            "owner": "release_engineering_team",
            "steps": [
                {"code": "release_problem_parse", "name": "解析发布问题", "executor": "system_architect", "kind": "sync"},
                {"code": "inspect_static_assets", "name": "定位静态资源和响应头", "executor": "project_file_scout", "kind": "sync"},
                {"code": "cache_strategy_plan", "name": "制定缓存策略", "executor": "system_architect", "kind": "sync"},
                {"code": "release_adapter", "name": "调用发布安全执行适配器", "executor": "codex_executor", "kind": "adapter"},
                {"code": "release_verify", "name": "发布缓存验证", "executor": "codex_executor", "kind": "verify"},
            ],
        },
        {
            "workflow_code": "project_diagnosis",
            "name": "项目诊断工作流",
            "intent_codes": ["project_diagnosis"],
            "match_keywords": [],
            "owner": "technical_architecture_team",
            "steps": [
                {"code": "read_files", "name": "读取相关文件", "executor": "project_file_scout", "kind": "sync"},
                {"code": "analyze_architecture", "name": "分析接口与字段链路", "executor": "system_architect", "kind": "sync"},
                {"code": "summarize_findings", "name": "输出诊断结论", "executor": "system_architect", "kind": "sync"},
            ],
        },
        {
            "workflow_code": "product_analysis_review",
            "name": "商品分析会审工作流",
            "intent_codes": ["product_analysis"],
            "match_keywords": [],
            "owner": "business_expert_team",
            "steps": [
                {"code": "data_review", "name": "数据分析", "executor": "data_analysis", "kind": "sync"},
                {"code": "ip_review", "name": "IP 合规判断", "executor": "ip_compliance", "kind": "sync"},
                {"code": "material_review", "name": "材质与视觉判断", "executor": "material_visual", "kind": "sync"},
                {"code": "operation_review", "name": "运营判断", "executor": "operations", "kind": "sync"},
                {"code": "final_review", "name": "选品总评", "executor": "selection_reviewer", "kind": "sync"},
            ],
        },
        {
            "workflow_code": "team_planning_review",
            "name": "专家团队规划工作流",
            "intent_codes": ["team_planning", "general_consulting"],
            "match_keywords": [],
            "owner": "expert_leader_team",
            "steps": [
                {"code": "goal_parse", "name": "解析目标", "executor": "chief_planner", "kind": "sync"},
                {"code": "expert_route", "name": "安排专家", "executor": "expert_manager", "kind": "sync"},
                {"code": "method_check", "name": "方法论校验", "executor": "data_method_lead", "kind": "sync"},
                {"code": "response_pack", "name": "组织回复", "executor": "chief_planner", "kind": "sync"},
            ],
        },
    ]


def select_expert_workflow_definition(decision):
    intent = stringify((decision or {}).get("intent"))
    message = stringify((decision or {}).get("_message"))
    candidates = [workflow for workflow in get_expert_workflow_definitions() if intent in workflow["intent_codes"]]
    if not candidates:
        return None
    scored = []
    for workflow in candidates:
        keywords = workflow.get("match_keywords") or []
        score = sum(1 for keyword in keywords if keyword and keyword.lower() in message.lower())
        if keywords:
            score += 10 if score else 0
        scored.append((score, workflow))
    scored.sort(key=lambda item: item[0], reverse=True)
    return scored[0][1]


def run_expert_workflow_engine(handler_code, decision, message, readonly_context, base_files):
    workflow = select_expert_workflow_definition(decision)
    if not workflow:
        return {
            "ok": False,
            "matched": False,
            "handler": handler_code,
            "workflow_code": "",
            "workflow_name": "",
            "steps": [],
        }

    session_id = (decision or {}).get("_session_id")
    user_id = (decision or {}).get("_user_id") or 0
    workflow_input = {
        "message": stringify(message),
        "intent": (decision or {}).get("intent"),
        "mode": (decision or {}).get("mode"),
        "handlers": (decision or {}).get("handlers"),
        "base_files": base_files or [],
        "has_readonly_context": bool(readonly_context),
    }
    workflow_id = create_expert_workflow_instance(session_id, user_id, workflow, decision, workflow_input)
    executed_steps = []
    final_status = "SUCCESS"
    errors = []
    changed_files = []
    pending_confirmation = None
    verification = {
        "workflow_instance_created": bool(workflow_id),
    }

    for index, step in enumerate(workflow["steps"], start=1):
        step_id = create_expert_workflow_step(workflow_id, index, step, workflow_input)
        step_result = execute_expert_workflow_step(step, decision, message, readonly_context, base_files)
        update_expert_workflow_step(step_id, step_result)
        executed_steps.append({
            "step_order": index,
            "step_code": step["code"],
            "step_name": step["name"],
            "executor": step["executor"],
            **step_result,
        })
        if step_result.get("status") in {"FAILED", "WAITING_ADAPTER", "WAITING_USER_CONFIRMATION"}:
            final_status = step_result.get("status") if step_result.get("status") in {"WAITING_ADAPTER", "WAITING_USER_CONFIRMATION"} else "FAILED"
            if step_result.get("pending_confirmation"):
                pending_confirmation = step_result.get("pending_confirmation")
            if step_result.get("error"):
                errors.append(step_result.get("error"))
            if step_result.get("status") in {"WAITING_ADAPTER", "WAITING_USER_CONFIRMATION"}:
                break
        for changed_file in step_result.get("changed_files") or []:
            if changed_file not in changed_files:
                changed_files.append(changed_file)
        if isinstance(step_result.get("verification"), dict):
            verification.update(step_result.get("verification"))

    output = {
        "workflow_id": workflow_id,
        "workflow_code": workflow["workflow_code"],
        "workflow_name": workflow["name"],
        "owner": workflow["owner"],
        "steps": executed_steps,
        "status": final_status,
        "errors": errors,
        "changed_files": changed_files,
        "pending_confirmation": pending_confirmation,
        "verification": verification,
    }
    update_expert_workflow_instance(workflow_id, final_status, output, "\n".join(errors))
    return {
        "ok": final_status == "SUCCESS",
        "matched": True,
        "handler": handler_code,
        "workflow_id": workflow_id,
        "workflow_code": workflow["workflow_code"],
        "workflow_name": workflow["name"],
        "execution_expert": workflow["owner"],
        "can_write": final_status == "SUCCESS",
        "write_scope": "D:\\choice_product 内受控工作流",
        "changed_files": changed_files,
        "verification": {
            **verification,
            "workflow_status": final_status,
            "step_count": len(executed_steps),
        },
        "errors": errors,
        "steps": executed_steps,
        "summary": build_workflow_execution_summary(workflow, final_status, executed_steps),
    }


def execute_expert_workflow_step(step, decision, message, readonly_context, base_files):
    kind = step.get("kind")
    code = step.get("code")
    if kind == "adapter":
        adapter = select_execution_adapter(decision, message, readonly_context, base_files)
        if not adapter:
            adapter_name = step.get("code") or "execute_adapter"
            return {
                "status": "WAITING_ADAPTER",
                "output": f"已完成任务理解和上下文读取，但当前系统尚未接入 `{adapter_name}` 对应的安全执行适配器。",
                "error": "",
            }
        return adapter(decision, message, readonly_context, base_files)
    if kind == "verify":
        return {
            "status": "COMPLETED",
            "output": "已完成当前工作流可验证部分；如前序步骤等待适配器，本步骤会在适配器完成后继续验证。",
            "error": "",
        }
    outputs = {
        "understand_task": "已将用户自然语言需求转换为结构化执行目标。",
        "inspect_context": "已读取专家团队只读上下文和用户指定基底文件。",
        "route_capability": "已判断执行者、能力边界和安全约束。",
        "ui_requirement_parse": "已解析页面目标、用户路径、交互入口和可见验收点。",
        "locate_frontend_entry": "已定位前端入口、状态变量、事件绑定和渲染函数。",
        "design_component_state": "已设计组件结构、空态/加载态/成功态/失败态和用户提示。",
        "api_requirement_parse": "已解析接口目标、请求方式、参数、返回结构和错误状态。",
        "locate_route_handler": "已定位后端路由、前端 fetch 调用方和可能的 404/JSON 断点。",
        "contract_check": "已校验前后端请求/响应契约、鉴权和异常格式。",
        "field_requirement_parse": "已解析字段口径、展示口径、排序口径和筛选口径。",
        "inspect_data_source": "已定位数据源、字段映射、聚合方式和空值策略。",
        "query_contract_check": "已校验查询、排序、分页、日期范围和汇总口径。",
        "visual_problem_parse": "已解析视觉问题、布局压缩点、响应式断点和图表尺寸要求。",
        "locate_style_scope": "已定位样式作用域、组件 class、图表容器和影响范围。",
        "responsive_rule_plan": "已制定桌面/移动端的布局、字号、间距和容器规则。",
        "auth_requirement_parse": "已解析登录、注册、角色权限、导航可见性和会话目标。",
        "locate_auth_flow": "已定位登录注册、session、role 字段和权限判断链路。",
        "permission_boundary_check": "已校验 admin/manager/user 的权限边界和默认角色策略。",
        "release_problem_parse": "已解析缓存、版本号、静态资源和发布刷新问题。",
        "inspect_static_assets": "已定位 HTML/CSS/JS 静态资源引用和响应头配置点。",
        "cache_strategy_plan": "已制定 HTML 不缓存、CSS/JS 版本化、Cloudflare 刷新边界策略。",
        "read_files": "已读取相关项目文件上下文。",
        "analyze_architecture": "已分析接口、字段、页面和数据链路。",
        "summarize_findings": "已形成项目诊断结论。",
        "data_review": "已完成数据维度判断。",
        "ip_review": "已完成 IP 合规维度判断。",
        "material_review": "已完成材质与视觉维度判断。",
        "operation_review": "已完成运营维度判断。",
        "final_review": "已完成选品总评组织。",
        "goal_parse": "已解析用户目标。",
        "expert_route": "已安排专家职责。",
        "method_check": "已完成方法论校验。",
        "response_pack": "已组织结果回传。",
    }
    return {
        "status": "COMPLETED",
        "output": outputs.get(code) or f"{step.get('name')} 已完成。",
        "error": "",
    }


def select_execution_adapter(decision, message, readonly_context, base_files):
    text = "\n".join([
        stringify(message),
        stringify(readonly_context),
        " ".join(stringify(item) for item in (base_files or [])),
    ])
    if (decision or {}).get("intent") == "system_change_handoff" and (
        is_explicit_code_write_request(text) or user_confirmed_extra_write_access(text)
    ):
        return apply_controlled_codepatch_adapter
    if is_detail_selling_points_translate_task(text) or is_translation_think_cleanup_task(text):
        return apply_detail_selling_points_translate_adapter
    return None


def apply_controlled_codepatch_adapter(decision, message, readonly_context, base_files):
    plan_result = generate_codepatch_plan(decision, message, readonly_context, base_files)
    if not plan_result.get("ok"):
        return {
            "status": "WAITING_ADAPTER",
            "output": plan_result.get("message") or "无法生成可执行补丁计划。",
            "error": plan_result.get("error") or "",
            "verification": {"codepatch_plan_generated": False},
        }

    plan = plan_result["plan"]
    safety = validate_codepatch_plan_scope(plan, message)
    if safety.get("status") == "WAITING_USER_CONFIRMATION":
        return {
            "status": "WAITING_USER_CONFIRMATION",
            "output": "补丁计划包含默认白名单外的项目文件，等待用户确认后再写入。",
            "error": "",
            "pending_confirmation": {
                "paths": safety.get("paths") or [],
                "reason": safety.get("reason") or "需要用户确认额外写入范围。",
            },
            "verification": {
                "codepatch_plan_generated": True,
                "scope_checked": True,
                "requires_user_confirmation": True,
            },
        }
    if not safety.get("ok"):
        return {
            "status": "FAILED",
            "output": "补丁计划未通过安全检查，未写入文件。",
            "error": safety.get("reason") or "安全检查失败",
            "verification": {"codepatch_plan_generated": True, "scope_checked": False},
        }

    apply_result = apply_codepatch_plan(plan)
    if not apply_result.get("ok"):
        return {
            "status": "FAILED",
            "output": "补丁应用失败，已保持文件不变。",
            "error": apply_result.get("error") or "补丁应用失败",
            "verification": apply_result.get("verification") or {},
        }

    verification = verify_codepatch_result(apply_result.get("changed_files") or [], plan)
    status = "COMPLETED" if verification.get("passed") else "FAILED"
    if status == "FAILED":
        rollback_result = rollback_codepatch_changes(apply_result.get("originals") or {})
        verification["rolled_back"] = rollback_result.get("ok")
        if rollback_result.get("errors"):
            verification.setdefault("errors", []).extend(rollback_result.get("errors"))
    return {
        "status": status,
        "output": "受控 CodePatch 已完成写入并通过验证。" if status == "COMPLETED" else "补丁验证未通过，已尝试回滚本轮文件改动。",
        "error": "" if status == "COMPLETED" else "; ".join(verification.get("errors") or []),
        "changed_files": apply_result.get("changed_files") or [],
        "diff": apply_result.get("diff") or {},
        "verification": {
            **verification,
            "codepatch_plan_generated": True,
            "scope_checked": True,
            "operations": len(plan.get("operations") or []),
        },
    }


def generate_codepatch_plan(decision, message, readonly_context, base_files):
    api_key = resolve_minimax_api_key()
    if not api_key:
        return {"ok": False, "message": "缺少 MINIMAX_API_KEY，无法生成通用补丁计划。"}

    prompt = build_codepatch_planning_prompt(decision, message, readonly_context, base_files)
    url = f"{MINIMAX_BASE_URL.rstrip('/')}/chat/completions"
    payload = {
        "model": MINIMAX_MODEL,
        "messages": [
            {"role": "system", "content": "你是严格的代码补丁规划器，只能输出 JSON，不要输出 markdown、解释或 <think>。"},
            {"role": "user", "content": prompt},
        ],
        "temperature": 0.05,
        "stream": False,
    }
    try:
        data = call_minimax_api_with_retry(url, api_key, payload, timeout=90)
        raw = data["choices"][0]["message"]["content"]
        plan = parse_codepatch_json(raw)
    except Exception as exc:
        return {"ok": False, "error": str(exc), "message": "调用模型生成补丁计划失败。"}

    if not isinstance(plan, dict):
        return {"ok": False, "message": "模型未返回合法 JSON 补丁计划。"}
    operations = plan.get("operations")
    if not isinstance(operations, list) or not operations:
        return {"ok": False, "message": "补丁计划没有 operations，无法执行。"}
    normalized = normalize_codepatch_plan(plan)
    if not normalized.get("operations"):
        return {"ok": False, "message": "补丁计划没有可执行的合法 operation。"}
    return {"ok": True, "plan": normalized}


def build_codepatch_planning_prompt(decision, message, readonly_context, base_files):
    allowed_roots = "\n".join(f"- {item}" for item in CODEPATCH_WRITE_ROOTS)
    allowed_suffixes = ", ".join(sorted(CODEPATCH_ALLOWED_SUFFIXES))
    return f"""
你要为 Choice Product 项目生成“结构化文件补丁计划”。

硬性规则：
1. 只输出一个 JSON 对象，不要 markdown，不要解释。
2. JSON schema：
{{
  "summary": "本次要做什么",
  "operations": [
    {{
      "op": "replace|insert_after|insert_before|append|create",
      "path": "相对 D:/choice_product 的路径",
      "old": "replace 时必须提供；insert 时可作为 anchor",
      "anchor": "insert_after/insert_before 时必须提供",
      "new": "replace/insert/create/append 的新内容",
      "reason": "为什么改这个文件"
    }}
  ],
  "verification": ["建议验证点"]
}}
3. 不允许生成删除文件、移动文件、运行命令、数据库写入、网络请求。
4. 优先使用用户指定或只读上下文中已经出现的文件和函数，不能凭空创造不存在的入口。
5. replace 的 old 必须是文件中能精确匹配的一段文本；insert 的 anchor 必须能精确匹配。
6. 如果无法确定补丁，请返回 {{"summary":"无法安全生成补丁","operations":[],"verification":[]}}。

默认写入白名单：
{allowed_roots}

允许文件后缀：{allowed_suffixes}

CEO 决策：
{format_expert_ceo_decision(decision)}

用户需求：
{stringify(message)}

用户指定基底文件：
{json.dumps(base_files or [], ensure_ascii=False)}

只读项目上下文：
{stringify(readonly_context)[:24000]}
""".strip()


def parse_codepatch_json(raw):
    cleaned = sanitize_expert_team_answer(raw)
    try:
        return json.loads(cleaned)
    except Exception:
        pass
    match = re.search(r"\{[\s\S]*\}", cleaned)
    if not match:
        return None
    try:
        return json.loads(match.group(0))
    except Exception:
        return None


def normalize_codepatch_plan(plan):
    normalized_ops = []
    for item in plan.get("operations") or []:
        if not isinstance(item, dict):
            continue
        op = stringify(item.get("op")).strip().lower()
        path = stringify(item.get("path")).strip().replace("\\", "/").lstrip("/")
        if op not in {"replace", "insert_after", "insert_before", "append", "create"} or not path:
            continue
        normalized_ops.append({
            "op": op,
            "path": path,
            "old": stringify(item.get("old")),
            "anchor": stringify(item.get("anchor")),
            "new": stringify(item.get("new")),
            "reason": stringify(item.get("reason")),
        })
    return {
        "summary": stringify(plan.get("summary")).strip(),
        "operations": normalized_ops,
        "verification": plan.get("verification") if isinstance(plan.get("verification"), list) else [],
    }


def validate_codepatch_plan_scope(plan, message):
    outside_whitelist = []
    for operation in plan.get("operations") or []:
        resolved, error = resolve_codepatch_path(operation.get("path"))
        if error:
            return {"ok": False, "reason": f"{operation.get('path')}: {error}"}
        rel = safe_relpath(resolved)
        if not is_codepatch_default_whitelisted(resolved):
            outside_whitelist.append(rel)
        danger = find_dangerous_codepatch_content(operation)
        if danger:
            return {"ok": False, "reason": f"{rel} 包含禁止内容：{danger}"}
    outside_whitelist = list(dict.fromkeys(outside_whitelist))
    if outside_whitelist and not user_confirmed_extra_write_access(message):
        return {
            "status": "WAITING_USER_CONFIRMATION",
            "paths": outside_whitelist,
            "reason": "这些文件在项目目录内，但不在默认 CodePatch 写入白名单中。",
        }
    return {"ok": True}


def resolve_codepatch_path(raw_path):
    text = stringify(raw_path).strip().replace("\\", "/").lstrip("/")
    if not text:
        return None, "路径为空"
    candidate = Path(text)
    if candidate.is_absolute():
        try:
            resolved = candidate.resolve(strict=False)
            resolved.relative_to(PROJECT_ROOT)
        except ValueError:
            return None, "禁止修改项目目录外文件"
        except OSError as exc:
            return None, f"路径解析失败：{exc}"
    else:
        resolved = (PROJECT_ROOT / candidate).resolve(strict=False)
        try:
            resolved.relative_to(PROJECT_ROOT)
        except ValueError:
            return None, "禁止修改项目目录外文件"
    rel_parts = resolved.relative_to(PROJECT_ROOT).parts
    if any(part in CODEPATCH_NEVER_WRITE_DIRS for part in rel_parts):
        return None, "位于禁止写入目录"
    if resolved.suffix.lower() not in CODEPATCH_ALLOWED_SUFFIXES:
        return None, f"不支持写入该文件类型：{resolved.suffix}"
    return resolved, None


def is_codepatch_default_whitelisted(file_path):
    rel = safe_relpath(file_path).replace("\\", "/")
    for root in CODEPATCH_WRITE_ROOTS:
        root_norm = root.replace("\\", "/").rstrip("/")
        if rel == root_norm or rel.startswith(root_norm + "/"):
            return True
    return False


def user_confirmed_extra_write_access(message):
    text = stringify(message)
    confirm_terms = ["允许修改", "确认修改", "同意修改", "可以修改", "放行", "允许写入", "确认写入"]
    scope_terms = ["白名单外", "这些文件", "上述文件", "项目内文件", "额外文件"]
    return any(term in text for term in confirm_terms) and any(term in text for term in scope_terms)


def find_dangerous_codepatch_content(operation):
    content = "\n".join([
        stringify(operation.get("old")),
        stringify(operation.get("anchor")),
        stringify(operation.get("new")),
    ])
    for pattern in CODEPATCH_DANGEROUS_PATTERNS:
        if re.search(pattern, content, flags=re.IGNORECASE):
            return pattern
    return ""


def apply_codepatch_plan(plan):
    operations = plan.get("operations") or []
    originals = {}
    updates = {}
    changed_files = []

    try:
        for operation in operations:
            resolved, error = resolve_codepatch_path(operation.get("path"))
            if error:
                return {"ok": False, "error": f"{operation.get('path')}: {error}"}
            rel = safe_relpath(resolved)
            if resolved not in originals:
                existed = resolved.exists()
                content = resolved.read_text(encoding="utf-8", errors="ignore") if existed else ""
                originals[resolved] = {"content": content, "existed": existed}
                updates[resolved] = content
            current = updates[resolved]
            updated, changed, error = apply_codepatch_operation(current, operation, resolved.exists())
            if error:
                return {"ok": False, "error": f"{rel}: {error}"}
            updates[resolved] = updated
            if changed and rel not in changed_files:
                changed_files.append(rel)

        diff_map = {}
        for resolved, updated in updates.items():
            original = originals[resolved]["content"]
            if original == updated:
                continue
            resolved.parent.mkdir(parents=True, exist_ok=True)
            resolved.write_text(updated, encoding="utf-8")
            rel = safe_relpath(resolved)
            diff_map[rel] = "\n".join(difflib.unified_diff(
                original.splitlines(),
                updated.splitlines(),
                fromfile=f"{rel} before",
                tofile=f"{rel} after",
                lineterm="",
            ))[:12000]
        return {"ok": True, "changed_files": list(diff_map.keys()), "diff": diff_map, "originals": originals}
    except Exception as exc:
        rollback_codepatch_changes(originals)
        return {"ok": False, "error": str(exc), "verification": {"rolled_back": True}}


def rollback_codepatch_changes(originals):
    errors = []
    for resolved, meta in (originals or {}).items():
        try:
            content = meta.get("content") if isinstance(meta, dict) else stringify(meta)
            existed = bool(meta.get("existed")) if isinstance(meta, dict) else True
            if existed:
                resolved.write_text(content, encoding="utf-8")
            elif resolved.exists():
                resolved.unlink()
        except Exception as exc:
            errors.append(f"{safe_relpath(resolved)}: 回滚失败：{exc}")
    return {"ok": not errors, "errors": errors}


def apply_codepatch_operation(current, operation, file_exists):
    op = operation.get("op")
    old = stringify(operation.get("old"))
    anchor = stringify(operation.get("anchor"))
    new = stringify(operation.get("new"))
    if op == "create":
        if file_exists and current.strip():
            return current, False, "create 目标文件已存在，拒绝覆盖"
        return new, bool(new != current), ""
    if op == "append":
        if not new:
            return current, False, "append 缺少 new 内容"
        if new in current:
            return current, False, ""
        separator = "" if not current or current.endswith("\n") else "\n"
        return current + separator + new, True, ""
    if op == "replace":
        if not old:
            return current, False, "replace 缺少 old 内容"
        if old not in current:
            return current, False, "replace 的 old 内容未在文件中找到"
        if current.count(old) > 1:
            return current, False, "replace 的 old 内容匹配多处，拒绝模糊替换"
        return current.replace(old, new, 1), old != new, ""
    if op in {"insert_after", "insert_before"}:
        if not anchor:
            return current, False, "insert 缺少 anchor"
        if anchor not in current:
            return current, False, "insert 的 anchor 未在文件中找到"
        if current.count(anchor) > 1:
            return current, False, "insert 的 anchor 匹配多处，拒绝模糊插入"
        if new in current:
            return current, False, ""
        replacement = anchor + new if op == "insert_after" else new + anchor
        return current.replace(anchor, replacement, 1), True, ""
    return current, False, f"不支持的操作：{op}"


def verify_codepatch_result(changed_files, plan):
    errors = []
    checked = []
    for rel in changed_files:
        resolved, error = resolve_codepatch_path(rel)
        if error:
            errors.append(f"{rel}: {error}")
            continue
        checked.append(rel)
        if resolved.suffix.lower() == ".py":
            try:
                py_compile.compile(str(resolved), doraise=True)
            except Exception as exc:
                errors.append(f"{rel}: Python 编译失败：{exc}")
        try:
            text = resolved.read_text(encoding="utf-8", errors="ignore")
        except OSError as exc:
            errors.append(f"{rel}: 读取验证失败：{exc}")
            continue
        if find_dangerous_codepatch_content({"new": text}):
            errors.append(f"{rel}: 验证发现禁止内容")
    if not changed_files:
        errors.append("补丁没有产生任何文件变化")
    return {
        "passed": not errors,
        "checked_files": checked,
        "errors": errors,
        "suggested_checks": plan.get("verification") or [],
    }


def is_translation_think_cleanup_task(message):
    text = stringify(message)
    lower = text.lower()
    has_translate = any(term in lower for term in ["翻译", "translation", "translate", "中文"])
    has_think = any(term in lower for term in ["<think>", "think", "思考"])
    wants_clean = any(term in lower for term in ["不要包括", "只需要", "删除", "去掉", "隐藏", "清理", "剥离", "strip", "remove"])
    return has_translate and has_think and wants_clean


def is_detail_selling_points_translate_task(message):
    text = stringify(message)
    lower = text.lower()
    selling_terms = ["卖点", "selling_points", "selling points"]
    translate_terms = ["翻译", "中文", "translate", "translation", "chinese"]
    detail_terms = [
        "站内详情",
        "站内详细",
        "商品详情",
        "商品详细",
        "详情",
        "详细",
        "detail",
        "product detail",
    ]
    return (
        any(term in lower for term in selling_terms)
        and any(term in lower for term in translate_terms)
        and any(term in lower for term in detail_terms)
    )


def apply_detail_selling_points_translate_adapter(decision, message, readonly_context, base_files):
    verified = detail_selling_points_translate_is_applied()
    return {
        "status": "COMPLETED" if verified else "FAILED",
        "output": "已接入站内详情卖点翻译按钮，并通过静态验证。" if verified else "卖点翻译按钮静态验证未通过。",
        "error": "" if verified else "未检测到 /api/translate-text、data-translate-selling-points 或翻译结果样式。",
        "changed_files": ["app.py", "static/index.html", "static/styles.css"] if verified else [],
        "verification": {
            "detail_selling_points_translate_is_applied": verified,
            "requires_restart": True,
        },
    }


def detail_selling_points_translate_is_applied():
    try:
        app_text = (PROJECT_ROOT / "app.py").read_text(encoding="utf-8", errors="ignore")
        index_text = (PROJECT_ROOT / "static" / "index.html").read_text(encoding="utf-8", errors="ignore")
        css_text = (PROJECT_ROOT / "static" / "styles.css").read_text(encoding="utf-8", errors="ignore")
    except OSError:
        return False
    return all([
        '"/api/translate-text"' in app_text or "@app.post(\"/api/translate-text\")" in app_text,
        "translateSellingPoints" in index_text,
        "stripThinkBlocks" in index_text,
        "data-translate-selling-points" in index_text,
        "detail-selling-points-translation" in index_text,
        ".detail-selling-points-translation" in css_text,
        "strip_llm_think_blocks" in app_text,
    ])


def build_workflow_execution_summary(workflow, status, steps):
    if status == "SUCCESS":
        return f"WorkflowEngine 已完成 `{workflow['workflow_code']}`，共执行 {len(steps)} 个步骤。"
    if status == "WAITING_USER_CONFIRMATION":
        return f"WorkflowEngine 已启动 `{workflow['workflow_code']}`，并等待用户确认额外写入范围。"
    if status == "WAITING_ADAPTER":
        return f"WorkflowEngine 已启动 `{workflow['workflow_code']}`，并执行到能力适配器边界；需要接入安全执行适配器后继续。"
    return f"WorkflowEngine 执行 `{workflow['workflow_code']}` 失败，已记录步骤错误。"


def create_expert_workflow_instance(session_id, user_id, workflow, decision, workflow_input):
    try:
        with db() as conn, conn.cursor() as cursor:
            cursor.execute(
                f"""
                INSERT INTO {EXPERT_WORKFLOW_INSTANCE_TABLE}
                  (session_id, user_id, workflow_code, workflow_name, intent, mode, status, input_json)
                VALUES (%s, %s, %s, %s, %s, %s, 'RUNNING', %s)
                """,
                (
                    session_id or 0,
                    user_id or 0,
                    workflow["workflow_code"],
                    workflow["name"],
                    (decision or {}).get("intent") or "",
                    (decision or {}).get("mode") or "",
                    json.dumps(workflow_input, ensure_ascii=False),
                ),
            )
            conn.commit()
            return cursor.lastrowid
    except Exception:
        return None


def create_expert_workflow_step(workflow_id, step_order, step, step_input):
    if not workflow_id:
        return None
    try:
        with db() as conn, conn.cursor() as cursor:
            cursor.execute(
                f"""
                INSERT INTO {EXPERT_WORKFLOW_STEP_TABLE}
                  (workflow_id, step_order, step_code, step_name, executor, status, input_json, started_at)
                VALUES (%s, %s, %s, %s, %s, 'RUNNING', %s, NOW())
                """,
                (
                    workflow_id,
                    step_order,
                    step["code"],
                    step["name"],
                    step["executor"],
                    json.dumps(step_input, ensure_ascii=False),
                ),
            )
            conn.commit()
            return cursor.lastrowid
    except Exception:
        return None


def update_expert_workflow_step(step_id, result):
    if not step_id:
        return
    status = result.get("status") or "COMPLETED"
    try:
        with db() as conn, conn.cursor() as cursor:
            cursor.execute(
                f"""
                UPDATE {EXPERT_WORKFLOW_STEP_TABLE}
                SET status = %s,
                    output_json = %s,
                    error_message = %s,
                    completed_at = IF(%s IN ('COMPLETED', 'FAILED', 'WAITING_ADAPTER', 'WAITING_USER_CONFIRMATION'), NOW(), completed_at)
                WHERE id = %s
                """,
                (
                    status,
                    json.dumps(result, ensure_ascii=False),
                    result.get("error") or "",
                    status,
                    step_id,
                ),
            )
            conn.commit()
    except Exception:
        return


def update_expert_workflow_instance(workflow_id, status, output, error_message=""):
    if not workflow_id:
        return
    db_status = "SUCCESS" if status == "SUCCESS" else ("FAILED" if status == "FAILED" else ("WAITING_USER_CONFIRMATION" if status == "WAITING_USER_CONFIRMATION" else "WAITING_ADAPTER"))
    try:
        with db() as conn, conn.cursor() as cursor:
            cursor.execute(
                f"""
                UPDATE {EXPERT_WORKFLOW_INSTANCE_TABLE}
                SET status = %s,
                    output_json = %s,
                    error_message = %s,
                    completed_at = IF(%s IN ('SUCCESS', 'FAILED'), NOW(), completed_at)
                WHERE id = %s
                """,
                (
                    db_status,
                    json.dumps(output, ensure_ascii=False),
                    error_message or "",
                    db_status,
                    workflow_id,
                ),
            )
            conn.commit()
    except Exception:
        return


def apply_detail_attribute_display_recipe():
    changed_files = []
    errors = []

    app_path = PROJECT_ROOT / "app.py"
    index_path = PROJECT_ROOT / "static" / "index.html"
    css_path = PROJECT_ROOT / "static" / "styles.css"

    try:
        app_text = app_path.read_text(encoding="utf-8", errors="ignore")
        app_marker = '"attributes": serialize_value(row.get(meta.get("attributes")))'
        if app_marker not in app_text:
            needle = '        "title": row.get(meta["title"]),\n'
            insert = (
                needle
                + '        "attributes": serialize_value(row.get(meta.get("attributes"))) if meta.get("attributes") else "",\n'
                + '        "selling_points": serialize_value(row.get(meta.get("selling_points"))) if meta.get("selling_points") else "",\n'
            )
            if needle in app_text:
                app_path.write_text(app_text.replace(needle, insert, 1), encoding="utf-8")
                changed_files.append("app.py")
            else:
                errors.append("app.py 中未找到 normalize_detail 的 title 插入点。")
    except OSError as exc:
        errors.append(f"app.py 写入失败：{exc}")

    try:
        index_text = index_path.read_text(encoding="utf-8", errors="ignore")
        if "attributesText" not in index_text:
            needle = "      const storeName = raw.shop_name || raw.store_name || raw['店铺名称'] || raw['店铺'] || '';\n"
            insert = (
                needle
                + "      const attributesText = detailValue(item.attributes, raw.attributes, raw['属性信息']);\n"
                + "      const sellingPointsText = detailValue(item.selling_points, raw.selling_points, raw['卖点']);\n"
            )
            if needle in index_text:
                index_text = index_text.replace(needle, insert, 1)
            else:
                errors.append("index.html 中未找到 renderProductDetail 的 storeName 插入点。")
        if "detail-extra-info" not in index_text:
            needle = "        <section class=\"detail-section\">\n          <div class=\"detail-section-head\">\n            <h4>数据总览</h4>"
            block = (
                "        ${attributesText || sellingPointsText ? `<section class=\"detail-section detail-extra-info\">\n"
                "          <div class=\"detail-section-head\"><h4>属性与卖点</h4></div>\n"
                "          <div class=\"detail-info-list\">\n"
                "            ${attributesText ? `<div class=\"detail-info-block\"><span>属性信息</span><p>${escapeHtml(attributesText)}</p></div>` : ''}\n"
                "            ${sellingPointsText ? `<div class=\"detail-info-block\"><span>卖点</span><p>${escapeHtml(sellingPointsText)}</p></div>` : ''}\n"
                "          </div>\n"
                "        </section>` : ''}\n\n"
            )
            if needle in index_text:
                index_text = index_text.replace(needle, block + needle, 1)
            else:
                errors.append("index.html 中未找到数据总览区块插入点。")
        if not errors or "index.html" not in " ".join(errors):
            original = index_path.read_text(encoding="utf-8", errors="ignore")
            if index_text != original:
                index_path.write_text(index_text, encoding="utf-8")
                changed_files.append("static/index.html")
    except OSError as exc:
        errors.append(f"index.html 写入失败：{exc}")

    try:
        css_text = css_path.read_text(encoding="utf-8", errors="ignore")
        if ".detail-info-block" not in css_text:
            css_addition = """

.detail-extra-info {
  padding: 20px 24px;
}

.detail-extra-info .detail-section-head {
  margin-bottom: 14px;
  padding-bottom: 12px;
}

.detail-info-list {
  display: grid;
  grid-template-columns: repeat(2, minmax(0, 1fr));
  gap: 14px;
}

.detail-info-block {
  min-width: 0;
  padding: 14px 16px;
  background: #f8fafc;
  border: 1px solid #e2e8f0;
  border-radius: 8px;
}

.detail-info-block span {
  display: block;
  margin-bottom: 8px;
  font-size: 12px;
  font-weight: 700;
  color: #64748b;
}

.detail-info-block p {
  margin: 0;
  color: #1e293b;
  font-size: 14px;
  line-height: 1.7;
  white-space: pre-wrap;
  word-break: break-word;
}
"""
            css_path.write_text(css_text.rstrip() + css_addition, encoding="utf-8")
            changed_files.append("static/styles.css")
    except OSError as exc:
        errors.append(f"styles.css 写入失败：{exc}")

    verified = detail_attribute_task_is_applied()
    ok = verified and not errors
    return {
        "ok": ok,
        "recipe": "detail_attribute_display",
        "can_write": True,
        "write_scope": "D:\\choice_product 内白名单文本文件",
        "changed_files": changed_files,
        "already_applied": not changed_files and verified,
        "verification": {
            "detail_attribute_task_is_applied": verified,
            "requires_restart": True,
        },
        "errors": errors,
        "summary": (
            "Codex执行Agent已完成商品详情属性与卖点展示写入并通过静态验证。"
            if changed_files else
            "Codex执行Agent已确认商品详情属性与卖点展示已存在，并通过静态验证。"
        ) if ok else "Codex执行Agent执行失败，已返回错误原因。",
    }


def handle_execution_handoff(handler_code, decision, message, readonly_context, base_files):
    current_workflow_id = (decision or {}).get("_current_workflow_id")
    workflow = fetch_expert_workflow_instance_by_id(current_workflow_id, (decision or {}).get("_session_id"), (decision or {}).get("_user_id"))
    if workflow:
        status = workflow.get("status")
        return {
            "handler": handler_code,
            "status": "COMPLETED" if status == "SUCCESS" else status,
            "summary": f"Response Delivery 已接收到工作流 `{workflow.get('workflow_code')}` 的执行状态：{status}。",
            "data": {
                "workflow_id": workflow.get("id"),
                "workflow_code": workflow.get("workflow_code"),
                "workflow_name": workflow.get("workflow_name"),
                "status": status,
                "output": parse_json(workflow.get("output_json")),
            },
        }
    return {
        "handler": handler_code,
        "status": "NO_EXECUTABLE_WORKFLOW",
        "summary": "业务执行层已接入，但当前任务没有命中可执行工作流，未写入文件。",
        "data": {
            "executor": decision.get("handoff_target") or "技术执行人员A / Codex",
            "user_action": "用户只需审核原因；不需要自己找文件或改代码。",
            "next_system_action": "为该类任务新增业务工作流后，业务执行层才能自动执行、写文件并回传验证结果。",
            "supported_workflows": list_expert_execution_workflows(),
        },
    }


def fetch_latest_expert_workflow_instance(session_id, user_id):
    if not session_id:
        return None
    try:
        with db() as conn, conn.cursor() as cursor:
            cursor.execute(
                f"""
                SELECT *
                FROM {EXPERT_WORKFLOW_INSTANCE_TABLE}
                WHERE session_id = %s AND user_id = %s
                ORDER BY id DESC
                LIMIT 1
                """,
                (session_id, user_id or 0),
            )
            return cursor.fetchone()
    except Exception:
        return None


def fetch_expert_workflow_instance_by_id(workflow_id, session_id, user_id):
    if not workflow_id or not session_id:
        return None
    try:
        with db() as conn, conn.cursor() as cursor:
            cursor.execute(
                f"""
                SELECT *
                FROM {EXPERT_WORKFLOW_INSTANCE_TABLE}
                WHERE id = %s AND session_id = %s AND user_id = %s
                LIMIT 1
                """,
                (workflow_id, session_id, user_id or 0),
            )
            return cursor.fetchone()
    except Exception:
        return None


def is_detail_attribute_task(message):
    text = stringify(message)
    return (
        ("attributes" in text or "属性信息" in text or "属性" in text)
        and ("selling_points" in text or "卖点" in text)
        and ("商品详情" in text or "站内详情" in text or "renderProductDetail" in text or "product_detail" in text)
    )


def detail_attribute_task_is_applied():
    try:
        app_text = (PROJECT_ROOT / "app.py").read_text(encoding="utf-8", errors="ignore")
        index_text = (PROJECT_ROOT / "static" / "index.html").read_text(encoding="utf-8", errors="ignore")
        css_text = (PROJECT_ROOT / "static" / "styles.css").read_text(encoding="utf-8", errors="ignore")
    except OSError:
        return False
    return all([
        '"attributes": serialize_value(row.get(meta.get("attributes")))' in app_text,
        '"selling_points": serialize_value(row.get(meta.get("selling_points")))' in app_text,
        "attributesText" in index_text,
        "sellingPointsText" in index_text,
        "detail-extra-info" in index_text,
        ".detail-info-block" in css_text,
    ])


def handle_business_expert(handler_code, decision, message, readonly_context, base_files):
    expert_names = {role["code"]: role["name"] for role in EXPERT_TEAM_ROLES}
    return {
        "handler": handler_code,
        "status": "COMPLETED",
        "summary": f"{expert_names.get(handler_code, handler_code)} 已完成本轮业务判断。",
        "data": {
            "uses_readonly_context": bool(readonly_context),
            "output_type": "business_review",
        },
    }


def handle_planning_expert(handler_code, decision, message, readonly_context, base_files):
    expert_names = {role["code"]: role["name"] for role in EXPERT_TEAM_ROLES}
    return {
        "handler": handler_code,
        "status": "COMPLETED",
        "summary": f"{expert_names.get(handler_code, handler_code)} 已完成规划判断。",
        "data": {
            "mode": decision.get("mode"),
            "intent": decision.get("intent"),
        },
    }


def build_agent_response_summary(decision, handler_results):
    waiting = [item for item in handler_results if item.get("status") in {"WAITING_EXECUTOR_BRIDGE", "WAITING_EXECUTOR_RECIPE", "WAITING_ADAPTER", "WAITING_USER_CONFIRMATION"}]
    failed = [item for item in handler_results if item.get("status") == "FAILED"]
    unsupported = [item for item in handler_results if item.get("status") in {"UNSUPPORTED_AUTOMATION", "NO_EXECUTABLE_WORKFLOW"}]
    quality = summarize_handler_quality(handler_results)
    if failed:
        success = False
        message = "部分 Handler 执行失败，已返回可见错误。"
    elif unsupported:
        success = True
        message = "专家团队已完成判断，但当前任务没有命中可自动执行的工作流，已标记为待补能力项。"
    elif waiting:
        success = True
        message = "专家团队已完成判断，但当前任务缺少可自动落地的执行适配器，已标记为待补能力项。"
    else:
        success = True
        message = "专家团队已完成分工、执行和结果校验，可反馈给用户审核。"
    return {
        "success": success,
        "message": message,
        "intent": decision.get("intent"),
        "mode": decision.get("mode"),
        "handler_statuses": [{"handler": item.get("handler"), "status": item.get("status")} for item in handler_results],
        "quality": quality,
    }


def summarize_handler_quality(handler_results):
    total = len(handler_results or [])
    with_evidence = 0
    closed_loop = 0
    low_confidence = []
    for item in handler_results or []:
        quality = item.get("quality") or {}
        if quality.get("has_evidence"):
            with_evidence += 1
        if quality.get("closed_loop"):
            closed_loop += 1
        if float(quality.get("confidence") or 0) < 0.5:
            low_confidence.append(item.get("handler"))
    return {
        "handler_count": total,
        "with_evidence": with_evidence,
        "closed_loop": closed_loop,
        "low_confidence_handlers": [item for item in low_confidence if item],
    }


def get_codex_executor_result(expert_execution):
    for item in (expert_execution or {}).get("results") or []:
        if item.get("handler") == "codex_executor":
            return item
    return None


def build_codex_executor_status_answer(ceo_decision, expert_execution):
    codex_result = get_codex_executor_result(expert_execution) or {}
    data = codex_result.get("data") or {}
    status = codex_result.get("status") or ""
    changed_files = data.get("changed_files") or []

    if status == "COMPLETED":
        changed_text = "、".join(changed_files) if changed_files else "相关代码已具备该能力，未重复写入"
        workflow_name = data.get("workflow_name") or data.get("workflow_code") or "本轮任务"
        verification = data.get("verification") or {}
        verified_text = "已完成验证" if verification else "已完成处理"
        return (
            "已完成。\n\n"
            f"本次已处理：{workflow_name}，{verified_text}。\n"
            f"涉及文件：{changed_text}\n"
            "你现在可以刷新页面或重新发起同类请求，检查本轮需求对应的效果。"
        )

    if status == "WAITING_USER_CONFIRMATION":
        pending = data.get("pending_confirmation") or {}
        paths = "、".join(pending.get("paths") or [])
        return (
            "这个需求已经生成了可执行修改计划，但其中包含默认白名单外的项目文件。\n\n"
            f"需要你确认是否允许修改：{paths or '待确认文件'}。\n"
            "你回复“允许修改这些文件”后，专家团队会继续执行。"
        )

    if status in {"UNSUPPORTED_AUTOMATION", "NO_EXECUTABLE_WORKFLOW", "WAITING_ADAPTER"}:
        return (
            "这个需求已经被专家团队理解，但当前系统还没有覆盖到对应的自动执行能力。\n\n"
            "我已记录为待补能力项；用户不需要提供代码或自己找文件。等执行能力补齐后，再次发送同类需求即可直接执行。"
        )

    return stringify(codex_result.get("summary") or "专家团队已完成本轮处理。")


def build_expert_post_learning(session_id, message, ceo_decision, expert_execution):
    return {
        "stage": "Post-Learning",
        "session_id": session_id,
        "recorded": True,
        "scope": "session_isolated",
        "insight": (
            "专家团队必须形成闭环：先查证据，再判断能力，再执行或标记待补能力，最后只给用户可审核结果；内部流程不直接展示给用户。"
        ),
        "intent": (ceo_decision or {}).get("intent"),
        "handler_count": (expert_execution or {}).get("handler_count"),
    }


def format_expert_execution(expert_execution):
    return json.dumps(expert_execution or {}, ensure_ascii=False, indent=2)


def format_expert_learning(learning_result):
    return json.dumps(learning_result or {}, ensure_ascii=False, indent=2)


def build_expert_team_system_prompt(user, project_code, project_context):
    role_lines = []
    for role in EXPERT_TEAM_ROLES:
        role_lines.append(
            f"- {role['name']}({role['code']}，{role['level']}): 擅长{role['focus']} 边界：{role['limits']}"
        )
    context = project_context or "当前项目是 choice_product 选品与数据监控系统，但专家团队要保持通用能力，不能只围绕单一项目。"
    return f"""
你是一个“AI专家团队”的专家领导层，直接服务当前用户，而不是只服务某个固定项目。

你的核心职责：
1. 先理解用户真实意图、长期目标和当前项目上下文。
2. 判断应该调用哪些专家，以及每个专家的任务边界。
3. 不让执行专家各说各话，要给出统一、清楚、可执行的团队结论。
4. 输出要短、清楚、有重点，优先降低用户阅读成本。
5. 如果用户问的是项目建设，你要从系统产品、技术架构、数据架构、AI效率、UX角度组织团队。
6. 如果用户问的是商品分析，你要调度数据、IP、材质、运营、商品产品经理、选品总评专家。
7. 如果用户问的是代码、接口、页面、字段、项目文件位置或为什么某功能缺失，你要优先调用“项目文件侦察员”读取系统提供的只读文件上下文，再让技术架构专家/用户体验负责人判断，不要先反问用户已经能从文件里看到的信息。
8. 必须遵守闭环纪律：没有文件证据、数据库证据、配置证据或 Handler 执行结果时，不能把推测说成事实。
9. 每次回答都要区分“已确认事实 / 合理推断 / 待验证事项 / 下一步动作”，但不要机械套模板；能直接给结论时先给结论。

当前用户：
- user_id: {user.get('id')}
- username: {user.get('username')}

当前项目上下文：
- project_code: {project_code}
- context: {context}

专家能力档案：
{chr(10).join(role_lines)}

回复格式：
- 面向用户时，不展示 WorkflowEngine、Dispatcher、Handler、CEO决策、IntentHandlerFactory、第6层、第7层等内部词。
- 优先输出用户能理解的结论：已经确认了什么、为什么、完成了什么、用户在哪里查看。
- 如果 Action Execution 已写入文件，要明确写“已执行/已修改”，并给修改文件和验证结果。
- 如果未写入，只能说“当前还没有覆盖到对应的自动执行能力，已记录为待补能力项”，不要再写“等待技术执行Agent接入”。
- 不要要求用户粘贴代码、找接口、找技术人员、判断能力或整理项目材料。

不要假装已经调用外部工具或数据库；但如果系统消息里提供了“只读工具结果/项目文件侦察结果”，你可以引用其中的文件路径、行号和结论。不得声称自己执行了写入、修改、删除、上线、提交代码等动作。
如果系统消息里出现“【权限缺口】”，你必须明确说明：没有完成该任务不是因为配置不存在，而是因为当前会话没有获得对应访问权限。你只能给用户两个选择：一是提供可复制执行的只读 SQL/命令并让用户把结果贴回；二是询问用户是否允许本轮开通一次性只读权限，用户授权后系统再继续执行。禁止在未授权时输出数据库结论、索引结论、慢查询结论或“已完成”。
如果系统消息里提供了“会话基底文件读取结果”，说明后端已经替你读取了用户指定的本地项目文件；你必须优先依据这些内容回答，不要再说“无法读取本地 Windows 路径”。
如果“会话基底文件读取结果”里写明“已对指定文件做全文件关键词检索”，你不能说“只看到前80行”，也不能要求用户粘贴同一个文件代码；如果上下文仍不够，只能要求用户补充更具体的关键词或把相关文件加入基底文件列表。
你必须遵循 TotalAgent 全生命周期：User Input -> TotalAgent Entry -> Security Gate -> RAG Enrichment -> Core Decision Layer -> Dispatcher -> Action Execution -> Response Delivery -> Post-Learning。用户只提供需求、查看结果、审核结果；其他文件检索、字段判断、能力判断、专家分发和执行交接由系统承担。
如果系统明确告诉你“执行器已经写入文件并通过验证”，你必须把修改文件、命中工作流和验证结果反馈给用户。如果系统告诉你“NO_EXECUTABLE_WORKFLOW”，你必须说明第 6 层已派发成功、第 7 层缺少对应业务工作流；不要说等待技术执行Agent接入。禁止让用户自己找技术人员、自己打开文件、自己改代码。
专家闭环纪律：
- 数据分析专家：没查库、没配置、没口径时，不能给数量和覆盖率结论。
- IP合规专家：必须区分明确IP命中、疑似风格、无法识别；无法识别不能直接打 A/S。
- 材质/视觉专家：必须区分工厂材质、非工厂材质、疑似材质；只给材质结论，不越权判断IP侵权。
- 技术架构专家：没读到文件/接口/配置时，不能说“缺失”或“已实现”，只能说“待验证”。
- 项目文件侦察员：只负责读文件和定位片段，不能替执行器宣布完成。
- Codex执行桥：只有返回 changed_files 或 verification 时，才允许说已执行或已验证。
""".strip()


def call_expert_team_ai(user, message, history, project_code, project_context, readonly_context="", ceo_decision=None, expert_execution=None, learning_result=None):
    api_key = resolve_minimax_api_key()
    if not api_key:
        raise RuntimeError("缺少 MINIMAX_API_KEY")
    messages = [{"role": "system", "content": build_expert_team_system_prompt(user, project_code, project_context)}]
    messages.append({"role": "system", "content": EXPERT_EXECUTION_HANDOFF_RULES})
    messages.append({
        "role": "system",
        "content": (
            "重要规则：专家团队接口是同步接口，不存在后台异步读取文件这一步。"
            "如果系统消息已经提供只读文件结果，你必须直接基于结果回答。"
            "禁止回复“请稍等、正在读取、侦察员正在读取、我稍后给你结果、你复制粘贴代码、请你提供接口代码、请你告诉我有哪些能力”。"
            "如果读到的信息仍不足，只能明确说明已读取哪些文件和行号、还缺哪个更具体的文件名或关键词。"
        ),
    })
    messages.append({
        "role": "system",
        "content": (
            "【CEO核心决策层派工单】\n"
            "下面是后端在调用大模型之前完成的本地决策。你必须按这张派工单组织回答，不能退回普通问答模式。\n"
            f"{format_expert_ceo_decision(ceo_decision)}\n\n"
            "硬性要求：need_user_input=false 时，不得让用户提供代码、接口、字段、截图、能力清单或项目结构；"
            "你应该使用只读工具结果和会话记忆给出结论、方案或执行交接单。"
            "同时必须明确执行状态：系统执行结果是 COMPLETED 时就是已执行；系统执行结果是 WAITING_USER_CONFIRMATION 时就是等待用户确认额外写入文件；系统执行结果是 NO_EXECUTABLE_WORKFLOW 时就是第 7 层缺少工作流，不要再写“等待技术执行Agent接入”。"
            "不要说“你需要找技术执行人员”或“你自己打开文件改”；应该说“执行对象是A/Codex，用户只需确认执行并审核结果”。"
        ),
    })
    messages.append({
        "role": "system",
        "content": (
            "【Dispatcher + Action Execution 执行结果】\n"
            "下面不是方案，而是后端 IntentHandlerFactory(local) 已经按 handlers 映射实际执行过的 Handler 结果。"
            "你必须把这些结果反馈给用户。\n"
            f"{format_expert_execution(expert_execution)}\n\n"
            "【Post-Learning 结果】\n"
            f"{format_expert_learning(learning_result)}"
        ),
    })
    if readonly_context:
        messages.append({
            "role": "system",
            "content": (
                "以下是系统白名单只读工具在本轮对话中读取到的结果。"
                "你可以引用这些结果，但不能声称自己执行了写入、修改、删除、上线、提交代码等动作。\n\n"
                f"{readonly_context}"
            ),
        })
    for item in history[-10:]:
        role = "assistant" if item.get("role") == "assistant" else "user"
        messages.append({"role": role, "content": stringify(item.get("content"))})
    messages.append({"role": "user", "content": message})
    url = f"{MINIMAX_BASE_URL.rstrip('/')}/chat/completions"
    payload = {"model": MINIMAX_MODEL, "messages": messages, "temperature": 0.2, "stream": False}
    try:
        data = call_minimax_api_with_retry(url, api_key, payload, timeout=60)
        return data["choices"][0]["message"]["content"].strip()
    except Exception as exc:
        return fallback_expert_team_answer(message, str(exc))


def sanitize_expert_team_answer(text):
    cleaned = stringify(text)
    cleaned = re.sub(r"<think>[\s\S]*?</think>", "", cleaned, flags=re.IGNORECASE).strip()
    cleaned = re.sub(r"^\s*<think>[\s\S]*?(?=(\*\*|团队判断|执行交接单|《执行交接单》|已完成|$))", "", cleaned, flags=re.IGNORECASE).strip()
    cleaned = remove_internal_workflow_diagnostics(cleaned)
    return cleaned or stringify(text).strip()


def remove_internal_workflow_diagnostics(text):
    cleaned = stringify(text)
    internal_terms = [
        "WorkflowEngine",
        "WAITING_ADAPTER",
        "WAITING_USER_CONFIRMATION",
        "IntentHandlerFactory",
        "Dispatcher",
        "Handler",
        "CEO决策",
        "专家分发",
        "Action Execution",
        "Response Delivery",
        "第 6 层",
        "第 7 层",
        "安全执行适配器",
        "codex_executor",
        "workflow_engine",
        "execution_handoff",
    ]
    if not any(term in cleaned for term in internal_terms):
        return cleaned
    lines = []
    for line in cleaned.splitlines():
        if any(term in line for term in internal_terms):
            continue
        if line.strip().startswith("- ") and (":" in line and any(key in line for key in ["handler", "status"])):
            continue
        lines.append(line)
    return "\n".join(lines).strip()


def sanitize_expert_message_for_display(text):
    cleaned = sanitize_expert_team_answer(text)
    if is_stale_detail_translation_answer(cleaned, {"status": "WAITING_ADAPTER"}):
        return (
            "这条历史回执来自旧版本的专家团队回复，已隐藏不相关的卖点翻译提示。\n\n"
            "请以当前最新回复为准。"
        )
    internal_terms = [
        "WorkflowEngine",
        "WAITING_ADAPTER",
        "IntentHandlerFactory",
        "Dispatcher",
        "Handler",
        "codex_executor",
        "workflow_engine",
        "execution_handoff",
        "CEO决策",
        "专家分发",
        "Action Execution",
        "Response Delivery",
        "第 6 层",
        "第 7 层",
        "安全执行适配器",
        "PLAN_AND_HANDOFF",
    ]
    if any(term in cleaned for term in internal_terms):
        if "COMPLETED" in cleaned or "SUCCESS" in cleaned:
            return (
                "已完成。\n\n"
                "本次功能已完成处理。你刷新页面后，回到对应页面查看效果即可。"
            )
        return (
            "这个需求已经被专家团队理解，但当前系统还没有覆盖到对应的自动执行能力。\n\n"
            "我已记录为待补能力项；用户不需要提供代码或自己找文件。等执行能力补齐后，再次发送同类需求即可直接执行。"
        )
    return cleaned


def is_expert_fake_wait_answer(text):
    cleaned = stringify(text)
    fake_wait_words = [
        "请稍等", "正在读取", "稍后", "请粘贴代码", "粘贴代码", "请提供",
        "你需要提供", "你需要找技术执行人员", "如果你自己是技术执行人员",
        "自己去改", "等待技术执行Agent接入", "等待 Codex 执行桥",
    ]
    return any(word in cleaned for word in fake_wait_words)


def ensure_expert_execution_status(answer, ceo_decision=None, expert_execution=None):
    text = sanitize_expert_team_answer(answer)
    codex_result = get_codex_executor_result(expert_execution)
    if codex_result and codex_result.get("status") in {"COMPLETED", "UNSUPPORTED_AUTOMATION", "NO_EXECUTABLE_WORKFLOW", "WAITING_ADAPTER", "WAITING_USER_CONFIRMATION"}:
        stale_words = ["等待技术执行Agent接入", "等待 Codex 执行桥", "未修改，仅生成", "执行交接单"]
        internal_words = ["WorkflowEngine", "WAITING_ADAPTER", "IntentHandlerFactory", "Dispatcher", "Handler", "CEO决策", "第 6 层", "第 7 层"]
        if any(word in text for word in stale_words + internal_words) or is_stale_detail_translation_answer(text, codex_result) or not text:
            return build_codex_executor_status_answer(ceo_decision, expert_execution)
    if "执行状态" in text:
        return text
    return text


def is_stale_detail_translation_answer(text, codex_result=None):
    cleaned = stringify(text)
    status = (codex_result or {}).get("status")
    if status == "COMPLETED":
        return False
    stale_markers = ["站内详情", "卖点", "翻译成中文"]
    return "已完成" in cleaned and all(marker in cleaned for marker in stale_markers)


class ExpertResponseValidator:
    """Checks the final user-facing answer before Response Delivery."""

    internal_terms = [
        "<think>",
        "</think>",
        "WorkflowEngine",
        "WAITING_ADAPTER",
        "WAITING_USER_CONFIRMATION",
        "IntentHandlerFactory",
        "Dispatcher",
        "Handler",
        "codex_executor",
        "workflow_engine",
        "execution_handoff",
        "CEO决策",
        "专家分发",
        "Action Execution",
        "Response Delivery",
        "第 6 层",
        "第 7 层",
        "安全执行适配器",
        "PLAN_AND_HANDOFF",
        "NO_EXECUTABLE_WORKFLOW",
    ]

    user_burden_terms = [
        "你需要找技术执行人员",
        "你自己打开文件",
        "你自己去改",
        "请粘贴代码",
        "请提供接口代码",
        "等待技术执行Agent接入",
        "等待 Codex 执行桥",
        "需要手动实现",
        "手动实现",
    ]

    def validate(self, original_input, raw_output, ceo_decision=None, expert_execution=None):
        output = sanitize_expert_team_answer(raw_output)
        reasons = []
        codex_result = get_codex_executor_result(expert_execution)
        codex_status = (codex_result or {}).get("status") or ""

        if not output.strip():
            reasons.append("输出为空")
        if any(term in output for term in self.internal_terms):
            reasons.append("包含内部工作流或调试信息")
        if any(term in output for term in self.user_burden_terms):
            reasons.append("把系统内部执行责任转嫁给用户")
        strong_claim_terms = ["缺失", "没有实现", "未实现", "不属实", "属实", "已实现", "已存在", "只有一套"]
        has_evidence = response_has_execution_evidence(expert_execution)
        if any(term in output for term in strong_claim_terms) and not has_evidence:
            reasons.append("包含强事实结论，但缺少文件/数据库/执行证据")

        score = estimate_answer_alignment_score(original_input, output)
        if not codex_result and score < 0.08:
            reasons.append(f"回答与用户问题相关性偏低（score={score:.2f}）")
        if codex_result and codex_status == "COMPLETED" and score < 0.08 and not is_explicit_code_write_request(original_input):
            reasons.append(f"执行完成回复与当前问题不匹配（score={score:.2f}）")

        if codex_status == "COMPLETED":
            if not any(term in output for term in ["已完成", "已处理", "可以看到", "刷新", "打开"]):
                reasons.append("执行完成类回复缺少结果和验收入口")
        if codex_status in {"WAITING_ADAPTER", "WAITING_USER_CONFIRMATION", "NO_EXECUTABLE_WORKFLOW", "UNSUPPORTED_AUTOMATION"}:
            if any(term in output for term in ["WorkflowEngine", "WAITING_ADAPTER", "第 6 层", "第 7 层"]):
                reasons.append("未完成类回复暴露内部诊断")
            if codex_status == "WAITING_USER_CONFIRMATION":
                if not any(term in output for term in ["确认", "允许", "白名单", "文件"]):
                    reasons.append("等待确认类回复缺少明确的用户确认动作")
            elif not any(term in output for term in ["当前系统", "自动执行能力", "待补能力", "再次发送"]):
                reasons.append("未完成类回复缺少用户可理解的后续说明")

        return {
            "passed": not reasons,
            "reasons": reasons,
            "alignment_score": round(score, 4),
        }


def validate_and_refine_expert_answer(original_input, raw_output, ceo_decision=None, expert_execution=None, max_retry=3):
    validator = ExpertResponseValidator()
    answer = sanitize_expert_team_answer(raw_output)
    attempts = []

    for retry_count in range(max_retry + 1):
        result = validator.validate(original_input, answer, ceo_decision, expert_execution)
        result["retry_count"] = retry_count
        attempts.append(dict(result))
        if result.get("passed"):
            final_result = dict(result)
            final_result["attempts"] = attempts
            return answer, final_result
        if retry_count >= max_retry:
            fallback = build_user_facing_execution_answer(ceo_decision, expert_execution, result.get("reasons") or [])
            final_result = validator.validate(original_input, fallback, ceo_decision, expert_execution)
            final_result.update({"retry_count": retry_count, "attempts": attempts, "fallback_used": True})
            return fallback, final_result
        answer = correct_expert_answer_for_user_experience(
            original_input,
            answer,
            result.get("reasons") or [],
            ceo_decision,
            expert_execution,
            retry_count + 1,
        )

    return answer, {"passed": True, "retry_count": 0, "attempts": attempts}


def correct_expert_answer_for_user_experience(original_input, raw_output, reasons, ceo_decision=None, expert_execution=None, retry_count=1):
    api_key = resolve_minimax_api_key()
    if not api_key:
        return build_user_facing_execution_answer(ceo_decision, expert_execution, reasons)

    prompt = (
        "你是“用户回答体验校验员”，专门把系统内部执行结果改写成普通用户能看懂的最终回复。\n"
        "你的目标不是重新分析任务，而是修正上一版回复的问题。\n\n"
        "硬性规则：\n"
        "1. 只输出最终给用户看的中文回复，不要解释你的校验过程。\n"
        "2. 禁止输出 <think>、WorkflowEngine、WAITING_ADAPTER、IntentHandlerFactory、codex_executor、CEO决策、第6层、第7层等内部词。\n"
        "3. 不要让用户自己找技术人员、自己粘贴代码、自己打开文件修改。\n"
        "4. 如果已经完成，说明完成了什么、涉及哪些文件、用户怎么查看结果。\n"
        "5. 如果没有完成，用用户能懂的话说明当前还不能自动执行，并说已经记录为待补能力项。\n"
        "6. 保持简短，最多 6 行。\n\n"
        f"重试次数：{retry_count}\n"
        f"原始用户需求：{stringify(original_input)}\n"
        f"失败原因：{'; '.join(stringify(item) for item in reasons)}\n"
        f"上一版回复：\n{stringify(raw_output)}\n\n"
        f"执行上下文摘要：\n{build_user_facing_execution_answer(ceo_decision, expert_execution, reasons)}"
    )

    try:
        url = f"{MINIMAX_BASE_URL.rstrip('/')}/chat/completions"
        payload = {
            "model": MINIMAX_MODEL,
            "messages": [
                {"role": "system", "content": "你是严格的用户体验回复改写器，只输出改写后的最终回复。"},
                {"role": "user", "content": prompt},
            ],
            "temperature": 0.1,
            "stream": False,
        }
        data = call_minimax_api_with_retry(url, api_key, payload, timeout=45)
        return sanitize_expert_team_answer(data["choices"][0]["message"]["content"].strip())
    except Exception:
        return build_user_facing_execution_answer(ceo_decision, expert_execution, reasons)


def build_user_facing_execution_answer(ceo_decision=None, expert_execution=None, reasons=None):
    codex_result = get_codex_executor_result(expert_execution)
    if codex_result:
        status = codex_result.get("status")
        data = codex_result.get("data") or {}
        changed_files = data.get("changed_files") or []
        if status == "COMPLETED":
            changed_text = "、".join(changed_files) if changed_files else "相关代码"
            workflow_name = data.get("workflow_name") or data.get("workflow_code") or "本轮任务"
            return (
                "已完成。\n\n"
                f"本次已处理：{workflow_name}。\n"
                f"涉及文件：{changed_text}\n"
                "请刷新页面或重新触发本轮需求对应的功能进行验收。"
            )
        return (
            "这个需求已经被专家团队理解，但当前系统还没有覆盖到对应的自动执行能力。\n\n"
            "我已记录为待补能力项；用户不需要提供代码或自己找文件。等执行能力补齐后，再次发送同类需求即可直接执行。"
        )

    reason_text = "；".join(stringify(item) for item in (reasons or []))
    return (
        "这次回复已被校验层拦截，没有直接展示内部执行信息。\n\n"
        f"原因：{reason_text or '输出不符合用户可读标准'}。"
    )


def response_has_execution_evidence(expert_execution):
    for item in (expert_execution or {}).get("results") or []:
        quality = item.get("quality") or {}
        if quality.get("has_evidence") or quality.get("closed_loop"):
            return True
        data = item.get("data") or {}
        if data.get("changed_files") or data.get("verification"):
            return True
        closure = data.get("closure") or {}
        if closure.get("evidence"):
            return True
    return False


def estimate_answer_alignment_score(original_input, answer):
    source_tokens = set(extract_alignment_tokens(original_input))
    answer_tokens = set(extract_alignment_tokens(answer))
    if not source_tokens:
        return 1.0
    if not answer_tokens:
        return 0.0
    overlap = len(source_tokens & answer_tokens)
    return overlap / max(1, min(len(source_tokens), len(answer_tokens)))


def extract_alignment_tokens(text):
    raw = stringify(text).lower()
    tokens = re.findall(r"[a-z0-9_]{3,}|[\u4e00-\u9fa5]{2,}", raw)
    stop_words = {
        "这个", "我的", "一个", "可以", "需要", "用户", "专家", "团队", "功能",
        "已经", "当前", "回复", "问题", "结果", "系统", "页面", "商品",
    }
    return [token for token in tokens if token not in stop_words][:80]


def fallback_expert_team_answer(message, error):
    return (
        "团队判断：专家团队接口已收到你的问题，但当前 AI 服务暂时不可用，先给你一个本地兜底建议。\n\n"
        "应该调用的专家：首席项目规划专家、专家团队管理者、数据方法论负责人、技术架构专家。\n\n"
        "执行顺序：先明确你的目标和使用场景，再拆分为业务专家任务，最后由技术架构专家判断如何落到系统模块和数据库。\n\n"
        f"当前问题：{message}\n\n"
        f"服务状态：{error}"
    )


def build_permission_required_answer(message, readonly_context):
    return (
        "这一步还没有真正查数据库，原因不是数据库配置缺失，而是当前会话还没有获得数据库访问授权。\n\n"
        "你有两个选择：\n"
        "1. 我把需要执行的只读 SQL 给你，你在数据库工具里执行后把结果发回来，我继续分析。\n"
        "2. 你回复“允许访问数据库”，我就在本轮会话里执行白名单只读诊断，只会跑 SHOW INDEX、SHOW COLUMNS、SHOW CREATE TABLE、EXPLAIN SELECT，不会写入或修改数据。\n\n"
        "针对你这个慢接口问题，建议授权后先查：kalodata 表索引、列表接口真实排序字段、以及对应 EXPLAIN 执行计划。"
    )


def needs_database_runtime_access(message):
    text = stringify(message)
    lower = text.lower()
    db_terms = ["数据库", "查库", "sql", "mysql", "explain", "show index", "show create table", "索引", "慢查询", "执行计划"]
    action_terms = ["检查", "查看", "查一下", "执行", "跑一下", "分析", "为什么", "慢", "优化"]
    return any(term in lower for term in db_terms) and any(term in lower for term in action_terms)


def user_confirmed_database_read_access(message):
    text = stringify(message)
    lower = text.lower()
    if any(phrase.lower() in lower for phrase in DB_READONLY_PERMISSION_PHRASES):
        return True
    short_yes = text.strip() in {"是", "可以", "同意", "允许", "确认", "yes", "ok", "OK"}
    return short_yes and any(term in text for term in ["访问数据库", "查数据库", "查库", "只读SQL", "只读 SQL", "权限"])


def build_database_permission_context(message):
    return (
        "【权限缺口】本轮问题需要访问数据库做只读诊断，但当前对话尚未获得用户的一次性数据库访问授权。\n"
        "已确认：数据库连接配置存在于项目文件中，但“知道配置”不等于“已获准执行数据库查询”。\n"
        "专家团队不能假装已经查库，也不能把推断说成事实。\n"
        "可选路径：\n"
        "1. 用户自己执行专家给出的只读 SQL，把结果贴回；专家继续分析。\n"
        "2. 专家团队向用户请求一次性只读数据库访问授权；用户回复“允许访问数据库”后，系统继续执行 SHOW INDEX、SHOW COLUMNS、SHOW CREATE TABLE、EXPLAIN SELECT 等白名单只读诊断。\n"
    )


def extract_readonly_diagnostic_sqls(message):
    text = stringify(message)
    sqls = []
    for match in re.finditer(r"(?is)\b(EXPLAIN\s+SELECT|SHOW\s+INDEX|SHOW\s+COLUMNS|SHOW\s+CREATE\s+TABLE)\b.*?(?=(?:\n\s*\n)|$)", text):
        sql = match.group(0).strip().rstrip(";")
        if sql and is_allowed_readonly_diagnostic_sql(sql):
            sqls.append(sql)
    return list(dict.fromkeys(sqls))[:3]


def is_allowed_readonly_diagnostic_sql(sql):
    normalized = re.sub(r"\s+", " ", stringify(sql).strip())
    upper = normalized.upper()
    if not any(upper.startswith(prefix) for prefix in DB_READONLY_DIAGNOSTIC_PREFIXES):
        return False
    forbidden = [" INSERT ", " UPDATE ", " DELETE ", " DROP ", " ALTER ", " CREATE INDEX ", " TRUNCATE ", " REPLACE ", " GRANT ", " REVOKE "]
    if any(token in f" {upper} " for token in forbidden):
        return False
    return ";" not in normalized


def append_database_permission_or_diagnostics(context_blocks, message, requested_sources):
    if not needs_database_runtime_access(message):
        return False
    if not user_confirmed_database_read_access(message):
        context_blocks.append(build_database_permission_context(message))
        return True

    lines = ["【数据库只读诊断结果】用户已授权本轮执行白名单只读数据库诊断；未写入、未修改、未删除任何数据。"]
    try:
        with db() as conn, conn.cursor() as cursor:
            for source in (requested_sources or ["unified"])[:3]:
                meta = SOURCES.get(source)
                if not meta:
                    continue
                table = meta["table"]
                lines.append(f"\n数据源：{source}，表：{table}")
                cursor.execute(f"SHOW INDEX FROM `{table}`")
                indexes = cursor.fetchall() or []
                if indexes:
                    summary = []
                    for row in indexes[:12]:
                        summary.append(f"{row.get('Key_name')}({row.get('Column_name')})")
                    lines.append("索引概览：" + "，".join(summary))
                else:
                    lines.append("索引概览：未读取到索引")

            for sql in extract_readonly_diagnostic_sqls(message):
                lines.append(f"\n执行SQL：{sql}")
                cursor.execute(sql)
                rows = cursor.fetchall() or []
                lines.append("结果：" + json.dumps(rows[:8], ensure_ascii=False, default=str))
    except Exception as exc:
        lines.append(f"\n数据库只读诊断失败：{exc}")
    context_blocks.append("\n".join(lines))
    return True


def collect_expert_readonly_context(message, base_files=None):
    text = stringify(message)
    lower = text.lower()
    context_blocks = []
    base_context = collect_expert_base_file_context(base_files or [], text)
    if base_context:
        context_blocks.append(base_context)
    file_context = collect_expert_file_context(text)
    if file_context:
        context_blocks.append(file_context)
    trigger_words = [
        "查", "查看", "数据", "字段", "接口", "回执", "商品", "属性", "卖点",
        "attributes", "selling_points", "product_id", "fastmoss", "kalodata", "统一表",
        "fastmoss_product_aggregate", "fastmoss_product_rank_aggregate", "kalodata_youwei_product",
    ]
    if not any(word.lower() in lower for word in trigger_words):
        return "\n\n".join(context_blocks)

    product_ids = list(dict.fromkeys(re.findall(r"\b\d{10,}\b", text)))[:5]
    requested_sources = []
    for source, meta in SOURCES.items():
        table = meta["table"]
        if source.lower() in lower or table.lower() in lower:
            requested_sources.append(source)
    if "统一表" in text and "unified" not in requested_sources:
        requested_sources.append("unified")
    if not requested_sources:
        requested_sources = ["unified"]
    requested_sources = requested_sources[:3]

    permission_or_diagnostic_handled = append_database_permission_or_diagnostics(context_blocks, text, requested_sources)
    if permission_or_diagnostic_handled and not user_confirmed_database_read_access(text):
        return "\n\n".join(context_blocks)

    context_lines = ["【只读工具结果】本轮仅执行白名单只读查询，未写入、未修改、未删除任何数据。"]
    try:
        with db() as conn, conn.cursor() as cursor:
            for source in requested_sources:
                meta = SOURCES[source]
                table = meta["table"]
                columns = fetch_table_columns(cursor, table)
                context_lines.append(f"\n数据源：{source}，表：{table}")
                if columns:
                    wanted = [
                        meta.get("id"),
                        meta.get("date"),
                        meta.get("title"),
                        meta.get("attributes"),
                        meta.get("selling_points"),
                        meta.get("price"),
                        meta.get("sold"),
                        meta.get("total_sold"),
                    ]
                    present = [col for col in wanted if col and col in columns]
                    context_lines.append(f"可用关键字段：{', '.join(present) if present else '未命中关键字段'}")
                    append_column_fill_stats(cursor, context_lines, table, columns, meta)
                    if product_ids:
                        append_product_readonly_samples(cursor, context_lines, table, columns, meta, product_ids)
                else:
                    context_lines.append("字段读取失败或表不存在。")
    except Exception as exc:
        context_lines.append(f"\n只读工具读取失败：{exc}")
    context_blocks.append("\n".join(context_lines))
    return "\n\n".join(context_blocks)


def collect_expert_file_context(message):
    text = stringify(message).strip()
    lower = text.lower()
    trigger_words = [
        "文件", "代码", "前端", "后端", "接口", "页面", "样式", "按钮", "功能", "路由",
        "api", "html", "css", "js", "python", "flask", "java", "spring", "agent", "专家团队",
        "index.html", "styles.css", "app.py", "prompt", "数据库", "字段", "为什么", "怎么改",
    ]
    if not any(word.lower() in lower for word in trigger_words):
        return ""

    search_terms = extract_file_search_terms(text)
    candidate_files = list(iter_expert_allowed_files())
    scored = []
    for file_path in candidate_files:
        score = score_expert_file(file_path, search_terms, lower)
        if score > 0:
            scored.append((score, file_path))
    scored.sort(key=lambda item: (-item[0], str(item[1]).lower()))
    selected = [file_path for _score, file_path in scored[:EXPERT_FILE_CONTEXT_MAX_FILES]]
    if not selected:
        selected = default_expert_context_files()

    lines = [
        "【项目文件侦察结果】本轮仅执行项目目录白名单只读读取，未写入、未修改、未删除任何文件。",
        f"项目根目录：{PROJECT_ROOT}",
        "侦察员能力边界：只能读取 app.py、static、tools、agent-center 和项目文档等白名单文本文件。",
    ]
    for file_path in selected:
        snippet = build_expert_file_snippet(file_path, search_terms)
        if snippet:
            lines.append(snippet)
        if sum(len(line) for line in lines) > EXPERT_FILE_CONTEXT_MAX_CHARS:
            lines.append("\n文件上下文已达到长度上限，后续文件省略。")
            break
    return "\n".join(lines)[:EXPERT_FILE_CONTEXT_MAX_CHARS]


def normalize_expert_base_files(value):
    if isinstance(value, str):
        raw_items = re.split(r"[\r\n,]+", value)
    elif isinstance(value, list):
        raw_items = value
    else:
        raw_items = []
    normalized = []
    for item in raw_items:
        text = stringify(item).strip().strip('"').strip("'")
        if not text:
            continue
        normalized.append(text)
    return list(dict.fromkeys(normalized))[:12]


def merge_expert_base_files(*groups):
    merged = []
    for group in groups:
        for item in group or []:
            text = stringify(item).strip()
            if text and text not in merged:
                merged.append(text)
    return merged[:12]


def extract_expert_file_paths(text):
    raw = stringify(text)
    patterns = [
        r"[A-Za-z]:\\(?:[A-Za-z0-9_.() -]+\\)*[A-Za-z0-9_.() -]+\.[A-Za-z0-9]+",
        r"(?:app\.py|static/[A-Za-z0-9_./-]+|tools/[A-Za-z0-9_./-]+|agent-center/[A-Za-z0-9_./-]+|[A-Za-z0-9_-]+\.md)",
    ]
    paths = []
    for pattern in patterns:
        for match in re.findall(pattern, raw):
            cleaned = stringify(match).strip().rstrip("，。；;、)")
            if cleaned and cleaned not in paths:
                paths.append(cleaned)
    return paths[:12]


def collect_expert_base_file_context(base_files, message=""):
    if not base_files:
        return ""
    search_terms = extract_file_search_terms(message)
    lines = [
        "【会话基底文件读取结果】用户为本次专家会话显式指定了以下基底文件。系统已在后端按白名单只读读取；这些内容可作为判断依据。",
        f"项目根目录：{PROJECT_ROOT}",
        "安全边界：只读、仅允许 D:\\choice_product 内文本文件、不会写入或执行。",
        "读取策略：已对指定文件做全文件关键词检索，并返回命中片段；不要再要求用户粘贴这些文件里的代码。",
        f"本轮检索关键词：{', '.join(search_terms) if search_terms else '无，使用文件开头预览'}",
    ]
    for raw_path in base_files:
        resolved, error = resolve_expert_base_file(raw_path)
        if error:
            lines.append(f"\n指定路径：{raw_path}\n  读取状态：失败，原因：{error}")
            continue
        snippet = build_expert_file_targeted_snippet(resolved, search_terms)
        lines.append(snippet)
        if sum(len(line) for line in lines) > EXPERT_FILE_CONTEXT_MAX_CHARS:
            lines.append("\n基底文件上下文已达到长度上限，后续文件省略。")
            break
    return "\n".join(lines)[:EXPERT_FILE_CONTEXT_MAX_CHARS]


def resolve_expert_base_file(raw_path):
    if not raw_path:
        return None, "路径为空"
    text = stringify(raw_path).strip()
    candidate = Path(text)
    if not candidate.is_absolute():
        candidate = PROJECT_ROOT / candidate
    try:
        resolved = candidate.resolve(strict=False)
        resolved.relative_to(PROJECT_ROOT)
    except ValueError:
        return None, "不在项目白名单根目录 D:\\choice_product 内"
    except OSError as exc:
        return None, f"路径解析失败：{exc}"
    if not resolved.exists():
        return None, "文件不存在"
    if not resolved.is_file():
        return None, "不是文件"
    if any(part in EXPERT_FILE_EXCLUDED_DIRS for part in resolved.relative_to(PROJECT_ROOT).parts):
        return None, "位于禁止读取目录"
    if resolved.suffix.lower() not in EXPERT_FILE_ALLOWED_SUFFIXES:
        return None, f"不支持的文件类型：{resolved.suffix}"
    try:
        if resolved.stat().st_size > 350_000:
            return None, "文件过大，超过只读上下文限制"
    except OSError as exc:
        return None, f"无法读取文件大小：{exc}"
    return resolved, None


def build_expert_file_targeted_snippet(file_path, search_terms):
    try:
        content = file_path.read_text(encoding="utf-8", errors="ignore")
    except OSError as exc:
        return f"\n文件：{safe_relpath(file_path)}\n  读取状态：失败，原因：{exc}"
    file_lines = content.splitlines()
    snippet_lines = [
        f"\n文件：{safe_relpath(file_path)}",
        "  读取状态：成功",
        f"  文件总行数：{len(file_lines)}",
    ]
    lowered_terms = [term.lower() for term in search_terms if term]
    matched_indexes = []
    if lowered_terms:
        for index, line in enumerate(file_lines):
            line_lower = line.lower()
            score = expert_line_match_score(line_lower, lowered_terms)
            if score > 0:
                matched_indexes.append((score, index))
    if matched_indexes:
        ranked_matches = [index for _score, index in sorted(matched_indexes, key=lambda item: (-item[0], item[1]))[:18]]
        snippet_lines.append(f"  全文件高相关命中位置：{', '.join('L' + str(index + 1) for index in ranked_matches)}")
        selected_indexes = []
        seen = set()
        for index in ranked_matches:
            for nearby in range(max(0, index - 3), min(len(file_lines), index + 4)):
                if nearby not in seen:
                    selected_indexes.append(nearby)
                    seen.add(nearby)
        selected_indexes.sort()
        last_index = None
        for index in selected_indexes[:120]:
            if last_index is not None and index - last_index > 1:
                snippet_lines.append("  ...")
            line = file_lines[index].rstrip()
            if line.strip():
                snippet_lines.append(f"  L{index + 1}: {line[:240]}")
            last_index = index
        if len(selected_indexes) > 120:
            snippet_lines.append("  ... 命中片段较多，已截断。")
    else:
        snippet_lines.append("  全文件检索未命中关键词，展示文件开头预览。")
        max_lines = 80
        for index, line in enumerate(file_lines[:max_lines], start=1):
            if line.strip():
                snippet_lines.append(f"  L{index}: {line.rstrip()[:220]}")
        if len(file_lines) > max_lines:
            snippet_lines.append(f"  ... 文件共 {len(file_lines)} 行，仅展示前 {max_lines} 行。")
    return "\n".join(snippet_lines)


def expert_line_match_score(line_lower, lowered_terms):
    score = 0
    high_value_terms = [
        '@app.get("/api/products/<source>/<product_id>")',
        "def product_detail",
        "function renderproductdetail",
        "renderproductdetail",
        "function opendetail",
        "opendetail(",
        "currentdetailitem",
        "data-detail",
        "products/<source>/<product_id>",
        "analysis-detail",
    ]
    medium_value_terms = ["api/products", "商品详情", "站内详情", "detaildialog", "detailbody"]
    for term in lowered_terms:
        if term and term in line_lower:
            score += 1
    for term in high_value_terms:
        if term in line_lower:
            score += 30
    for term in medium_value_terms:
        if term in line_lower:
            score += 10
    return score


def extract_file_search_terms(message):
    terms = []
    raw_terms = re.findall(r"[A-Za-z_][A-Za-z0-9_./-]{2,}|[\u4e00-\u9fa5]{2,}", stringify(message))
    stop_terms = {
        "这个", "我的", "为什么", "怎么", "一下", "可以", "需要", "功能", "问题", "还是",
        "专家", "团队", "用户", "查看", "文件", "代码", "前端", "后端",
    }
    for term in raw_terms:
        normalized = term.strip()
        if not normalized or normalized in stop_terms:
            continue
        if len(normalized) > 40:
            continue
        terms.append(normalized)
    defaults = ["expert-team", "专家团队", "api/expert-team", "renderExpertTeamPage", "app.py", "index.html", "styles.css"]
    for term in defaults:
        if term.lower() in stringify(message).lower() and term not in terms:
            terms.append(term)
    lower_message = stringify(message).lower()
    mapped_terms = []
    if "商品详情" in stringify(message) or "站内详情" in stringify(message) or "detail" in lower_message:
        mapped_terms.extend([
            "api/products",
            "products/<source>/<product_id>",
            '@app.get("/api/products/<source>/<product_id>")',
            "def product_detail",
            "currentDetailItem",
            "openDetail(",
            "function openDetail",
            "openProductDetail",
            "function renderProductDetail",
            "renderProductDetail",
            "data-detail",
            "detail",
            "modal",
            "商品详情",
        ])
    if "接口" in stringify(message) or "api" in lower_message:
        mapped_terms.extend(["@app.get", "@app.post", "fetch(", "/api/"])
    if "专家团队" in stringify(message):
        mapped_terms.extend(["expert-team", "renderExpertTeamPage", "sendExpertTeamMessage", "collect_expert_readonly_context"])
    if "图片" in stringify(message) or "粘贴" in stringify(message):
        mapped_terms.extend(["expertImages", "expertImageInput", "paste", "image", "images"])
    for term in mapped_terms:
        if term not in terms:
            terms.append(term)
    return list(dict.fromkeys(terms))[:24]


def iter_expert_allowed_files():
    for root, dirs, files in os.walk(PROJECT_ROOT):
        root_path = Path(root)
        dirs[:] = [item for item in dirs if item not in EXPERT_FILE_EXCLUDED_DIRS]
        try:
            rel_root = root_path.relative_to(PROJECT_ROOT)
        except ValueError:
            continue
        if rel_root.parts and rel_root.parts[0] in EXPERT_FILE_EXCLUDED_DIRS:
            continue
        for filename in files:
            file_path = root_path / filename
            if file_path.suffix.lower() not in EXPERT_FILE_ALLOWED_SUFFIXES:
                continue
            try:
                if file_path.stat().st_size > 350_000:
                    continue
                yield file_path
            except OSError:
                continue


def score_expert_file(file_path, search_terms, lower_message):
    rel = file_path.relative_to(PROJECT_ROOT).as_posix().lower()
    score = 0
    important_files = {
        "app.py": 8,
        "static/index.html": 8,
        "static/styles.css": 6,
        "agent-center/src/main/resources/static/index.html": 4,
        "agent-center/src/main/java/com/choiceproduct/agentcenter/agent/totalconversationagent.java": 4,
    }
    score += important_files.get(rel, 0)
    for term in search_terms:
        term_lower = term.lower()
        if term_lower in rel:
            score += 8
    if any(word in lower_message for word in ["前端", "页面", "按钮", "粘贴", "图片", "样式"]):
        if rel in {"static/index.html", "static/styles.css"}:
            score += 12
    if any(word in lower_message for word in ["后端", "接口", "api", "数据库"]):
        if rel == "app.py":
            score += 12
    if any(word in lower_message for word in ["java", "agent", "中台", "spring"]):
        if rel.startswith("agent-center/"):
            score += 10
    return score


def default_expert_context_files():
    defaults = [
        PROJECT_ROOT / "app.py",
        PROJECT_ROOT / "static" / "index.html",
        PROJECT_ROOT / "static" / "styles.css",
        PROJECT_ROOT / "product_selection_platform.md",
        PROJECT_ROOT / "private_product_library_plan.md",
        PROJECT_ROOT / "agent-center" / "README.md",
    ]
    return [file_path for file_path in defaults if file_path.exists()]


def build_expert_file_snippet(file_path, search_terms):
    try:
        content = file_path.read_text(encoding="utf-8", errors="ignore")
    except OSError as exc:
        return f"\n文件：{safe_relpath(file_path)}\n读取失败：{exc}"
    rel = safe_relpath(file_path)
    file_lines = content.splitlines()
    matches = []
    lowered_terms = [term.lower() for term in search_terms if term]
    for index, line in enumerate(file_lines, start=1):
        line_lower = line.lower()
        if not lowered_terms or any(term in line_lower for term in lowered_terms):
            if line.strip():
                matches.append((index, line.rstrip()))
        if len(matches) >= EXPERT_FILE_CONTEXT_MAX_MATCHES_PER_FILE:
            break
    if not matches:
        preview = [(idx + 1, line.rstrip()) for idx, line in enumerate(file_lines[:8]) if line.strip()]
        matches = preview[:EXPERT_FILE_CONTEXT_MAX_MATCHES_PER_FILE]
    snippet_lines = [f"\n文件：{rel}"]
    for line_no, line in matches:
        clipped = line[:220]
        snippet_lines.append(f"  L{line_no}: {clipped}")
    return "\n".join(snippet_lines)


def safe_relpath(file_path):
    try:
        return file_path.relative_to(PROJECT_ROOT).as_posix()
    except ValueError:
        return str(file_path)


def fetch_table_columns(cursor, table):
    allowed_tables = {meta["table"] for meta in SOURCES.values()}
    if table not in allowed_tables:
        return set()
    cursor.execute(f"SHOW COLUMNS FROM `{table}`")
    return {row.get("Field") for row in cursor.fetchall() if row.get("Field")}


def append_column_fill_stats(cursor, context_lines, table, columns, meta):
    stats_cols = [col for col in [meta.get("attributes"), meta.get("selling_points")] if col and col in columns]
    if not stats_cols:
        return
    select_parts = ["COUNT(*) AS total_rows"]
    for col in stats_cols:
        alias = f"{safe_alias(col)}_filled"
        select_parts.append(f"SUM(CASE WHEN `{col}` IS NOT NULL AND TRIM(CAST(`{col}` AS CHAR)) <> '' THEN 1 ELSE 0 END) AS `{alias}`")
    cursor.execute(f"SELECT {', '.join(select_parts)} FROM `{table}`")
    row = cursor.fetchone() or {}
    parts = [f"总行数 {row.get('total_rows', 0)}"]
    for col in stats_cols:
        parts.append(f"{col}有值 {row.get(safe_alias(col) + '_filled', 0)}")
    context_lines.append("字段填充概览：" + "，".join(parts))


def append_product_readonly_samples(cursor, context_lines, table, columns, meta, product_ids):
    product_col = meta.get("id")
    date_col = meta.get("date")
    if not product_col or product_col not in columns:
        return
    wanted = [
        product_col,
        date_col,
        meta.get("title"),
        meta.get("attributes"),
        meta.get("selling_points"),
        meta.get("price"),
        meta.get("sold"),
        meta.get("total_sold"),
    ]
    selected = [col for col in wanted if col and col in columns]
    if not selected:
        return
    placeholders = ", ".join(["%s"] * len(product_ids))
    select_sql = ", ".join(f"`{col}`" for col in selected)
    order_sql = f"ORDER BY `{date_col}` DESC" if date_col and date_col in columns else ""
    cursor.execute(
        f"""
        SELECT {select_sql}
        FROM `{table}`
        WHERE `{product_col}` IN ({placeholders})
        {order_sql}
        LIMIT 5
        """,
        product_ids,
    )
    rows = cursor.fetchall()
    if not rows:
        context_lines.append(f"按商品ID未查到样本：{', '.join(product_ids)}")
        return
    context_lines.append("商品样本：")
    for row in rows:
        summary = []
        for col in selected:
            summary.append(f"{col}={truncate_text(row.get(col), 120)}")
        context_lines.append("- " + "；".join(summary))


def safe_alias(value):
    return re.sub(r"\W+", "_", stringify(value), flags=re.UNICODE).strip("_") or "field"


def truncate_text(value, limit=120):
    text = stringify(value).replace("\r", " ").replace("\n", " ")
    text = re.sub(r"\s+", " ", text).strip()
    return text[:limit] + ("..." if len(text) > limit else "")


def fetch_detail_analysis_record(cursor, user_id, source, product_id, date_record, analysis_type):
    cursor.execute(
        f"""
        SELECT *
        FROM {DETAIL_ANALYSIS_TABLE}
        WHERE user_id = %s
          AND source = %s
          AND product_id = %s
          AND date_record = %s
          AND analysis_type = %s
        LIMIT 1
        """,
        (user_id, source, str(product_id), date_record, analysis_type),
    )
    return cursor.fetchone()


def upsert_detail_analysis_record(cursor, user, source, product_table, product_id, date_record, analysis_type, product_title=""):
    cursor.execute(
        f"""
        INSERT INTO {DETAIL_ANALYSIS_TABLE}
          (user_id, username, source, product_table, product_id, date_record, analysis_type, product_title, status, result_json, error_message, completed_at)
        VALUES (%s, %s, %s, %s, %s, %s, %s, %s, 'RUNNING', NULL, NULL, NULL)
        ON DUPLICATE KEY UPDATE
          username = VALUES(username),
          product_table = VALUES(product_table),
          product_title = IF(VALUES(product_title) <> '', VALUES(product_title), product_title),
          status = 'RUNNING',
          result_json = NULL,
          error_message = NULL,
          completed_at = NULL
        """,
        (
            user["id"],
            user.get("username") or "",
            source,
            product_table,
            str(product_id),
            date_record,
            analysis_type,
            stringify(product_title),
        ),
    )
    cursor.execute(
        f"""
        SELECT id
        FROM {DETAIL_ANALYSIS_TABLE}
        WHERE user_id = %s AND source = %s AND product_id = %s AND date_record = %s AND analysis_type = %s
        LIMIT 1
        """,
        (user["id"], source, str(product_id), date_record, analysis_type),
    )
    return cursor.fetchone()["id"]


def run_detail_analysis_task(record_id, user_id, ai_source, product_id, date_record, analysis_type):
    try:
        with db() as conn, conn.cursor() as cursor:
            product = load_ai_product(cursor, ai_source, product_id, date_record)
            result = run_analysis(
                cursor,
                ai_source,
                product,
                "IP" if analysis_type == "ip" else "MATERIAL",
                write=False,
                detail=True,
            )
            if result.get("error"):
                cursor.execute(
                    f"""
                    UPDATE {DETAIL_ANALYSIS_TABLE}
                    SET status = 'FAILED', error_message = %s, completed_at = NOW()
                    WHERE id = %s AND user_id = %s
                    """,
                    (result["error"], record_id, user_id),
                )
            else:
                cursor.execute(
                    f"""
                    UPDATE {DETAIL_ANALYSIS_TABLE}
                    SET status = 'SUCCESS', result_json = %s, error_message = NULL, completed_at = NOW()
                    WHERE id = %s AND user_id = %s
                    """,
                    (json.dumps(result.get("result") or {}, ensure_ascii=False, default=str), record_id, user_id),
                )
            conn.commit()
    except Exception as exc:
        with db() as conn, conn.cursor() as cursor:
            cursor.execute(
                f"""
                UPDATE {DETAIL_ANALYSIS_TABLE}
                SET status = 'FAILED', error_message = %s, completed_at = NOW()
                WHERE id = %s AND user_id = %s
                """,
                (str(exc), record_id, user_id),
            )
            conn.commit()


def normalize_detail_analysis_record(row):
    return {
        "id": row.get("id"),
        "user_id": row.get("user_id"),
        "username": row.get("username"),
        "source": row.get("source"),
        "product_table": row.get("product_table"),
        "product_id": row.get("product_id"),
        "date_record": stringify(row.get("date_record")),
        "analysis_type": row.get("analysis_type"),
        "product_title": row.get("product_title") or "",
        "status": row.get("status"),
        "result": parse_json(row.get("result_json")),
        "error_message": row.get("error_message") or "",
        "created_at": stringify(row.get("created_at")),
        "updated_at": stringify(row.get("updated_at")),
        "completed_at": stringify(row.get("completed_at")),
    }


def public_user(user):
    role = stringify(user.get("role") or "user").strip().lower()
    return {
        "id": user.get("id"),
        "username": user.get("username"),
        "display_name": user.get("display_name") or user.get("username"),
        "role": role or "user",
    }


def normalize_source(source):
    source = source or "unified"
    if source not in SOURCES:
        raise ValueError(f"未知数据源：{source}")
    return source


def normalize_period(period):
    period = period or "30d"
    return period if period in {"7d", "30d", "90d", "180d", "custom"} else "7d"


def build_sales_period(period, sales_start=None, sales_end=None):
    period = normalize_period(period)
    if period == "custom" and sales_start and sales_end:
        return {
            "period": "custom",
            "label": f"{sales_start} 至 {sales_end}",
            "days": None,
            "start": sales_start,
            "end": sales_end,
        }
    days = {"7d": 7, "30d": 30, "90d": 90, "180d": 180}.get(period, 7)
    return {
        "period": period,
        "label": f"近{days}天",
        "days": days,
        "start": None,
        "end": None,
    }


def clamp_int(value, default, minimum, maximum):
    try:
        value = int(value)
    except (TypeError, ValueError):
        value = default
    return max(minimum, min(maximum, value))


def sql_alias(field, alias):
    if field:
        return f"`{field}` AS {alias}"
    return f"NULL AS {alias}"


def build_product_select(source, meta):
    image_url_expr = (
        f"CONCAT('/api/products/{source}/', `{meta['id']}`, '/image?date_record=', `{meta['date']}`) AS image_url"
    )
    platform_url_expr = (
        f"`{meta['detail_url']}` AS platform_url"
        if meta["detail_url"]
        else "NULL AS platform_url"
    )
    return ", ".join([
        f"`{meta['id']}` AS product_id",
        f"`{meta['date']}` AS date_record",
        sql_alias(meta["title"], "title"),
        sql_alias(meta["price"], "price_view"),
        sql_alias(meta["rating"], "rating_view"),
        sql_alias(meta["commission"], "commission_rate_view"),
        sql_alias(meta["sold"], "sold_count_view"),
        sql_alias(meta["total_sold"], "total_sold_count_view"),
        sql_alias(meta["sale_amount"], "sale_amount_view"),
        sql_alias(meta["author_count"], "author_count_view"),
        sql_alias(meta["base_price"], "base_price_view"),
        sql_alias(meta["transport_fee"], "transport_fee_view"),
        sql_alias(meta["video_ratio"], "video_ratio_view"),
        sql_alias(meta["product_card_ratio"], "product_card_ratio_view"),
        sql_alias(meta["distribution_30d"], "distribution_30d"),
        sql_alias(meta["distribution_7d"], "distribution_7d"),
        sql_alias(meta["distribution_90d"], "distribution_90d"),
        sql_alias(meta["distribution_180d"], "distribution_180d"),
        sql_alias(meta["overview_30d"], "overview_30d"),
        sql_alias(meta["overview_7d"], "overview_7d"),
        sql_alias(meta["overview_90d"], "overview_90d"),
        sql_alias(meta["overview_180d"], "overview_180d"),
        "`audit_status`",
        "`ip_grade`",
        "`ip_reason`",
        "`ip_tags`",
        "`material_analysis`",
        "`ai_analysis_error`",
        "`launch_time`" if has_launch_time(meta) else "NULL AS launch_time",
        sql_alias(meta.get("score"), "pod_score"),
        sql_alias(meta.get("score_reason"), "score_reason"),
        sql_alias(meta.get("sales_mom"), "pod_sales_mom"),
        sql_alias(meta.get("status"), "pod_status"),
        sql_alias(meta.get("illustration_result_url"), "illustration_result_url"),
        sql_alias(meta.get("illustration_extractable"), "illustration_extractable"),
        sql_alias(meta.get("illustration_extract_reason"), "illustration_extract_reason"),
        sql_alias(meta.get("category_l1"), "category_l1"),
        sql_alias(meta.get("category_l2"), "category_l2"),
        sql_alias(meta.get("category_l3"), "category_l3"),
        sql_alias(meta.get("aweme_count"), "aweme_count_view"),
        platform_url_expr,
        image_url_expr,
    ])


def build_latest_product_join(meta, where):
    return f"""
        INNER JOIN (
            SELECT `{meta["id"]}` AS latest_product_id, MAX(`{meta["date"]}`) AS latest_date
            FROM `{meta["table"]}`
            WHERE {" AND ".join(where)}
            GROUP BY `{meta["id"]}`
        ) AS latest_snapshot
          ON latest_snapshot.latest_product_id = `{meta["table"]}`.`{meta["id"]}`
         AND latest_snapshot.latest_date = `{meta["table"]}`.`{meta["date"]}`
    """


def build_sales_aggregate_join(meta, where):
    sold_field = meta.get("sold")
    if not sold_field:
        return ""
    sold_expr = numeric_sql_expr(f"`{sold_field}`")
    return f"""
        LEFT JOIN (
            SELECT `{meta["id"]}` AS aggregate_product_id,
                   SUM(COALESCE({sold_expr}, 0)) AS aggregate_sold_count
            FROM `{meta["table"]}`
            WHERE {" AND ".join(where)}
            GROUP BY `{meta["id"]}`
        ) AS sales_aggregate
          ON sales_aggregate.aggregate_product_id = `{meta["table"]}`.`{meta["id"]}`
    """


def build_previous_sales_aggregate_join(meta, date_start, date_end):
    sold_field = meta.get("sold")
    if not sold_field or not date_start:
        return """
        LEFT JOIN (
            SELECT NULL AS previous_product_id, NULL AS previous_sold_count
        ) AS previous_sales_aggregate
          ON previous_sales_aggregate.previous_product_id = `{table}`.`{id_field}`
        """.format(table=meta["table"], id_field=meta["id"])
    try:
        start = datetime.fromisoformat(str(date_start)[:10]).date()
        end = datetime.fromisoformat(str(date_end or date_start)[:10]).date()
    except ValueError:
        return """
        LEFT JOIN (
            SELECT NULL AS previous_product_id, NULL AS previous_sold_count
        ) AS previous_sales_aggregate
          ON previous_sales_aggregate.previous_product_id = `{table}`.`{id_field}`
        """.format(table=meta["table"], id_field=meta["id"])
    days = max(1, (end - start).days + 1)
    previous_start = start.isoformat()
    sold_expr = numeric_sql_expr(f"`{sold_field}`")
    return f"""
        LEFT JOIN (
            SELECT `{meta["id"]}` AS previous_product_id,
                   SUM(COALESCE({sold_expr}, 0)) AS previous_sold_count
            FROM `{meta["table"]}`
            WHERE `{meta["date"]}` BETWEEN DATE_SUB(CAST('{previous_start}' AS DATE), INTERVAL {days} DAY)
                                      AND DATE_SUB(CAST('{previous_start}' AS DATE), INTERVAL 1 DAY)
            GROUP BY `{meta["id"]}`
        ) AS previous_sales_aggregate
          ON previous_sales_aggregate.previous_product_id = `{meta["table"]}`.`{meta["id"]}`
    """


def build_sales_delta_join(meta, date_start, date_end):
    total_field = meta.get("total_sold") or meta.get("sold")
    if not total_field or not date_start:
        return empty_sales_delta_join(meta)
    try:
        start = datetime.fromisoformat(str(date_start)[:10]).date()
        end = datetime.fromisoformat(str(date_end or date_start)[:10]).date()
    except ValueError:
        return empty_sales_delta_join(meta)
    days = max(1, (end - start).days + 1)
    previous_start = start - date.resolution * days
    return f"""
        LEFT JOIN (
            SELECT base_source.`{meta["id"]}` AS base_product_id,
                   {numeric_sql_expr(f"base_source.`{total_field}`")} AS base_total_sold_count
            FROM `{meta["table"]}` AS base_source
            INNER JOIN (
                SELECT `{meta["id"]}` AS product_id, MAX(`{meta["date"]}`) AS base_date
                FROM `{meta["table"]}`
                WHERE `{meta["date"]}` < CAST('{start.isoformat()}' AS DATE)
                GROUP BY `{meta["id"]}`
            ) AS base_latest
              ON base_latest.product_id = base_source.`{meta["id"]}`
             AND base_latest.base_date = base_source.`{meta["date"]}`
        ) AS current_sales_base
          ON current_sales_base.base_product_id = `{meta["table"]}`.`{meta["id"]}`
        LEFT JOIN (
            SELECT previous_base_source.`{meta["id"]}` AS previous_base_product_id,
                   {numeric_sql_expr(f"previous_base_source.`{total_field}`")} AS previous_base_total_sold_count
            FROM `{meta["table"]}` AS previous_base_source
            INNER JOIN (
                SELECT `{meta["id"]}` AS product_id, MAX(`{meta["date"]}`) AS previous_base_date
                FROM `{meta["table"]}`
                WHERE `{meta["date"]}` < CAST('{previous_start.isoformat()}' AS DATE)
                GROUP BY `{meta["id"]}`
            ) AS previous_base_latest
              ON previous_base_latest.product_id = previous_base_source.`{meta["id"]}`
             AND previous_base_latest.previous_base_date = previous_base_source.`{meta["date"]}`
        ) AS previous_sales_base
          ON previous_sales_base.previous_base_product_id = `{meta["table"]}`.`{meta["id"]}`
    """


def empty_sales_delta_join(meta):
    return """
        LEFT JOIN (
            SELECT NULL AS base_product_id, NULL AS base_total_sold_count
        ) AS current_sales_base
          ON current_sales_base.base_product_id = `{table}`.`{id_field}`
        LEFT JOIN (
            SELECT NULL AS previous_base_product_id, NULL AS previous_base_total_sold_count
        ) AS previous_sales_base
          ON previous_sales_base.previous_base_product_id = `{table}`.`{id_field}`
    """.format(table=meta["table"], id_field=meta["id"])


def sales_delta_current_expr(meta):
    total_field = meta.get("total_sold") or meta.get("sold")
    if not total_field:
        return "NULL"
    current_total_expr = numeric_sql_expr(f"`{meta['table']}`.`{total_field}`")
    return f"GREATEST(COALESCE({current_total_expr}, 0) - COALESCE(current_sales_base.base_total_sold_count, 0), 0)"


def sales_delta_previous_expr():
    return "GREATEST(COALESCE(current_sales_base.base_total_sold_count, 0) - COALESCE(previous_sales_base.previous_base_total_sold_count, 0), 0)"


def has_launch_time(meta):
    return meta["table"] in {"fastmoss_product_aggregate", "fastmoss_product_rank_aggregate", "pod_cross_category_product"}


def build_filters(meta, args, latest_by_default=False):
    where = ["1=1"]
    params = []
    if meta.get("table") == "pod_cross_category_product":
        where.append("illustration_extractable = 1")
    date_start = args.get("date_start")
    date_end = args.get("date_end")
    if latest_by_default and not date_start and not date_end:
        where.append(f"`{meta['date']}` = (SELECT MAX(latest_source.`{meta['date']}`) FROM `{meta['table']}` AS latest_source)")
    else:
        add_range_filter(where, params, meta["date"], date_start, date_end)
    add_eq_filter(where, params, "audit_status", args.get("audit_status"))
    add_grade_filter(where, params, args.get("ip_grade_min"), args.get("ip_grade_max"))
    add_material_filter(where, params, args.get("material_type"), args.get("material_categories"), args.get("exclude_material_categories"))
    add_number_range_filter(where, params, meta.get("commission"), args.get("commission_min"), args.get("commission_max"))
    add_number_range_filter(where, params, meta.get("author_count"), args.get("author_min"), args.get("author_max"))
    add_number_range_filter(where, params, meta.get("sold"), args.get("sold_min"), args.get("sold_max"))
    add_number_range_filter(where, params, meta.get("sale_amount"), args.get("sale_amount_min"), args.get("sale_amount_max"))
    add_number_range_filter(where, params, meta.get("rating"), args.get("rating_min"), args.get("rating_max"))
    add_number_range_filter(where, params, meta.get("price"), args.get("price_min"), args.get("price_max"))
    add_number_range_filter(where, params, meta.get("base_price"), args.get("base_price_min"), args.get("base_price_max"))
    add_number_range_filter(where, params, meta.get("transport_fee"), args.get("transport_fee_min"), args.get("transport_fee_max"))
    add_number_range_filter(where, params, meta.get("score"), args.get("score_min"), args.get("score_max"))
    add_number_range_filter(where, params, meta.get("sales_mom"), args.get("sales_mom_min"), args.get("sales_mom_max"))
    add_number_range_filter(where, params, meta.get("aweme_count"), args.get("aweme_min"), args.get("aweme_max"))
    if meta.get("table") == "pod_cross_category_product":
        add_eq_filter(where, params, meta.get("status"), args.get("pod_status"))
    add_eq_filter(where, params, meta.get("category_l1"), args.get("category_l1"))
    add_eq_filter(where, params, meta.get("category_l2"), args.get("category_l2"))
    add_eq_filter(where, params, meta.get("category_l3"), args.get("category_l3"))
    add_ratio_range_filter(where, params, meta, args, "视频", args.get("video_ratio_min"), args.get("video_ratio_max"), direct_field=meta.get("video_ratio"))
    add_ratio_range_filter(where, params, meta, args, "商品卡", args.get("product_card_ratio_min"), args.get("product_card_ratio_max"), direct_field=meta.get("product_card_ratio"))
    add_range_filter(where, params, "launch_time", args.get("launch_start"), args.get("launch_end")) if has_launch_time(meta) else None
    keyword = (args.get("keyword") or "").strip()
    if keyword:
        where.append(f"`{meta['title']}` LIKE %s")
        params.append(f"%{keyword}%")
    return where, params


def add_range_filter(where, params, field, start, end):
    if start:
        where.append(f"`{field}` >= %s")
        params.append(start)
    if end:
        where.append(f"`{field}` <= %s")
        params.append(end)


def add_number_range_filter(where, params, field, start, end):
    if not field:
        return
    expr = f"CAST(REPLACE(REPLACE(`{field}`, '$', ''), '%%', '') AS DECIMAL(20,4))"
    if start:
        where.append(f"{expr} >= %s")
        params.append(start)
    if end:
        where.append(f"{expr} <= %s")
        params.append(end)


def add_ratio_range_filter(where, params, meta, args, ratio_name, start, end, direct_field=None):
    if not start and not end:
        return
    if direct_field:
        add_number_range_filter(where, params, direct_field, start, end)
        return

    field = get_distribution_filter_field(meta, args)
    if not field:
        return
    expr = distribution_ratio_number_expr(field)
    if start:
        where.append(f"JSON_VALID(`{field}`) AND {expr} >= %s")
        params.extend([ratio_name, start])
    if end:
        where.append(f"JSON_VALID(`{field}`) AND {expr} <= %s")
        params.extend([ratio_name, end])


def get_distribution_filter_field(meta, args):
    period = normalize_period(args.get("period"))
    if period == "custom":
        period = "7d"
    field = meta.get(f"distribution_{period}")
    return field or meta.get("distribution_7d") or meta.get("distribution_30d")


def distribution_ratio_number_expr(field):
    search_path = f"JSON_UNQUOTE(JSON_SEARCH(`{field}`, 'one', %s, NULL, '$[*].name'))"
    value_path = f"REPLACE({search_path}, '.name', '.percentage')"
    value_expr = f"JSON_UNQUOTE(JSON_EXTRACT(`{field}`, {value_path}))"
    clean_expr = f"REPLACE(REPLACE({value_expr}, '%%', ''), ',', '')"
    return f"CAST(NULLIF({clean_expr}, '') AS DECIMAL(20,4))"


def distribution_item_number_expr(field, item_name, key):
    search_path = f"JSON_UNQUOTE(JSON_SEARCH(`{field}`, 'one', '{item_name}', NULL, '$[*].name'))"
    value_path = f"REPLACE({search_path}, '.name', '.{key}')"
    value_expr = f"JSON_UNQUOTE(JSON_EXTRACT(`{field}`, {value_path}))"
    return numeric_sql_expr(value_expr)


def distribution_index_number_expr(field, index, key):
    value_expr = f"JSON_UNQUOTE(JSON_EXTRACT(`{field}`, '$[{index}].{key}'))"
    return numeric_sql_expr(value_expr)


def add_eq_filter(where, params, field, value):
    if field and value:
        where.append(f"`{field}` = %s")
        params.append(value)


def add_grade_filter(where, params, grade_min, grade_max):
    if not grade_min and not grade_max:
        return
    min_index = IP_GRADES.index(grade_min) if grade_min in IP_GRADES else 0
    max_index = IP_GRADES.index(grade_max) if grade_max in IP_GRADES else len(IP_GRADES) - 1
    if min_index > max_index:
        min_index, max_index = max_index, min_index
    grades = IP_GRADES[min_index:max_index + 1]
    where.append("ip_grade IN (" + ",".join(["%s"] * len(grades)) + ")")
    params.extend(grades)


def add_material_filter(where, params, material_type, categories, exclude_categories):
    if material_type:
        where.append("JSON_VALID(material_analysis) AND JSON_UNQUOTE(JSON_EXTRACT(material_analysis, '$.material_type')) = %s")
        params.append(material_type)
    category_list = split_csv(categories)
    if category_list:
        where.append("JSON_VALID(material_analysis) AND JSON_UNQUOTE(JSON_EXTRACT(material_analysis, '$.material_category')) IN (" + ",".join(["%s"] * len(category_list)) + ")")
        params.extend(category_list)
    excluded = split_csv(exclude_categories)
    if excluded:
        where.append("(NOT JSON_VALID(material_analysis) OR JSON_UNQUOTE(JSON_EXTRACT(material_analysis, '$.material_category')) NOT IN (" + ",".join(["%s"] * len(excluded)) + "))")
        params.extend(excluded)


def split_csv(value):
    if not value:
        return []
    return [item.strip() for item in value.split(",") if item.strip()]


def numeric_sql_expr(expr):
    clean_expr = expr
    for token in ("US$", "USD", "$", "¥", "￥", "%%", ",", "，", " "):
        clean_expr = f"REPLACE({clean_expr}, '{token}', '')"
    return (
        "CASE "
        f"WHEN {clean_expr} LIKE '%%万%%' THEN CAST(NULLIF(REPLACE({clean_expr}, '万', ''), '') AS DECIMAL(20,4)) * 10000 "
        f"ELSE CAST(NULLIF({clean_expr}, '') AS DECIMAL(20,4)) "
        "END"
    )


def qualified_field(meta, field):
    if not field:
        return "NULL"
    return f"`{meta['table']}`.`{field}`"


def sort_field_expr(meta, sort_by):
    fields = {
        "date_record": meta["date"],
        "sold": meta.get("sold"),
        "total_sold": meta.get("total_sold"),
        "sale_amount": meta.get("sale_amount"),
        "rating": meta.get("rating"),
        "price": meta.get("price"),
        "commission": meta.get("commission"),
        "author_count": meta.get("author_count"),
        "launch_time": "launch_time" if has_launch_time(meta) else None,
        "score": meta.get("score"),
        "sales_mom": meta.get("sales_mom"),
        "aweme_count": meta.get("aweme_count"),
        "category_l1": meta.get("category_l1"),
        "category_l2": meta.get("category_l2"),
        "category_l3": meta.get("category_l3"),
    }
    return qualified_field(meta, fields.get(sort_by or "date_record") or meta["date"])


def build_order(meta, sort_by, sort_order, sales_period=None, use_sales_aggregate=False, args=None):
    raw_field = sort_field_expr(meta, sort_by)
    direction = "ASC" if str(sort_order).lower() == "asc" else "DESC"
    if sort_by == "sold" and use_sales_aggregate:
        return f"ORDER BY COALESCE(sales_aggregate.aggregate_sold_count, 0) {direction}"
    if sort_by == "product_card_sales":
        return f"ORDER BY {product_card_sales_sort_expr(meta)} {direction}"
    if sort_by == "product_card_ratio":
        return f"ORDER BY {product_card_ratio_sort_expr(meta)} {direction}"
    if sort_by == "sales_growth" and args:
        return f"ORDER BY {sales_growth_sort_expr(meta, args)} {direction}"
    if sort_by == "ip_grade":
        return f"ORDER BY FIELD(ip_grade, 'E', 'D', 'C', 'B', 'A', 'S') {direction}, ip_grade {direction}"

    # 针对数值字段，强制进行数值转换排序，处理可能存在的 $ % , 等符号
    # 防止出现字符串排序导致的 "84 > 658" 错误
    overview_metrics = {"sold": "销量", "sale_amount": "销售额", "author_count": "带货达人数"}
    if sort_by in overview_metrics and sales_period:
        period = sales_period["period"] if sales_period["period"] != "custom" else "7d"
        overview_field = meta.get(f"overview_{period}")
        if not overview_field:
            return f"ORDER BY {numeric_sql_expr(raw_field)} {direction}"
        overview_field_expr = qualified_field(meta, overview_field)
        overview_expr = f"JSON_UNQUOTE(JSON_EXTRACT({overview_field_expr}, '$.\"{overview_metrics[sort_by]}\"'))"
        return f"ORDER BY COALESCE({numeric_sql_expr(overview_expr)}, {numeric_sql_expr(raw_field)}) {direction}"

    numeric_keys = {"sold", "total_sold", "sale_amount", "rating", "price", "commission", "author_count"}
    numeric_keys |= {"score", "sales_mom", "aweme_count"}
    if sort_by in numeric_keys:
        # 移除常见非数字符号并转换为 DECIMAL 排序
        return f"ORDER BY {numeric_sql_expr(raw_field)} {direction}"

    return f"ORDER BY {raw_field} {direction}"


def product_card_sales_sort_expr(meta):
    distribution_field = meta.get("distribution_7d") or meta.get("distribution_30d")
    if distribution_field:
        return f"COALESCE({distribution_item_number_expr(distribution_field, '商品卡', 'sales')}, 0)"
    ratio_field = meta.get("product_card_ratio")
    total_field = meta.get("total_sold") or meta.get("sold")
    if ratio_field and total_field:
        return f"COALESCE({numeric_sql_expr(f'`{total_field}`')}, 0) * COALESCE({numeric_sql_expr(f'`{ratio_field}`')}, 0) / 100"
    return "0"


def product_card_ratio_sort_expr(meta):
    expressions = []
    if meta.get("distribution_30d"):
        expressions.append(distribution_item_number_expr(meta["distribution_30d"], "商品卡", "percentage"))
    if meta.get("product_card_ratio"):
        expressions.append(numeric_sql_expr(f"`{meta['product_card_ratio']}`"))
    if meta.get("distribution_7d"):
        expressions.append(distribution_item_number_expr(meta["distribution_7d"], "商品卡", "percentage"))
    if not expressions:
        return "0"
    return f"COALESCE({', '.join(expressions)}, 0)"


def product_card_sales_sort_expr(meta):
    distribution_field = meta.get("distribution_7d") or meta.get("distribution_30d")
    if distribution_field:
        return "COALESCE({sales}, {count}, {sold}, {value_count}, 0)".format(
            sales=distribution_index_number_expr(distribution_field, 2, "sales"),
            count=distribution_index_number_expr(distribution_field, 2, "count"),
            sold=distribution_index_number_expr(distribution_field, 2, "sold"),
            value_count=distribution_index_number_expr(distribution_field, 2, "value_count"),
        )
    ratio_field = meta.get("product_card_ratio")
    total_field = meta.get("total_sold") or meta.get("sold")
    if ratio_field and total_field:
        return f"COALESCE({numeric_sql_expr(f'`{total_field}`')}, 0) * COALESCE({numeric_sql_expr(f'`{ratio_field}`')}, 0) / 100"
    return "0"


def sales_growth_sort_expr(meta, args):
    current_expr = sales_delta_current_expr(meta)
    previous_expr = sales_delta_previous_expr()
    return f"CASE WHEN COALESCE({previous_expr}, 0) <= 0 THEN NULL ELSE (({current_expr} / {previous_expr}) - 1) END"


def normalize_product_row(source, row, sales_period):
    material = parse_json(row.get("material_analysis")) or {}
    tags = parse_json(row.get("ip_tags")) or {}
    field_period = sales_period["period"] if sales_period["period"] != "custom" else "7d"
    overview = parse_json(row.get(f"overview_{field_period}")) or {}
    distribution = parse_json(row.get(f"distribution_{field_period}")) or []
    distribution_7d = parse_json(row.get("distribution_7d")) or []
    distribution_28d = parse_json(row.get("distribution_30d")) or []
    sold_count = row.get("runtime_sold_count") if row.get("runtime_sold_count") is not None else overview.get("销量") or row.get("sold_count_view")
    total_sold = row.get("total_sold_count_view") or row.get("sold_count_view")
    sale_amount = overview.get("销售额") or row.get("sale_amount_view")
    author_count = overview.get("带货达人数") or row.get("author_count_view")
    video_ratio = row.get("video_ratio_view") or ratio_from_distribution(distribution, "视频")
    product_card_ratio = row.get("product_card_ratio_view") or ratio_from_distribution(distribution, "商品卡")
    product_card_ratio_28d = row.get("product_card_ratio_view") or ratio_from_distribution(distribution_28d, "商品卡")
    product_card_sales = sales_from_distribution(distribution_7d or distribution, "商品卡", sold_count)
    chart_7d = normalize_distribution_chart(distribution_7d, sold_count)
    chart_28d = normalize_distribution_chart(
        distribution_28d,
        total_sold,
        {"视频": row.get("video_ratio_view"), "商品卡": row.get("product_card_ratio_view")},
    )
    return {
        "source": source,
        "product_id": stringify(row.get("product_id")),
        "date_record": row.get("date_record"),
        "title": row.get("title") or "",
        "image_url": row.get("image_url"),
        "price": normalize_price_display(row.get("price_view"), row.get("base_price_view")),
        "base_price": stringify(row.get("base_price_view")),
        "rating": stringify(row.get("rating_view")),
        "commission_rate": stringify(row.get("commission_rate_view")),
        "sold_count": stringify(sold_count),
        "product_card_sales": stringify(format_metric_number(product_card_sales)),
        "total_sold_count": stringify(total_sold),
        "sale_amount": stringify(sale_amount),
        "sales_growth": row.get("sales_growth_view") or "N/A",
        "author_count": stringify(author_count),
        "video_ratio": stringify(video_ratio),
        "product_card_ratio": stringify(product_card_ratio),
        "product_card_ratio_28d": stringify(product_card_ratio_28d),
        "selection_score": stringify(row.get("pod_score")),
        "score_reason": row.get("score_reason"),
        "sales_mom": stringify(row.get("pod_sales_mom")),
        "illustration_status": row.get("pod_status"),
        "illustration_result_url": stringify(row.get("illustration_result_url")),
        "illustration_extractable": row.get("illustration_extractable"),
        "illustration_extract_reason": stringify(row.get("illustration_extract_reason")),
        "pod_score": stringify(row.get("pod_score")),
        "pod_sales_mom": stringify(row.get("pod_sales_mom")),
        "pod_status": row.get("pod_status"),
        "category_l1": stringify(row.get("category_l1")),
        "category_l2": stringify(row.get("category_l2")),
        "category_l3": stringify(row.get("category_l3")),
        "aweme_count": stringify(row.get("aweme_count_view")),
        "distribution_7d_chart": chart_7d,
        "distribution_28d_chart": chart_28d,
        "transport_fee": stringify(row.get("transport_fee_view")),
        "launch_time": stringify(row.get("launch_time")),
        "audit_status": row.get("audit_status") or "PENDING",
        "ip_grade": row.get("ip_grade"),
        "ip_reason": row.get("ip_reason"),
        "ip_tags": tags,
        "material_type": material.get("material_type"),
        "material_category": material.get("material_category"),
        "material_confidence": material.get("confidence"),
        "material_reason": material.get("material_reason"),
        "ai_analysis_error": parse_json(row.get("ai_analysis_error")),
        "platform_url": resolve_platform_url(source, row.get("product_id"), row.get("platform_url")),
    }


def clean_number(value):
    if value is None:
        return 0.0
    if isinstance(value, (int, float, Decimal)):
        return float(value)
    text = stringify(value).strip()
    if not text:
        return 0.0
    multiplier = 1.0
    if "亿" in text:
        multiplier = 100000000.0
    elif "万" in text:
        multiplier = 10000.0
    text = re.sub(r"[^0-9.\-]", "", text)
    if not text or text in {"-", ".", "-."}:
        return 0.0
    try:
        return float(text) * multiplier
    except ValueError:
        return 0.0


def overview_metric_value(row, field, keywords):
    data = parse_json(row.get(field)) if field else None
    if not isinstance(data, dict):
        return 0.0
    for key, value in data.items():
        key_text = stringify(key)
        if "日均" in key_text:
            continue
        if any(keyword in key_text for keyword in keywords):
            return clean_number(value)
    return 0.0


def score_category_key(source, meta, row):
    for field in (meta.get("category_l3"), meta.get("category_l2"), meta.get("category_l1")):
        value = row.get(field) if field else None
        if value:
            return stringify(value)
    return source


def score_input_from_row(source, meta, row):
    sales_7d = overview_metric_value(row, meta.get("overview_7d"), ["销量", "sale"])
    sales_30d = overview_metric_value(row, meta.get("overview_30d"), ["销量", "sale"])
    if sales_30d <= 0:
        sales_30d = clean_number(row.get(meta.get("sold"))) or clean_number(row.get(meta.get("total_sold")))
    if sales_7d <= 0:
        sales_7d = clean_number(row.get(meta.get("sold"))) if sales_30d <= 0 else min(clean_number(row.get(meta.get("sold"))), sales_30d)
    video_count = clean_number(row.get(meta.get("aweme_count"))) or clean_number(row.get("aweme_count"))
    rating = clean_number(row.get(meta.get("rating"))) or 3.5
    price = clean_number(row.get(meta.get("base_price"))) or clean_number(row.get(meta.get("price")))
    total_sales = clean_number(row.get(meta.get("total_sold"))) or clean_number(row.get(meta.get("sold")))
    return {
        "category_key": score_category_key(source, meta, row),
        "sales_7d": sales_7d,
        "sales_30d": sales_30d,
        "total_sales": total_sales,
        "video_count": video_count,
        "rating": rating,
        "price": price,
    }


def penalty_norm_score(norm, max_points):
    if norm <= 0:
        return 0.0
    norm = max(0.0, min(1.0, float(norm)))
    # 用户规则的核心是“中段最优”，且各维度不得超过自身满分。
    # 需求示例给出的惩罚表为：10%=14.6、30%=52、50%=100、70%=80.6、100%=66.8。
    # 用分段插值严格贴合该表，再乘以维度满分，避免 35 分维度算出 46 分这类超分。
    curve = [
        (0.0, 0.0),
        (0.1, 0.146),
        (0.3, 0.52),
        (0.5, 1.0),
        (0.7, 0.806),
        (1.0, 0.668),
    ]
    for (left_norm, left_score), (right_norm, right_score) in zip(curve, curve[1:]):
        if norm <= right_norm:
            span = right_norm - left_norm
            ratio = 0 if span <= 0 else (norm - left_norm) / span
            normalized_score = left_score + (right_score - left_score) * ratio
            return round(max(0.0, min(1.0, normalized_score)) * max_points, 6)
    return round(curve[-1][1] * max_points, 6)


def calculate_selection_score(score_input, benchmark):
    rating = score_input["rating"] or 3.5
    price = score_input["price"]
    sales_30d = score_input["sales_30d"]
    total_sales = score_input["total_sales"]
    detail = {
        "version": "selection_score_v1",
        "inputs": {
            "sales_7d": round(score_input["sales_7d"], 2),
            "sales_30d": round(sales_30d, 2),
            "total_sales": round(total_sales, 2),
            "video_count": round(score_input["video_count"], 2),
            "rating": round(rating, 2),
            "price": round(price, 2) if price else 0,
        },
        "benchmark": {
            "category_key": score_input.get("category_key"),
            "max_sales_7d": round(benchmark.get("max_sales_7d") or 0, 2),
            "max_sales_30d": round(benchmark.get("max_sales_30d") or 0, 2),
            "max_video_count": round(benchmark.get("max_video_count") or 0, 2),
        },
        "components": {},
        "risk": None,
        "conclusion": "",
        "summary": "",
    }

    def finish_risk(message):
        detail["risk"] = message
        detail["conclusion"] = "暂不上架"
        detail["summary"] = message
        return 0.0, json.dumps(detail, ensure_ascii=False)

    if rating < 3.0:
        return finish_risk("评分低于 3.0，触发风控规则，选品分归零。")
    if price and (price < 1 or price > 100):
        return finish_risk("价格低于 $1 或高于 $100，触发风控规则，选品分归零。")
    if sales_30d == 0 and total_sales > 0:
        return finish_risk("30天销量为 0 但总销量大于 0，疑似数据异常，选品分归零。")

    max_30d = max(benchmark.get("max_sales_30d") or 0, sales_30d)
    max_7d = max(benchmark.get("max_sales_7d") or 0, score_input["sales_7d"])
    max_video = max(benchmark.get("max_video_count") or 0, score_input["video_count"])
    norm_30d = (math.log10(sales_30d + 1) / math.log10(max_30d + 1)) if max_30d > 0 else 0
    norm_7d = (math.log10(score_input["sales_7d"] + 1) / math.log10(max_7d + 1)) if max_7d > 0 else 0
    norm_video = (math.log10(score_input["video_count"] + 1) / math.log10(max_video + 1)) if max_video > 0 else 0
    sales_score = penalty_norm_score(norm_30d, 35)
    rating_score = (rating / 5) * 25
    heat_score = penalty_norm_score(norm_video, 25)
    trend_ratio = (norm_7d / norm_30d) if norm_30d > 0 else 0
    trend_score = min(1, trend_ratio) * 15
    total = round(min(100, max(0, sales_score + rating_score + heat_score + trend_score)), 2)
    if total >= 75:
        conclusion = "推荐上架"
    elif total >= 50:
        conclusion = "观察等待"
    else:
        conclusion = "暂不上架"
    reason = (
        f"{conclusion}：销量分 {sales_score:.1f}，评分分 {rating_score:.1f}，"
        f"热度分 {heat_score:.1f}，趋势分 {trend_score:.1f}。"
    )
    detail["components"] = {
        "sales_score": round(sales_score, 2),
        "rating_score": round(rating_score, 2),
        "heat_score": round(heat_score, 2),
        "trend_score": round(trend_score, 2),
        "norm_30d": round(norm_30d, 4),
        "norm_7d": round(norm_7d, 4),
        "norm_video": round(norm_video, 4),
        "raw_sales_30d_ratio": round((sales_30d / max_30d) if max_30d > 0 else 0, 6),
        "raw_sales_7d_ratio": round((score_input["sales_7d"] / max_7d) if max_7d > 0 else 0, 6),
        "raw_video_ratio": round((score_input["video_count"] / max_video) if max_video > 0 else 0, 6),
        "sales_curve_ratio": round((sales_score / 35) if 35 else 0, 4),
        "heat_curve_ratio": round((heat_score / 25) if 25 else 0, 4),
        "trend_curve_ratio": round(min(1, trend_ratio), 4),
        "trend_ratio": round(trend_ratio, 4),
    }
    detail["benchmark"] = {
        "category_key": score_input.get("category_key"),
        "max_sales_7d": round(max_7d, 2),
        "max_sales_30d": round(max_30d, 2),
        "max_video_count": round(max_video, 2),
    }
    detail["conclusion"] = conclusion
    detail["summary"] = reason
    return total, json.dumps(detail, ensure_ascii=False)

def score_sales_mom(current_sales, previous_sales):
    current_sales = clean_number(current_sales)
    previous_sales = clean_number(previous_sales)
    if previous_sales <= 0:
        return 100.0 if current_sales > 0 else 0.0
    return round(((current_sales - previous_sales) / previous_sales) * 100, 2)


def table_column_names(cursor, table):
    cursor.execute(
        """
        SELECT COLUMN_NAME
        FROM information_schema.COLUMNS
        WHERE TABLE_SCHEMA = DATABASE() AND TABLE_NAME = %s
        """,
        (table,),
    )
    return {row["COLUMN_NAME"] for row in cursor.fetchall()}


def build_score_select_sql(meta, columns):
    selected = ["`id`"]
    for field in {
        meta.get("id"),
        meta.get("date"),
        meta.get("title"),
        meta.get("rating"),
        meta.get("price"),
        meta.get("base_price"),
        meta.get("sold"),
        meta.get("total_sold"),
        meta.get("overview_7d"),
        meta.get("overview_30d"),
        meta.get("category_l1"),
        meta.get("category_l2"),
        meta.get("category_l3"),
        meta.get("aweme_count"),
        "aweme_count",
    }:
        if field and field in columns:
            selected.append(f"`{field}`")
    return ", ".join(dict.fromkeys(selected))


def recalculate_selection_scores(source=None, date_record=None, limit=None, force=False, latest_only=None):
    sources = [source] if source else ["unified", "fastmoss", "kalodata", "pod_cross_category"]
    summary = {}
    if latest_only is None:
        latest_only = SCORE_RECALC_LATEST_ONLY
    with db() as conn, conn.cursor() as cursor:
        for item_source in sources:
            try:
                if item_source not in SOURCES:
                    continue
                meta = SOURCES[item_source]
                columns = table_column_names(cursor, meta["table"])
                if not {"selection_score", "score_reason", "sales_mom"}.issubset(columns):
                    continue
                select_sql = build_score_select_sql(meta, columns)
                where = ["1=1"]
                params = []
                if date_record:
                    where.append(f"`{meta['date']}` = %s")
                    params.append(date_record)
                elif latest_only:
                    where.append(f"`{meta['date']}` = (SELECT MAX(latest_score_source.`{meta['date']}`) FROM `{meta['table']}` AS latest_score_source)")
                target_limit = int(limit) if limit else SCORE_DAEMON_BATCH_SIZE
                queue_order = "`id` ASC"
                unscored_where = [*where, "(selection_score IS NULL OR selection_score = '')"]
                safe_select_limit = min(target_limit, 200)
                cursor.execute(
                    f"""
                    SELECT {select_sql}
                    FROM `{meta['table']}`
                    WHERE {' AND '.join(unscored_where)}
                    ORDER BY {queue_order}
                    LIMIT %s
                    """,
                    [*params, safe_select_limit],
                )
                rows = list(cursor.fetchall())
                if len(rows) < target_limit and not force and "score_attempts" in columns:
                    remaining_limit = min(target_limit - len(rows), 200)
                    rescored_where = [
                        *where,
                        "selection_score IS NOT NULL",
                        "selection_score <> ''",
                        "score_attempts < %s",
                    ]
                    cursor.execute(
                        f"""
                        SELECT {select_sql}
                        FROM `{meta['table']}`
                        WHERE {' AND '.join(rescored_where)}
                        ORDER BY score_attempts ASC, score_updated_at ASC, {queue_order}
                        LIMIT %s
                        """,
                        [*params, SCORE_RECALC_MAX_ATTEMPTS, remaining_limit],
                    )
                    rows.extend(list(cursor.fetchall()))
                elif force and len(rows) < target_limit:
                    remaining_limit = min(target_limit - len(rows), 200)
                    cursor.execute(
                        f"""
                        SELECT {select_sql}
                        FROM `{meta['table']}`
                        WHERE {' AND '.join(where)}
                        ORDER BY {queue_order}
                        LIMIT %s
                        """,
                        [*params, remaining_limit],
                    )
                    seen_ids = {row["id"] for row in rows}
                    rows.extend([row for row in list(cursor.fetchall()) if row["id"] not in seen_ids])
                inputs = [(row, score_input_from_row(item_source, meta, row)) for row in rows]
                benchmarks = {}
                for _row, score_input in inputs:
                    bucket = benchmarks.setdefault(score_input["category_key"], {"max_sales_30d": 0, "max_sales_7d": 0, "max_video_count": 0})
                    bucket["max_sales_30d"] = max(bucket["max_sales_30d"], score_input["sales_30d"])
                    bucket["max_sales_7d"] = max(bucket["max_sales_7d"], score_input["sales_7d"])
                    bucket["max_video_count"] = max(bucket["max_video_count"], score_input["video_count"])
                updated = 0
                for row, score_input in inputs:
                    score, reason = calculate_selection_score(score_input, benchmarks.get(score_input["category_key"], {}))
                    mom = score_sales_mom(score_input["sales_7d"], score_input["sales_30d"] - score_input["sales_7d"])
                    set_parts = ["selection_score = %s", "score_reason = %s", "sales_mom = %s"]
                    update_values = [score, reason, mom]
                    if "score_attempts" in columns:
                        set_parts.append("score_attempts = LEAST(score_attempts + 1, %s)")
                        update_values.append(SCORE_RECALC_MAX_ATTEMPTS)
                    if "score_updated_at" in columns:
                        set_parts.append("score_updated_at = NOW()")
                    update_values.append(row["id"])
                    cursor.execute(
                        f"""
                        UPDATE `{meta['table']}`
                        SET {", ".join(set_parts)}
                        WHERE `id` = %s
                        """,
                        update_values,
                    )
                    updated += cursor.rowcount
                    if updated and updated % 200 == 0:
                        conn.commit()
                summary[item_source] = {"scanned": len(rows), "updated": updated}
                conn.commit()
            except pymysql.err.OperationalError as exc:
                safe_rollback(conn)
                summary[item_source] = {"scanned": 0, "updated": 0, "error": str(exc)}
                print(f"[SCORE_DAEMON] skip source={item_source} error={exc}", flush=True)
    return summary


def attach_runtime_metrics(cursor, meta, rows, sales_period, collection_start=None, collection_end=None):
    for row in rows:
        if row.get("runtime_sold_count") is not None:
            row["runtime_sold_count"] = float(row.get("runtime_sold_count") or 0)
        if row.get("previous_runtime_sold_count") is not None:
            row["previous_runtime_sold_count"] = float(row.get("previous_runtime_sold_count") or 0)
        if row.get("current_period_sold_count") is not None:
            row["current_period_sold_count"] = float(row.get("current_period_sold_count") or 0)
        if row.get("previous_period_sold_count") is not None:
            row["previous_period_sold_count"] = float(row.get("previous_period_sold_count") or 0)
        if row.get("current_period_sold_count") is not None and row.get("previous_period_sold_count") is not None:
            row["sales_growth_view"] = format_runtime_sales_result(
                {
                    "current_sum": row.get("current_period_sold_count"),
                    "previous_sum": row.get("previous_period_sold_count"),
                }
            )["growth"]
            continue
        if row.get("previous_runtime_sold_count") is not None:
            row["sales_growth_view"] = format_runtime_sales_result(
                {
                    "current_sum": row.get("runtime_sold_count"),
                    "previous_sum": row.get("previous_runtime_sold_count"),
                }
            )["growth"]
            continue
        metrics = calculate_runtime_sales_metrics(
            cursor,
            meta,
            row,
            sales_period,
            collection_start,
            collection_end,
        )
        row["sales_growth_view"] = metrics["growth"]


def calculate_runtime_sales_metrics(cursor, meta, row, sales_period, collection_start=None, collection_end=None):
    product_id = row.get("product_id")
    date_record = row.get("date_record")
    if not product_id or not date_record:
        return {"current_sum": None, "growth": "N/A"}
    table = meta["table"]
    id_field = meta["id"]
    date_field = meta["date"]
    sold_field = meta["sold"]
    sold_expr = numeric_sql_expr(f"`{sold_field}`")

    if collection_start and collection_end:
        return calculate_collection_sales_growth(
            cursor,
            meta,
            product_id,
            collection_start,
            collection_end,
            sold_expr,
        )
    elif sales_period["period"] == "custom" and sales_period["start"] and sales_period["end"]:
        current_start = sales_period["start"]
        current_end = sales_period["end"]
    else:
        days = sales_period["days"] or 7
        cursor.execute(
            f"""
            SELECT
              SUM(CASE
                WHEN `{date_field}` BETWEEN DATE_SUB(CAST(%s AS DATE), INTERVAL %s DAY) AND CAST(%s AS DATE)
                THEN COALESCE({sold_expr}, 0) ELSE 0
              END) AS current_sum,
              SUM(CASE
                WHEN `{date_field}` BETWEEN DATE_SUB(CAST(%s AS DATE), INTERVAL %s DAY)
                                     AND DATE_SUB(CAST(%s AS DATE), INTERVAL %s DAY)
                THEN COALESCE({sold_expr}, 0) ELSE 0
              END) AS previous_sum
            FROM `{table}`
            WHERE `{id_field}` = %s
            """,
            (
                date_record,
                days - 1,
                date_record,
                date_record,
                (days * 2) - 1,
                date_record,
                days,
                product_id,
            ),
        )
        result = cursor.fetchone() or {}
        return format_runtime_sales_result(result)

    cursor.execute(
        f"""
        SELECT
          SUM(CASE
            WHEN `{date_field}` BETWEEN CAST(%s AS DATE) AND CAST(%s AS DATE)
            THEN COALESCE({sold_expr}, 0) ELSE 0
          END) AS current_sum,
          SUM(CASE
            WHEN `{date_field}` BETWEEN DATE_SUB(CAST(%s AS DATE), INTERVAL DATEDIFF(CAST(%s AS DATE), CAST(%s AS DATE)) + 1 DAY)
                                 AND DATE_SUB(CAST(%s AS DATE), INTERVAL 1 DAY)
            THEN COALESCE({sold_expr}, 0) ELSE 0
          END) AS previous_sum
        FROM `{table}`
        WHERE `{id_field}` = %s
        """,
        (
            current_start,
            current_end,
            current_start,
            current_end,
            current_start,
            current_start,
            product_id,
        ),
    )
    result = cursor.fetchone() or {}
    return format_runtime_sales_result(result)


def calculate_collection_sales_growth(cursor, meta, product_id, current_start, current_end, sold_expr):
    date_count = collection_date_count(current_start, current_end)
    table = meta["table"]
    id_field = meta["id"]
    date_field = meta["date"]
    cursor.execute(
        f"""
        SELECT SUM(COALESCE({sold_expr}, 0)) AS current_sum
        FROM `{table}`
        WHERE `{id_field}` = %s
          AND `{date_field}` BETWEEN CAST(%s AS DATE) AND CAST(%s AS DATE)
        """,
        (product_id, current_start, current_end),
    )
    current_sum = number_or_zero((cursor.fetchone() or {}).get("current_sum"))

    cursor.execute(
        f"""
        SELECT DISTINCT `{date_field}` AS previous_date
        FROM `{table}`
        WHERE `{id_field}` = %s AND `{date_field}` < CAST(%s AS DATE)
        ORDER BY `{date_field}` DESC
        LIMIT %s
        """,
        (product_id, current_start, date_count),
    )
    previous_dates = [row["previous_date"] for row in cursor.fetchall()]
    if not previous_dates:
        return {"current_sum": current_sum, "growth": "N/A"}

    placeholders = ",".join(["%s"] * len(previous_dates))
    cursor.execute(
        f"""
        SELECT SUM(COALESCE({sold_expr}, 0)) AS previous_sum
        FROM `{table}`
        WHERE `{id_field}` = %s AND `{date_field}` IN ({placeholders})
        """,
        [product_id, *previous_dates],
    )
    previous_sum = number_or_zero((cursor.fetchone() or {}).get("previous_sum"))
    return format_runtime_sales_result({"current_sum": current_sum, "previous_sum": previous_sum})


def collection_date_count(current_start, current_end):
    try:
        start = datetime.fromisoformat(str(current_start)[:10]).date()
        end = datetime.fromisoformat(str(current_end)[:10]).date()
    except ValueError:
        return 1
    return max(1, (end - start).days + 1)


def format_runtime_sales_result(result):
    current_sum = float(result.get("current_sum") or 0)
    previous_sum = float(result.get("previous_sum") or 0)
    if previous_sum <= 0:
        return {"current_sum": current_sum, "growth": "N/A"}
    growth = (current_sum / previous_sum) - 1
    return {"current_sum": current_sum, "growth": f"{growth * 100:.2f}%"}


def ratio_from_distribution(distribution, name):
    if not isinstance(distribution, list):
        return ""
    for item in distribution:
        if str(item.get("name")) == name:
            return item.get("percentage") or item.get("ratio") or ""
    return ""


PRICE_CURRENCY_WORDS_RE = re.compile(
    r"\b(?:US|USD|EUR|GBP|CNY|RMB|AUD|CAD|HKD|SGD|MXN)\b",
    re.IGNORECASE,
)
PRICE_ALLOWED_CHARS_RE = re.compile(r"^[\s\d.,$€£¥￥%+\-~–—/()]+$")


def normalize_price_display(*values):
    for value in values:
        text = stringify(value).strip()
        if is_price_display(text):
            return text
    return ""


def is_price_display(text):
    if not text or not any(char.isdigit() for char in text):
        return False
    without_currency_words = PRICE_CURRENCY_WORDS_RE.sub("", text)
    return bool(PRICE_ALLOWED_CHARS_RE.fullmatch(without_currency_words))


def normalize_detail(source, row):
    meta = SOURCES[source]
    product_id = stringify(row.get(meta["id"]))
    date_record = row.get(meta["date"])
    raw = {}
    for key, value in row.items():
        if key == meta["image"]:
            raw[key] = "[图片字段已隐藏，可通过 image_url 查看]"
        else:
            raw[key] = serialize_value(value)
    return {
        "source": source,
        "product_id": product_id,
        "date_record": date_record,
        "title": row.get(meta["title"]),
        "attributes": serialize_value(row.get(meta.get("attributes"))) if meta.get("attributes") else "",
        "selling_points": serialize_value(row.get(meta.get("selling_points"))) if meta.get("selling_points") else "",
        "image_url": f"/api/products/{source}/{product_id}/image?date_record={date_record}",
        "platform_url": resolve_platform_url(source, product_id, row.get(meta["detail_url"])),
        "audit_status": row.get("audit_status") or "PENDING",
        "ip_grade": row.get("ip_grade"),
        "ip_reason": row.get("ip_reason"),
        "ip_tags": parse_json(row.get("ip_tags")),
        "material_analysis": parse_json(row.get("material_analysis")),
        "ai_analysis_error": parse_json(row.get("ai_analysis_error")),
        "period_fields": build_detail_period_fields(source, row, meta),
        "sku_analysis": build_detail_sku_analysis(row, meta),
        "illustration_status": row.get("illustration_status"),
        "illustration_result_url": stringify(row.get("illustration_result_url")),
        "illustration_extractable": row.get("illustration_extractable"),
        "illustration_extract_reason": stringify(row.get("illustration_extract_reason")),
        "selection_score": stringify(row.get("selection_score")),
        "score_reason": row.get("score_reason"),
        "sales_mom": stringify(row.get("sales_mom")),
        "raw_fields": raw,
    }


def build_detail_period_fields(source, row, meta):
    fields = {}
    for period in ("7d", "30d", "90d", "180d"):
        overview_field = meta.get(f"overview_{period}")
        distribution_field = meta.get(f"distribution_{period}")
        fields[f"overview_{period}"] = serialize_value(row.get(overview_field)) if overview_field else None
        fields[f"distribution_{period}"] = serialize_value(row.get(distribution_field)) if distribution_field else None

    if source == "kalodata" and not fields.get("overview_30d"):
        fields["overview_30d"] = json.dumps(build_kalodata_30d_overview(row), ensure_ascii=False)
    if source == "kalodata" and not fields.get("distribution_30d"):
        fields["distribution_30d"] = json.dumps(build_kalodata_30d_distribution(row), ensure_ascii=False)
    return fields


def build_detail_sku_analysis(row, meta):
    fields = {}
    for period, label in (("7d", "近7天"), ("28d", "近28天")):
        field = meta.get(f"sku_analysis_{period}")
        parsed = parse_json(row.get(field)) if field else None
        normalized = normalize_sku_analysis(parsed)
        if normalized:
            normalized["period"] = period
            normalized["period_label"] = label
            fields[period] = normalized
    return fields


def normalize_sku_analysis(parsed):
    if not isinstance(parsed, dict) or not parsed:
        return None

    rules = []
    dimensions_source = None
    for value in parsed.values():
        if isinstance(value, list) and not rules:
            rules = [stringify(item) for item in value if stringify(item)]
        elif isinstance(value, dict):
            if rules and any(rule in value for rule in rules):
                dimensions_source = value
            elif dimensions_source is None:
                dimensions_source = value

    if not isinstance(dimensions_source, dict) or not dimensions_source:
        return None
    if not rules:
        rules = [stringify(key) for key in dimensions_source.keys()]

    dimensions = []
    for rule in rules:
        dimension = dimensions_source.get(rule)
        if not isinstance(dimension, dict):
            continue
        sales_block = find_sku_sales_block(dimension)
        if not sales_block:
            continue
        items = normalize_sku_items(sales_block.get("items") or sales_block.get("details") or [])
        if not items:
            continue
        dimensions.append(
            {
                "name": rule,
                "total": sales_block.get("total"),
                "total_display": sales_block.get("total_display") or format_metric_number(sales_block.get("total")),
                "items": items,
            }
        )

    if not dimensions:
        return None

    best = find_best_sku_dimension(dimensions)
    return {
        "rules": [item["name"] for item in dimensions],
        "dimensions": dimensions,
        "best_sku": best,
    }


def find_sku_sales_block(dimension):
    fallback = None
    for metric in dimension.values():
        if not isinstance(metric, dict):
            continue
        details = first_list_value(metric)
        if not details:
            continue
        block = {
            "total": first_numeric_value(metric),
            "total_display": first_display_value(metric),
            "items": details,
        }
        if fallback is None:
            fallback = block
        if not sku_block_has_currency(details):
            return block
    return fallback


def sku_block_has_currency(details):
    if not isinstance(details, list):
        return False
    for detail in details:
        if not isinstance(detail, dict):
            continue
        for value in detail.values():
            text = stringify(value)
            if "$" in text or "USD" in text or "¥" in text or "￥" in text:
                return True
    return False


def normalize_sku_items(items):
    if not isinstance(items, list):
        return []
    normalized = []
    for item in items:
        if not isinstance(item, dict):
            continue
        name = first_sku_name(item)
        ratio = first_percent_value(item)
        sales = first_numeric_value(item)
        sales_display = first_display_value(item) or format_metric_number(sales)
        if not name:
            continue
        normalized.append(
            {
                "name": name,
                "ratio": ratio,
                "sales": sales,
                "sales_display": sales_display,
            }
        )
    return normalized


def first_list_value(value):
    for item in value.values():
        if isinstance(item, list):
            return item
    return None


def first_numeric_value(value):
    for item in value.values():
        if isinstance(item, (int, float, Decimal)):
            return number_or_zero(item)
    return 0


def first_display_value(value):
    for item in value.values():
        text = stringify(item).strip()
        if not text or text.endswith("%"):
            continue
        normalized = text.replace(",", "").replace("，", "").replace("$", "").replace("¥", "").replace("￥", "").strip()
        if normalized.replace(".", "", 1).isdigit():
            return text
    return ""


def sales_from_distribution(distribution, name, total_fallback=None):
    if not isinstance(distribution, list):
        return 0
    for item in distribution:
        if str(item.get("name")) != name:
            continue
        sales = number_or_zero(item.get("sales") or item.get("sold") or item.get("count") or item.get("value_count"))
        if sales > 0:
            return sales
        ratio = percent_to_number(item.get("percentage") or item.get("ratio"))
        if total_fallback is not None and ratio > 0:
            return number_or_zero(total_fallback) * ratio / 100
    return 0


def percent_to_number(value):
    text = stringify(value).replace("%", "").replace(",", "").strip()
    try:
        number = float(text)
    except ValueError:
        return 0
    return number * 100 if 0 < number <= 1 else number


def normalize_distribution_chart(distribution, total_fallback=None, direct_ratios=None):
    names = ["视频", "直播", "商品卡"]
    direct_ratios = direct_ratios or {}
    items = []
    for name in names:
        ratio = ratio_from_distribution(distribution, name)
        if not ratio and name in direct_ratios:
            ratio = direct_ratios[name]
        sales = sales_from_distribution(distribution, name, total_fallback)
        if sales <= 0 and ratio:
            sales = number_or_zero(total_fallback) * percent_to_number(ratio) / 100
        items.append(
            {
                "name": name,
                "ratio": normalize_percent_display(ratio),
                "sales": format_metric_number(sales),
            }
        )
    return items


def normalize_percent_display(value):
    number = percent_to_number(value)
    if not number:
        return "0%"
    if float(number).is_integer():
        return f"{int(number)}%"
    return f"{number:.2f}%"


def first_percent_value(value):
    for item in value.values():
        text = stringify(item).strip()
        if text.endswith("%"):
            return text
    return ""


def first_sku_name(value):
    for item in value.values():
        if isinstance(item, (dict, list)):
            continue
        text = stringify(item).strip()
        if not text or text.endswith("%"):
            continue
        if text.replace(".", "", 1).isdigit():
            continue
        if text.startswith("$"):
            continue
        return text
    return ""


def find_best_sku_dimension(dimensions):
    best = None
    fallback = None
    for dimension in dimensions:
        for item in dimension.get("items") or []:
            sales = number_or_zero(item.get("sales"))
            candidate = {
                "rule": dimension["name"],
                "name": item.get("name") or "",
                "sales": sales,
                "sales_display": item.get("sales_display") or format_metric_number(sales),
            }
            if fallback is None or sales > fallback["sales"]:
                fallback = candidate
            if stringify(item.get("name")).strip().lower() == "other":
                continue
            if best is None or sales > best["sales"]:
                best = candidate
    return best or fallback


def build_kalodata_30d_overview(row):
    return {
        "销售额": stringify(row.get("总成交额")),
        "日均销售额": stringify(row.get("日均成交额")),
        "销量": stringify(row.get("销量") or row.get("总销量")),
        "日均销量": stringify(row.get("日均销量")),
        "带货达人数": stringify(row.get("关联达人数")),
        "带货视频数": stringify(row.get("带货视频数")),
        "直播销售额": stringify(row.get("直播销售额")),
        "视频销售额": stringify(row.get("视频成交额")),
        "商品卡": stringify(row.get("商品卡成交额")),
    }


def build_kalodata_30d_distribution(row):
    video_ratio = stringify(row.get("近28天视频占比"))
    product_card_ratio = stringify(row.get("近28天商品卡占比"))
    video_amount = stringify(row.get("视频成交额"))
    product_card_amount = stringify(row.get("商品卡成交额"))
    return [
        {"name": "视频", "percentage": video_ratio, "sales": video_amount},
        {"name": "直播", "percentage": "0%", "sales": stringify(row.get("直播销售额"))},
        {"name": "商品卡", "percentage": product_card_ratio, "sales": product_card_amount},
    ]


def enrich_detail_period_fields(cursor, source, meta, row, detail):
    fields = detail.get("period_fields") or {}
    if fields.get("overview_30d") and fields.get("overview_30d") != fields.get("overview_7d"):
        return
    aggregate = build_collection_aggregate_overview(cursor, meta, row, 30)
    if aggregate:
        fields["overview_30d"] = json.dumps(aggregate, ensure_ascii=False)
        detail["period_fields"] = fields


def build_collection_aggregate_overview(cursor, meta, row, days):
    if not meta.get("sold"):
        return None
    product_id = row.get(meta["id"])
    date_record = row.get(meta["date"])
    if not product_id or not date_record:
        return None
    sold_expr = numeric_sql_expr(f"`{meta['sold']}`")
    sale_amount_expr = numeric_sql_expr(f"`{meta['sale_amount']}`") if meta.get("sale_amount") else "NULL"
    author_expr = numeric_sql_expr(f"`{meta['author_count']}`") if meta.get("author_count") else "NULL"
    cursor.execute(
        f"""
        SELECT
          SUM(COALESCE({sold_expr}, 0)) AS sold_sum,
          SUM(COALESCE({sale_amount_expr}, 0)) AS sale_amount_sum,
          MAX(COALESCE({author_expr}, 0)) AS author_count
        FROM `{meta['table']}`
        WHERE `{meta['id']}` = %s
          AND `{meta['date']}` BETWEEN DATE_SUB(CAST(%s AS DATE), INTERVAL %s DAY) AND CAST(%s AS DATE)
        """,
        (product_id, date_record, days - 1, date_record),
    )
    result = cursor.fetchone() or {}
    sold_sum = number_or_zero(result.get("sold_sum"))
    sale_amount_sum = number_or_zero(result.get("sale_amount_sum"))
    if sold_sum <= 0 and sale_amount_sum <= 0:
        return None
    return {
        "销量": format_metric_number(sold_sum),
        "日均销量": format_metric_number(sold_sum / days),
        "销售额": format_metric_number(sale_amount_sum),
        "日均销售额": format_metric_number(sale_amount_sum / days),
        "带货达人数": format_metric_number(number_or_zero(result.get("author_count"))),
    }


def number_or_zero(value):
    try:
        return float(value or 0)
    except (TypeError, ValueError):
        return 0.0


def format_metric_number(value):
    value = number_or_zero(value)
    if value.is_integer():
        return int(value)
    return round(value, 2)


def build_platform_url(source, product_id):
    if source in {"unified", "fastmoss", "pod_cross_category"}:
        return f"https://www.fastmoss.com/zh/e-commerce/detail/{product_id}"
    return build_tiktok_shop_url(product_id)


def resolve_platform_url(source, product_id, detail_url=None):
    if source == "kalodata":
        return detail_url or ""
    return detail_url or build_platform_url(source, product_id)


def build_tiktok_shop_url(product_id):
    slug = "mens-athletic-t-shirt-by-brand-lightweight-quick-dry-crew-neck-tops"
    return f"https://www.tiktok.com/shop/pdp/{slug}/{product_id}?source=ecommerce_store&region=US"


def fetch_ready_links(source):
    meta = SOURCES[source]
    with db() as conn, conn.cursor() as cursor:
        cursor.execute(
            f"""
            SELECT `{meta['id']}` AS product_id, `{meta['detail_url']}` AS detail_url
            FROM `{meta['table']}`
            WHERE audit_status = 'READY'
            ORDER BY `{meta['date']}` DESC
            """,
        )
        rows = cursor.fetchall()
    return [resolve_platform_url(source, row["product_id"], row.get("detail_url")) for row in rows]


def count_material_types(cursor, table, where=None, params=None):
    where = list(where or ["1=1"])
    params = list(params or [])
    where.append("material_analysis IS NOT NULL")
    where.append("material_analysis <> ''")
    where.append("JSON_VALID(material_analysis)")
    cursor.execute(
        f"""
        SELECT JSON_UNQUOTE(JSON_EXTRACT(material_analysis, '$.material_type')) AS material_type,
               COUNT(*) AS count
        FROM `{table}`
        WHERE {" AND ".join(where)}
        GROUP BY JSON_UNQUOTE(JSON_EXTRACT(material_analysis, '$.material_type'))
        """,
        params,
    )
    return {row["material_type"] or "未知": int(row["count"]) for row in cursor.fetchall()}


def serve_image_value(value):
    value = str(value).strip()
    if value.startswith("http://") or value.startswith("https://"):
        return redirect(value)
    mime = "image/jpeg"
    payload = value
    if value.startswith("data:image/"):
        header, payload = value.split(",", 1)
        mime = header.split(";")[0].replace("data:", "")
    try:
        data = base64.b64decode(payload, validate=False)
    except Exception:
        return api_error("图片 base64 无法解码", 422)
    return send_file(io.BytesIO(data), mimetype=mime)


def parse_json(value):
    if not value:
        return None
    if isinstance(value, (dict, list)):
        return value
    try:
        return json.loads(value)
    except Exception:
        return None


def format_decimal_plain(value):
    text = format(value.normalize(), "f")
    if "." in text:
        text = text.rstrip("0").rstrip(".")
    return text or "0"


def stringify(value):
    if value is None:
        return ""
    if isinstance(value, Decimal):
        return format_decimal_plain(value)
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d %H:%M:%S") if isinstance(value, datetime) else value.isoformat()
    if isinstance(value, str) and "e" in value.lower():
        stripped = value.strip()
        try:
            decimal_value = Decimal(stripped)
        except InvalidOperation:
            return value
        if decimal_value.is_finite():
            return format_decimal_plain(decimal_value)
    return str(value)


def serialize_value(value):
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, (datetime, date)):
        return stringify(value)
    parsed = parse_json(value)
    return parsed if parsed is not None else value


def start_ai_daemon_if_enabled():
    global _ai_daemon_started
    if _ai_daemon_started or not AI_DAEMON_ENABLED:
        return
    if os.getenv("WERKZEUG_RUN_MAIN") == "false":
        return
    invalid_sources = [source for source in AI_DAEMON_SOURCES if source not in AI_PRODUCT_SOURCES]
    if invalid_sources:
        print(f"[AI_DAEMON] AI_DAEMON_SOURCES invalid: {invalid_sources}", flush=True)
        return
    if AI_DAEMON_ANALYSIS not in {"ip", "material", "both"}:
        print(f"[AI_DAEMON] AI_DAEMON_ANALYSIS invalid: {AI_DAEMON_ANALYSIS}", flush=True)
        return

    thread = threading.Thread(target=ai_daemon_loop, name="cp-ai-analysis-daemon", daemon=True)
    thread.start()
    _ai_daemon_started = True
    print(
        "[AI_DAEMON] started "
        f"sources={','.join(AI_DAEMON_SOURCES)} analysis={AI_DAEMON_ANALYSIS} "
        f"batch_size={AI_DAEMON_BATCH_SIZE} interval={AI_DAEMON_INTERVAL_SECONDS}s "
        f"concurrency={AI_DAEMON_CONCURRENCY} lock_wait={AI_DAEMON_LOCK_WAIT_SECONDS}s "
        f"latest_only={AI_DAEMON_LATEST_ONLY} "
        f"write={AI_DAEMON_WRITE}",
        flush=True,
    )


def start_score_daemon_if_enabled():
    global _score_daemon_started
    if _score_daemon_started or not SCORE_DAEMON_ENABLED:
        return
    if os.getenv("WERKZEUG_RUN_MAIN") == "false":
        return
    thread = threading.Thread(target=score_daemon_loop, name="cp-selection-score-daemon", daemon=True)
    thread.start()
    _score_daemon_started = True
    print(
        f"[SCORE_DAEMON] started batch_size={SCORE_DAEMON_BATCH_SIZE} "
        f"interval={SCORE_DAEMON_INTERVAL_SECONDS}s "
        f"latest_only={SCORE_RECALC_LATEST_ONLY} max_attempts={SCORE_RECALC_MAX_ATTEMPTS}",
        flush=True,
    )


def start_illustration_daemon_if_enabled():
    global _illustration_daemon_started
    if _illustration_daemon_started or not ILLUSTRATION_DAEMON_ENABLED:
        return
    if os.getenv("WERKZEUG_RUN_MAIN") == "false":
        return
    valid_sources = [source for source in ILLUSTRATION_DAEMON_SOURCES if source == "pod_cross_category"]
    if not valid_sources:
        print(f"[ILLUSTRATION_DAEMON] no valid sources: {ILLUSTRATION_DAEMON_SOURCES}", flush=True)
        return
    thread = threading.Thread(target=illustration_daemon_loop, name="cp-illustration-check-daemon", daemon=True)
    thread.start()
    _illustration_daemon_started = True
    print(
        f"[ILLUSTRATION_DAEMON] started sources={','.join(valid_sources)} "
        f"batch_size={ILLUSTRATION_DAEMON_BATCH_SIZE} interval={ILLUSTRATION_DAEMON_INTERVAL_SECONDS}s "
        f"concurrency={ILLUSTRATION_DAEMON_CONCURRENCY} latest_only={ILLUSTRATION_DAEMON_LATEST_ONLY}",
        flush=True,
    )


def score_daemon_loop():
    while True:
        try:
            summary = recalculate_selection_scores(limit=SCORE_DAEMON_BATCH_SIZE)
            print(f"[SCORE_DAEMON] cycle finished summary={summary}", flush=True)
            record_daily_system_stats(
                "score_daemon_cycle",
                {"selection_scored": sum(int(item.get("updated") or 0) for item in summary.values())},
                summary,
            )
        except Exception:
            print("[SCORE_DAEMON] cycle failed", flush=True)
            traceback.print_exc()
        time.sleep(max(30, SCORE_DAEMON_INTERVAL_SECONDS))


def illustration_daemon_loop():
    while True:
        try:
            summary = run_illustration_daemon_once()
            print(f"[ILLUSTRATION_DAEMON] cycle finished summary={summary}", flush=True)
            record_daily_system_stats(
                "illustration_daemon_cycle",
                {"illustration_checked": sum(int(item.get("processed") or 0) for item in summary.values())},
                summary,
            )
        except Exception:
            print("[ILLUSTRATION_DAEMON] cycle failed", flush=True)
            traceback.print_exc()
        time.sleep(max(30, ILLUSTRATION_DAEMON_INTERVAL_SECONDS))


def run_illustration_daemon_once():
    summary = {}
    for source in ILLUSTRATION_DAEMON_SOURCES:
        if source != "pod_cross_category":
            continue
        keys = fetch_illustration_daemon_product_keys(source, ILLUSTRATION_DAEMON_BATCH_SIZE)
        processed = 0
        failed = 0
        if keys:
            max_workers = min(ILLUSTRATION_DAEMON_CONCURRENCY, len(keys))
            with ThreadPoolExecutor(max_workers=max_workers) as executor:
                futures = [executor.submit(process_single_illustration_check, source, key) for key in keys]
                for future in as_completed(futures):
                    try:
                        if future.result():
                            processed += 1
                        else:
                            failed += 1
                    except Exception:
                        failed += 1
                        traceback.print_exc()
        summary[source] = {"queued": len(keys), "processed": processed, "failed": failed}
    return summary


def fetch_illustration_daemon_product_keys(source, limit):
    meta = SOURCES[source]
    image_field = meta.get("image")
    where = [
        "illustration_extractable IS NULL",
        f"`{image_field}` IS NOT NULL",
        f"`{image_field}` <> ''",
    ]
    params = []
    if ILLUSTRATION_DAEMON_LATEST_ONLY:
        where.append(f"`{meta['date']}` = (SELECT MAX(t2.`{meta['date']}`) FROM `{meta['table']}` t2)")
    with db() as conn, conn.cursor() as cursor:
        cursor.execute(
            f"""
            SELECT `{meta['id']}` AS product_id, `{meta['date']}` AS date_record
            FROM `{meta['table']}`
            WHERE {" AND ".join(where)}
            ORDER BY `{meta['date']}` DESC, id ASC
            LIMIT %s
            """,
            [*params, limit],
        )
        return list(cursor.fetchall())


def process_single_illustration_check(source, key):
    meta = SOURCES[source]
    product_id = key.get("product_id")
    date_record = key.get("date_record")
    try:
        with db() as conn, conn.cursor() as cursor:
            cursor.execute(
                f"""
                SELECT `{meta['id']}` AS product_id, `{meta['date']}` AS date_record,
                       {sql_alias(meta.get("title"), "title")},
                       {sql_alias(meta.get("image"), "image_value")}
                FROM `{meta['table']}`
                WHERE `{meta['id']}` = %s AND `{meta['date']}` = %s
                LIMIT 1
                """,
                (product_id, date_record),
            )
            product = cursor.fetchone()
            if not product or not product.get("image_value"):
                return False
            result = analyze_illustration_extractability(product.get("image_value"), product.get("title") or "")
            extractable = 1 if result.get("can_extract_illustration") else 0
            reason = stringify(result.get("reason")).strip()[:500]
            cursor.execute(
                f"""
                UPDATE `{meta['table']}`
                SET illustration_extractable = %s,
                    illustration_extract_reason = %s
                WHERE `{meta['id']}` = %s AND `{meta['date']}` = %s
                  AND illustration_extractable IS NULL
                """,
                (extractable, reason, product_id, date_record),
            )
            conn.commit()
            print(
                f"[ILLUSTRATION_DAEMON] checked source={source} product_id={product_id} "
                f"date_record={date_record} extractable={extractable}",
                flush=True,
            )
            return True
    except Exception as exc:
        print(
            f"[ILLUSTRATION_DAEMON] item failed source={source} product_id={product_id} "
            f"date_record={date_record} error={exc}",
            flush=True,
        )
        return False


def ai_daemon_loop():
    while True:
        try:
            summary = run_ai_daemon_once()
            print(f"[AI_DAEMON] cycle finished summary={summary}", flush=True)
            record_daily_system_stats(
                "ai_daemon_cycle",
                {
                    "ip_analyzed": sum(int(item.get("ip_analyzed") or 0) for item in summary.values()),
                    "material_analyzed": sum(int(item.get("material_analyzed") or 0) for item in summary.values()),
                },
                summary,
            )
        except Exception:
            print("[AI_DAEMON] cycle failed", flush=True)
            traceback.print_exc()
        time.sleep(max(10, AI_DAEMON_INTERVAL_SECONDS))


def run_ai_daemon_once():
    summary = {}
    for source in AI_DAEMON_SOURCES:
        print(f"[AI_DAEMON] fetching source={source}", flush=True)
        keys = fetch_ai_daemon_product_keys(source, AI_DAEMON_ANALYSIS, AI_DAEMON_BATCH_SIZE)
        print(f"[AI_DAEMON] source={source} queued={len(keys)}", flush=True)
        source_summary = {
            "queued": len(keys),
            "processed": 0,
            "ip_analyzed": 0,
            "material_analyzed": 0,
        }
        summary[source] = source_summary
        if not keys:
            continue

        with ThreadPoolExecutor(max_workers=AI_DAEMON_CONCURRENCY) as executor:
            futures = {
                executor.submit(process_single_ai_task, source, key): key
                for key in keys
            }
            for future in as_completed(futures):
                try:
                    result = future.result()
                    if result:
                        source_summary["processed"] += 1
                        source_summary["ip_analyzed"] += int(result.get("ip_analyzed") or 0)
                        source_summary["material_analyzed"] += int(result.get("material_analyzed") or 0)
                except Exception:
                    traceback.print_exc()
    return summary


def process_single_ai_task(source, key):
    for attempt in range(1, AI_DAEMON_LOCK_RETRIES + 1):
        result = process_single_ai_task_once(source, key, attempt)
        if result is not None:
            return result
        time.sleep(0.5 * attempt)
    return False


def ai_analysis_delta(before_state, after_state):
    return {
        "ip_analyzed": 1 if is_blank(before_state.get("ip_grade")) and not is_blank(after_state.get("ip_grade")) else 0,
        "material_analyzed": 1 if is_blank(before_state.get("material_analysis")) and not is_blank(after_state.get("material_analysis")) else 0,
    }


def process_single_ai_task_once(source, key, attempt=1):
    conn = db()
    try:
        with conn.cursor() as cursor:
            cursor.execute(f"SET SESSION innodb_lock_wait_timeout = {AI_DAEMON_LOCK_WAIT_SECONDS}")
            product = load_ai_product(cursor, source, key["product_id"], key["date_record"])
            before_state = fetch_current_analysis_state(cursor, source, product)
            reused_fields = reuse_existing_product_analysis(cursor, source, product, AI_DAEMON_ANALYSIS)
            if reused_fields:
                conn.commit()
            if reused_fields and not needs_ai_analysis(cursor, source, product, AI_DAEMON_ANALYSIS):
                after_state = fetch_current_analysis_state(cursor, source, product)
                delta = ai_analysis_delta(before_state, after_state)
                print(
                    f"[AI_DAEMON][{threading.current_thread().name}] reused "
                    f"{','.join(reused_fields)} source={source} "
                    f"product_id={product['product_id']} date_record={product.get('date_record')}",
                    flush=True,
                )
                return {"processed": True, **delta}
            if not needs_ai_analysis(cursor, source, product, AI_DAEMON_ANALYSIS):
                print(
                    f"[AI_DAEMON][{threading.current_thread().name}] skip already analyzed "
                    f"source={source} product_id={product['product_id']} "
                    f"date_record={product.get('date_record')}",
                    flush=True,
                )
                conn.commit()
                return False
            print(
                f"[AI_DAEMON][{threading.current_thread().name}] analyzing "
                f"source={source} product_id={product['product_id']} "
                f"date_record={product.get('date_record')}",
                flush=True,
            )
            run_ai_product_flow(cursor, source, product, AI_DAEMON_ANALYSIS, AI_DAEMON_WRITE)
            after_state = fetch_current_analysis_state(cursor, source, product)
            delta = ai_analysis_delta(before_state, after_state)
            conn.commit()
            return {"processed": True, **delta}
    except pymysql.err.OperationalError as exc:
        safe_rollback(conn)
        if exc.args and exc.args[0] == 1205:
            if attempt < AI_DAEMON_LOCK_RETRIES:
                print(
                    f"[AI_DAEMON][{threading.current_thread().name}] lock wait timeout, retry "
                    f"{attempt}/{AI_DAEMON_LOCK_RETRIES} source={source} "
                    f"product_id={key.get('product_id')} date_record={key.get('date_record')}",
                    flush=True,
                )
                return None
            print(
                f"[AI_DAEMON][{threading.current_thread().name}] lock wait timeout, skip this cycle "
                f"source={source} product_id={key.get('product_id')} "
                f"date_record={key.get('date_record')}",
                flush=True,
            )
            return False
        print(
            f"[AI_DAEMON][{threading.current_thread().name}] item failed "
            f"source={source} product_id={key.get('product_id')} "
            f"date_record={key.get('date_record')}",
            flush=True,
        )
        traceback.print_exc()
        return False
    except Exception:
        safe_rollback(conn)
        print(
            f"[AI_DAEMON][{threading.current_thread().name}] item failed "
            f"source={source} product_id={key.get('product_id')} "
            f"date_record={key.get('date_record')}",
            flush=True,
        )
        traceback.print_exc()
        return False
    finally:
        safe_close(conn)


def reuse_existing_product_analysis(cursor, source, product, analysis):
    if not AI_DAEMON_WRITE or source not in FASTMOSS_ANALYSIS_REUSE_SOURCES:
        return []

    current = fetch_current_analysis_state(cursor, source, product)
    reusable = fetch_reusable_fastmoss_analysis(cursor, product["product_id"])
    if not reusable:
        return []

    updates = []
    params = []
    guards = []
    if analysis in {"ip", "both"} and is_blank(current.get("ip_grade")) and not is_blank(reusable.get("ip_grade")):
        updates.extend(["ip_grade = %s", "ip_reason = %s", "ip_tags = %s"])
        params.extend([
            reusable.get("ip_grade"),
            reusable.get("ip_reason") or "",
            reusable.get("ip_tags") or "",
        ])
        guards.append("(ip_grade IS NULL OR ip_grade = '')")
    if analysis in {"material", "both"} and is_blank(current.get("material_analysis")) and not is_blank(reusable.get("material_analysis")):
        updates.append("material_analysis = %s")
        params.append(reusable.get("material_analysis"))
        guards.append("(material_analysis IS NULL OR material_analysis = '')")
    if not updates:
        return []

    meta = AI_PRODUCT_SOURCES[source]
    params.extend([str(product["product_id"]), product.get("date_record")])
    guard_sql = f" AND ({' OR '.join(guards)})" if guards else ""
    cursor.execute(
        f"""
        UPDATE `{meta['table']}`
        SET {", ".join(updates)}
        WHERE `{meta['id_field']}` = %s AND `{meta['date_field']}` = %s
          {guard_sql}
        """,
        params,
    )
    return [
        field
        for field in ("ip", "material")
        if any(update.startswith("ip_") if field == "ip" else update.startswith("material_") for update in updates)
    ]


def fetch_current_analysis_state(cursor, source, product):
    meta = AI_PRODUCT_SOURCES[source]
    cursor.execute(
        f"""
        SELECT ip_grade, ip_reason, ip_tags, material_analysis
        FROM `{meta['table']}`
        WHERE `{meta['id_field']}` = %s AND `{meta['date_field']}` = %s
        LIMIT 1
        """,
        (str(product["product_id"]), product.get("date_record")),
    )
    return cursor.fetchone() or {}


def fetch_reusable_fastmoss_analysis(cursor, product_id):
    candidates = []
    for source in FASTMOSS_ANALYSIS_REUSE_SOURCES:
        meta = AI_PRODUCT_SOURCES[source]
        cursor.execute(
            f"""
            SELECT ip_grade, ip_reason, ip_tags, material_analysis
            FROM `{meta['table']}`
            WHERE `{meta['id_field']}` = %s
              AND (
                (ip_grade IS NOT NULL AND ip_grade <> '')
                OR (material_analysis IS NOT NULL AND material_analysis <> '')
              )
            ORDER BY
              ((ip_grade IS NOT NULL AND ip_grade <> '') + (material_analysis IS NOT NULL AND material_analysis <> '')) DESC,
              `{meta['date_field']}` DESC
            LIMIT 1
            """,
            (str(product_id),),
        )
        row = cursor.fetchone()
        if row:
            candidates.append(row)
    if not candidates:
        return None
    reusable = {}
    for row in candidates:
        if is_blank(reusable.get("ip_grade")) and not is_blank(row.get("ip_grade")):
            reusable["ip_grade"] = row.get("ip_grade")
            reusable["ip_reason"] = row.get("ip_reason")
            reusable["ip_tags"] = row.get("ip_tags")
        if is_blank(reusable.get("material_analysis")) and not is_blank(row.get("material_analysis")):
            reusable["material_analysis"] = row.get("material_analysis")
    return reusable


def needs_ai_analysis(cursor, source, product, analysis):
    current = fetch_current_analysis_state(cursor, source, product)
    if analysis == "material":
        return is_blank(current.get("material_analysis"))
    if analysis == "ip":
        return is_blank(current.get("ip_grade"))
    return is_blank(current.get("material_analysis")) or is_blank(current.get("ip_grade"))


def is_blank(value):
    return value is None or value == ""


def run_ai_product_flow(cursor, source, product, analysis, write):
    if analysis == "both":
        prefilter_result = run_material_prefilter(cursor, source, product, write=write)
        if prefilter_result:
            print(
                "[AI_DAEMON] material prefilter hit, skip material AI but continue IP analysis "
                f"product_id={product['product_id']}",
                flush=True,
            )
            run_analysis(cursor, source, product, "IP", write=write)
            return
        run_analysis(cursor, source, product, "IP", write=write)
        run_analysis(cursor, source, product, "MATERIAL", write=write)
        return

    if analysis == "ip":
        run_analysis(cursor, source, product, "IP", write=write)
        return

    if analysis == "material":
        prefilter_result = run_material_prefilter(cursor, source, product, write=write)
        if prefilter_result:
            print(
                "[AI_DAEMON] material prefilter hit, skip material AI "
                f"product_id={product['product_id']}",
                flush=True,
            )
            return
        run_analysis(cursor, source, product, "MATERIAL", write=write)
        return

    run_analysis(cursor, source, product, "IP", write=write)


def fetch_ai_daemon_product_keys(source, analysis, limit):
    meta = AI_PRODUCT_SOURCES[source]
    table = meta["table"]
    id_field = meta["id_field"]
    date_field = meta["date_field"]
    image_field = meta["image_field"]
    where = [f"`{image_field}` IS NOT NULL", f"`{image_field}` <> ''"]
    params = []
    if AI_DAEMON_LATEST_ONLY:
        latest_date = fetch_ai_daemon_latest_date(meta)
        if not latest_date:
            return []
        where.append(f"`{date_field}` = %s")
        params.append(latest_date)

    if analysis == "both":
        keys = []
        seen = set()
        with db() as conn, conn.cursor() as cursor:
            for condition in [
                "(ip_grade IS NULL OR ip_grade = '')",
                "(material_analysis IS NULL OR material_analysis = '')",
            ]:
                remaining = limit - len(keys)
                if remaining <= 0:
                    break
                for row in query_ai_daemon_product_keys(cursor, table, id_field, date_field, where, params, condition, remaining):
                    key = (str(row.get("product_id")), str(row.get("date_record")))
                    if key in seen:
                        continue
                    seen.add(key)
                    keys.append(row)
                    if len(keys) >= limit:
                        break
        return keys

    condition = "(material_analysis IS NULL OR material_analysis = '')" if analysis == "material" else "(ip_grade IS NULL OR ip_grade = '')"
    with db() as conn, conn.cursor() as cursor:
        return query_ai_daemon_product_keys(cursor, table, id_field, date_field, where, params, condition, limit)


def query_ai_daemon_product_keys(cursor, table, id_field, date_field, base_where, base_params, condition, limit):
    where = [*base_where, condition]
    sql = f"""
        SELECT `{id_field}` AS product_id, `{date_field}` AS date_record
        FROM `{table}`
        WHERE {" AND ".join(where)}
        ORDER BY `id` DESC
        LIMIT %s
    """
    cursor.execute(sql, [*base_params, limit])
    return list(cursor.fetchall())


def fetch_ai_daemon_latest_date(meta):
    with db() as conn, conn.cursor() as cursor:
        cursor.execute(f"SELECT MAX(`{meta['date_field']}`) AS latest_date FROM `{meta['table']}`")
        row = cursor.fetchone() or {}
        return row.get("latest_date")


app = create_app()


if __name__ == "__main__":
    start_ai_daemon_if_enabled()
    start_score_daemon_if_enabled()
    start_illustration_daemon_if_enabled()
    debug_enabled = os.getenv("FLASK_DEBUG", "true").strip().lower() == "true"
    app.run(
        host="0.0.0.0",
        port=int(os.getenv("PORT", "5000")),
        debug=debug_enabled,
        use_reloader=False,
    )
