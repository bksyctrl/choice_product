import base64
import hashlib
import json
import mimetypes
import os
import posixpath
import zipfile
import xml.etree.ElementTree as ET
from datetime import datetime
from pathlib import Path

import openpyxl
import pymysql


DB_CONFIG = {
    "host": "192.168.0.168",
    "port": 3306,
    "user": "ITaimysql",
    "password": "Ai12345678@",
    "database": "ecommerce_workflow",
    "charset": "utf8mb4",
    "cursorclass": pymysql.cursors.DictCursor,
}

WORKBOOKS = [
    {
        "rule_type": "IP",
        "source_name": "美区IP款侵权插画",
        "source_path": r"D:\choice_product\美区IP款侵权插画.xlsx",
        "output_dir": r"D:\choice_product\ai_rule_images\ip",
    },
    {
        "rule_type": "MATERIAL",
        "source_name": "工厂材质",
        "source_path": r"D:\choice_product\工厂材质.xlsx",
        "output_dir": r"D:\choice_product\ai_rule_images\material",
    },
]

NS = {
    "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
}


DDL = [
    """
    CREATE TABLE IF NOT EXISTS `ai_rule_source_file` (
      `id` bigint(20) NOT NULL AUTO_INCREMENT,
      `rule_type` varchar(20) NOT NULL COMMENT 'IP=IP规则, MATERIAL=材质规则',
      `source_name` varchar(255) NOT NULL COMMENT '规则来源名称',
      `source_path` varchar(1000) NOT NULL COMMENT '原始Excel路径',
      `file_sha256` char(64) NOT NULL COMMENT '原始文件SHA256',
      `file_size_bytes` bigint(20) NOT NULL DEFAULT 0 COMMENT '原始文件大小',
      `sheet_names` longtext DEFAULT NULL COMMENT '工作表名称JSON',
      `title_text` varchar(255) DEFAULT NULL COMMENT 'Excel首个非空标题文本',
      `total_images` int(11) NOT NULL DEFAULT 0 COMMENT '内嵌图片数量',
      `imported_at` datetime NOT NULL DEFAULT CURRENT_TIMESTAMP,
      `updated_at` datetime NOT NULL DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP,
      PRIMARY KEY (`id`),
      UNIQUE KEY `uk_rule_source_file_sha` (`file_sha256`),
      KEY `idx_rule_source_file_type` (`rule_type`)
    ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COMMENT='AI规则Excel来源文件';
    """,
    """
    CREATE TABLE IF NOT EXISTS `ai_rule_image` (
      `id` bigint(20) NOT NULL AUTO_INCREMENT,
      `rule_type` varchar(20) NOT NULL COMMENT 'IP=IP规则图, MATERIAL=材质规则图',
      `source_file_id` bigint(20) NOT NULL COMMENT '来源文件ID',
      `source_sheet` varchar(255) DEFAULT NULL COMMENT '来源工作表',
      `image_index` int(11) NOT NULL COMMENT '来源文件内图片序号',
      `anchor_row` int(11) DEFAULT NULL COMMENT 'Excel中图片锚点行',
      `anchor_col` int(11) DEFAULT NULL COMMENT 'Excel中图片锚点列',
      `original_media_name` varchar(255) DEFAULT NULL COMMENT 'xlsx内部媒体文件名',
      `image_mime` varchar(100) DEFAULT NULL COMMENT '图片MIME类型',
      `image_ext` varchar(20) DEFAULT NULL COMMENT '图片扩展名',
      `image_sha256` char(64) NOT NULL COMMENT '图片SHA256',
      `image_size_bytes` bigint(20) NOT NULL DEFAULT 0 COMMENT '图片大小',
      `local_path` varchar(1000) NOT NULL COMMENT '抽取后的本地图片路径',
      `image_base64` longtext DEFAULT NULL COMMENT '图片base64，默认不填，后续需要时再补',
      `status` varchar(20) NOT NULL DEFAULT 'PENDING' COMMENT 'PENDING=待识别, DONE=已识别, FAILED=识别失败',
      `created_at` datetime NOT NULL DEFAULT CURRENT_TIMESTAMP,
      `updated_at` datetime NOT NULL DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP,
      PRIMARY KEY (`id`),
      UNIQUE KEY `uk_rule_image_sha` (`image_sha256`),
      KEY `idx_rule_image_type_status` (`rule_type`, `status`),
      KEY `idx_rule_image_source` (`source_file_id`)
    ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COMMENT='AI规则图片样本';
    """,
    """
    CREATE TABLE IF NOT EXISTS `ai_rule_image_analysis` (
      `id` bigint(20) NOT NULL AUTO_INCREMENT,
      `rule_image_id` bigint(20) NOT NULL COMMENT '规则图片ID',
      `provider` varchar(50) DEFAULT NULL COMMENT '识别服务商',
      `model_name` varchar(100) DEFAULT NULL COMMENT '识别模型',
      `status` varchar(20) NOT NULL DEFAULT 'PENDING' COMMENT 'PENDING=待识别, SUCCESS=成功, FAILED=失败',
      `image_description` longtext DEFAULT NULL COMMENT '图片整体描述',
      `ocr_text` longtext DEFAULT NULL COMMENT 'OCR文字',
      `logo_clues` longtext DEFAULT NULL COMMENT 'Logo线索JSON',
      `character_clues` longtext DEFAULT NULL COMMENT '角色/IP线索JSON',
      `material_clues` longtext DEFAULT NULL COMMENT '材质线索JSON',
      `visual_keywords` longtext DEFAULT NULL COMMENT '视觉关键词JSON',
      `raw_response` longtext DEFAULT NULL COMMENT '模型原始返回',
      `error_message` longtext DEFAULT NULL COMMENT '失败原因',
      `analyzed_at` datetime DEFAULT NULL,
      `created_at` datetime NOT NULL DEFAULT CURRENT_TIMESTAMP,
      `updated_at` datetime NOT NULL DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP,
      PRIMARY KEY (`id`),
      UNIQUE KEY `uk_rule_image_analysis_image` (`rule_image_id`),
      KEY `idx_rule_image_analysis_status` (`status`)
    ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COMMENT='规则图片AI识别结果';
    """,
    """
    CREATE TABLE IF NOT EXISTS `ai_ip_rule` (
      `id` bigint(20) NOT NULL AUTO_INCREMENT,
      `rule_image_id` bigint(20) DEFAULT NULL COMMENT '关联规则图片',
      `keyword` varchar(255) DEFAULT NULL COMMENT '关键词',
      `aliases` longtext DEFAULT NULL COMMENT '别名JSON',
      `risk_grade` varchar(10) DEFAULT NULL COMMENT 'S/A/B/C/D/E',
      `ip_type` varchar(100) DEFAULT NULL COMMENT 'IP类型',
      `description` longtext DEFAULT NULL COMMENT '规则说明',
      `visual_keywords` longtext DEFAULT NULL COMMENT '视觉关键词JSON',
      `confidence` decimal(5,2) DEFAULT NULL COMMENT '规则置信度',
      `verified` tinyint(1) NOT NULL DEFAULT 0 COMMENT '是否人工确认',
      `enabled` tinyint(1) NOT NULL DEFAULT 1 COMMENT '是否启用',
      `created_at` datetime NOT NULL DEFAULT CURRENT_TIMESTAMP,
      `updated_at` datetime NOT NULL DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP,
      PRIMARY KEY (`id`),
      KEY `idx_ai_ip_rule_grade` (`risk_grade`),
      KEY `idx_ai_ip_rule_enabled` (`enabled`)
    ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COMMENT='IP规则库';
    """,
    """
    CREATE TABLE IF NOT EXISTS `ai_material_rule` (
      `id` bigint(20) NOT NULL AUTO_INCREMENT,
      `rule_image_id` bigint(20) DEFAULT NULL COMMENT '关联规则图片',
      `keyword` varchar(255) DEFAULT NULL COMMENT '关键词',
      `aliases` longtext DEFAULT NULL COMMENT '别名JSON',
      `material_type` varchar(100) DEFAULT NULL COMMENT '材质大类',
      `material_category` varchar(100) DEFAULT NULL COMMENT '材质细分类',
      `description` longtext DEFAULT NULL COMMENT '规则说明',
      `visual_keywords` longtext DEFAULT NULL COMMENT '视觉关键词JSON',
      `confidence` decimal(5,2) DEFAULT NULL COMMENT '规则置信度',
      `verified` tinyint(1) NOT NULL DEFAULT 0 COMMENT '是否人工确认',
      `enabled` tinyint(1) NOT NULL DEFAULT 1 COMMENT '是否启用',
      `created_at` datetime NOT NULL DEFAULT CURRENT_TIMESTAMP,
      `updated_at` datetime NOT NULL DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP,
      PRIMARY KEY (`id`),
      KEY `idx_ai_material_rule_category` (`material_type`, `material_category`),
      KEY `idx_ai_material_rule_enabled` (`enabled`)
    ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COMMENT='材质规则库';
    """,
    """
    CREATE TABLE IF NOT EXISTS `ai_experience` (
      `id` bigint(20) NOT NULL AUTO_INCREMENT,
      `experience_type` varchar(30) NOT NULL COMMENT 'SUCCESS=成功经验, FAILURE=失败经验, CORRECTION=人工纠错',
      `analysis_type` varchar(20) NOT NULL COMMENT 'IP/MATERIAL/IMAGE',
      `title` varchar(255) DEFAULT NULL COMMENT '经验标题',
      `content` longtext NOT NULL COMMENT '经验内容',
      `source_table` varchar(100) DEFAULT NULL COMMENT '来源商品表',
      `product_id` varchar(100) DEFAULT NULL COMMENT '来源商品ID',
      `rule_image_id` bigint(20) DEFAULT NULL COMMENT '来源规则图片ID',
      `confidence` decimal(5,2) NOT NULL DEFAULT 0.30 COMMENT '可信度',
      `verified` tinyint(1) NOT NULL DEFAULT 0 COMMENT '是否人工确认',
      `enabled` tinyint(1) NOT NULL DEFAULT 1 COMMENT '是否启用',
      `created_at` datetime NOT NULL DEFAULT CURRENT_TIMESTAMP,
      `updated_at` datetime NOT NULL DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP,
      PRIMARY KEY (`id`),
      KEY `idx_ai_experience_type` (`experience_type`, `analysis_type`, `enabled`),
      KEY `idx_ai_experience_product` (`source_table`, `product_id`)
    ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COMMENT='AI成功经验、失败经验、人工纠错记忆';
    """,
    """
    CREATE TABLE IF NOT EXISTS `ai_product_image_analysis` (
      `id` bigint(20) NOT NULL AUTO_INCREMENT,
      `source_table` varchar(100) NOT NULL COMMENT '商品来源表',
      `product_id` varchar(100) NOT NULL COMMENT '商品ID',
      `date_record` varchar(10) DEFAULT NULL COMMENT '采集日期',
      `image_sha256` char(64) DEFAULT NULL COMMENT '商品主图SHA256',
      `provider` varchar(50) DEFAULT NULL COMMENT '识别服务商',
      `model_name` varchar(100) DEFAULT NULL COMMENT '识别模型',
      `status` varchar(20) NOT NULL DEFAULT 'PENDING' COMMENT 'PENDING=待识别, SUCCESS=成功, FAILED=失败',
      `image_description` longtext DEFAULT NULL COMMENT '图片整体描述',
      `ocr_text` longtext DEFAULT NULL COMMENT 'OCR文字',
      `logo_clues` longtext DEFAULT NULL COMMENT 'Logo线索JSON',
      `character_clues` longtext DEFAULT NULL COMMENT '角色/IP线索JSON',
      `material_clues` longtext DEFAULT NULL COMMENT '材质线索JSON',
      `visual_keywords` longtext DEFAULT NULL COMMENT '视觉关键词JSON',
      `raw_response` longtext DEFAULT NULL COMMENT '模型原始返回',
      `error_message` longtext DEFAULT NULL COMMENT '失败原因',
      `analyzed_at` datetime DEFAULT NULL,
      `created_at` datetime NOT NULL DEFAULT CURRENT_TIMESTAMP,
      `updated_at` datetime NOT NULL DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP,
      PRIMARY KEY (`id`),
      UNIQUE KEY `uk_product_image_analysis` (`source_table`, `product_id`, `date_record`),
      KEY `idx_product_image_analysis_status` (`status`)
    ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COMMENT='商品主图AI识别结果';
    """,
    """
    CREATE TABLE IF NOT EXISTS `ai_analysis_log` (
      `id` bigint(20) NOT NULL AUTO_INCREMENT,
      `source_table` varchar(100) NOT NULL COMMENT '商品来源表',
      `product_id` varchar(100) NOT NULL COMMENT '商品ID',
      `date_record` varchar(10) DEFAULT NULL COMMENT '采集日期',
      `analysis_type` varchar(20) NOT NULL COMMENT 'IP/MATERIAL/IMAGE',
      `provider` varchar(50) DEFAULT NULL COMMENT 'AI服务商',
      `model_name` varchar(100) DEFAULT NULL COMMENT '模型名称',
      `prompt_version` varchar(50) DEFAULT NULL COMMENT '提示词版本',
      `input_snapshot` longtext DEFAULT NULL COMMENT '输入商品快照JSON',
      `retrieved_rules` longtext DEFAULT NULL COMMENT '召回规则JSON',
      `retrieved_experiences` longtext DEFAULT NULL COMMENT '召回经验JSON',
      `request_payload` longtext DEFAULT NULL COMMENT '请求体JSON',
      `raw_response` longtext DEFAULT NULL COMMENT '原始返回',
      `parsed_result` longtext DEFAULT NULL COMMENT '解析后的JSON结果',
      `status` varchar(20) NOT NULL COMMENT 'SUCCESS=成功, FAILED=失败',
      `error_message` longtext DEFAULT NULL COMMENT '失败原因',
      `created_at` datetime NOT NULL DEFAULT CURRENT_TIMESTAMP,
      PRIMARY KEY (`id`),
      KEY `idx_ai_analysis_log_product` (`source_table`, `product_id`, `date_record`),
      KEY `idx_ai_analysis_log_type_status` (`analysis_type`, `status`)
    ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COMMENT='AI商品分析日志';
    """,
]


def sha256_file(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def first_non_empty_cell(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                for value in row:
                    if value is not None and str(value).strip():
                        return str(value).strip()
    finally:
        wb.close()
    return None


def workbook_sheet_names(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        return wb.sheetnames
    finally:
        wb.close()


def load_rels(zip_file, drawing_path):
    rel_path = "xl/drawings/_rels/" + posixpath.basename(drawing_path) + ".rels"
    rels = {}
    if rel_path not in zip_file.namelist():
        return rels
    root = ET.fromstring(zip_file.read(rel_path))
    for rel in root:
        rels[rel.attrib.get("Id")] = rel.attrib.get("Target")
    return rels


def extract_image_anchors(path):
    anchors = []
    with zipfile.ZipFile(path) as z:
        drawings = [
            n
            for n in z.namelist()
            if n.startswith("xl/drawings/drawing") and n.endswith(".xml")
        ]
        for drawing_path in drawings:
            rels = load_rels(z, drawing_path)
            root = ET.fromstring(z.read(drawing_path))
            for tag in ("twoCellAnchor", "oneCellAnchor"):
                for anchor in root.findall("xdr:" + tag, NS):
                    frm = anchor.find("xdr:from", NS)
                    pic = anchor.find("xdr:pic", NS)
                    if frm is None or pic is None:
                        continue
                    blip = pic.find(".//a:blip", NS)
                    if blip is None:
                        continue
                    rid = blip.attrib.get(
                        "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
                    )
                    target = rels.get(rid)
                    if not target:
                        continue
                    media_path = posixpath.normpath(
                        posixpath.join(posixpath.dirname(drawing_path), target)
                    )
                    anchors.append(
                        {
                            "drawing_path": drawing_path,
                            "media_path": media_path,
                            "anchor_row": int(frm.findtext("xdr:row", default="0", namespaces=NS)) + 1,
                            "anchor_col": int(frm.findtext("xdr:col", default="0", namespaces=NS)) + 1,
                        }
                    )
    anchors.sort(key=lambda x: (x["anchor_row"], x["anchor_col"], x["media_path"]))
    return anchors


def create_schema(cursor):
    for ddl in DDL:
        cursor.execute(ddl)


def upsert_source(cursor, config, file_hash, sheet_names, title_text, total_images):
    source_path = config["source_path"]
    cursor.execute(
        """
        INSERT INTO `ai_rule_source_file`
          (`rule_type`, `source_name`, `source_path`, `file_sha256`, `file_size_bytes`, `sheet_names`, `title_text`, `total_images`)
        VALUES (%s, %s, %s, %s, %s, %s, %s, %s)
        ON DUPLICATE KEY UPDATE
          `rule_type` = VALUES(`rule_type`),
          `source_name` = VALUES(`source_name`),
          `source_path` = VALUES(`source_path`),
          `file_size_bytes` = VALUES(`file_size_bytes`),
          `sheet_names` = VALUES(`sheet_names`),
          `title_text` = VALUES(`title_text`),
          `total_images` = VALUES(`total_images`)
        """,
        (
            config["rule_type"],
            config["source_name"],
            source_path,
            file_hash,
            os.path.getsize(source_path),
            json.dumps(sheet_names, ensure_ascii=False),
            title_text,
            total_images,
        ),
    )
    cursor.execute("SELECT `id` FROM `ai_rule_source_file` WHERE `file_sha256` = %s", (file_hash,))
    return cursor.fetchone()["id"]


def save_rule_images(cursor, config, source_file_id, sheet_name, anchors):
    output_dir = Path(config["output_dir"])
    output_dir.mkdir(parents=True, exist_ok=True)
    inserted = 0
    with zipfile.ZipFile(config["source_path"]) as z:
        for index, anchor in enumerate(anchors, start=1):
            data = z.read(anchor["media_path"])
            image_hash = hashlib.sha256(data).hexdigest()
            ext = Path(anchor["media_path"]).suffix.lower() or ".bin"
            mime = mimetypes.types_map.get(ext, "application/octet-stream")
            filename = f"{config['rule_type'].lower()}_{index:04d}_{image_hash[:12]}{ext}"
            local_path = str(output_dir / filename)
            if not os.path.exists(local_path):
                with open(local_path, "wb") as f:
                    f.write(data)

            cursor.execute(
                """
                INSERT INTO `ai_rule_image`
                  (`rule_type`, `source_file_id`, `source_sheet`, `image_index`, `anchor_row`, `anchor_col`,
                   `original_media_name`, `image_mime`, `image_ext`, `image_sha256`, `image_size_bytes`, `local_path`)
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
                ON DUPLICATE KEY UPDATE
                  `rule_type` = VALUES(`rule_type`),
                  `source_file_id` = VALUES(`source_file_id`),
                  `source_sheet` = VALUES(`source_sheet`),
                  `image_index` = VALUES(`image_index`),
                  `anchor_row` = VALUES(`anchor_row`),
                  `anchor_col` = VALUES(`anchor_col`),
                  `original_media_name` = VALUES(`original_media_name`),
                  `image_mime` = VALUES(`image_mime`),
                  `image_ext` = VALUES(`image_ext`),
                  `image_size_bytes` = VALUES(`image_size_bytes`),
                  `local_path` = VALUES(`local_path`)
                """,
                (
                    config["rule_type"],
                    source_file_id,
                    sheet_name,
                    index,
                    anchor["anchor_row"],
                    anchor["anchor_col"],
                    anchor["media_path"],
                    mime,
                    ext.lstrip("."),
                    image_hash,
                    len(data),
                    local_path,
                ),
            )
            inserted += 1
    return inserted


def main():
    conn = pymysql.connect(**DB_CONFIG)
    try:
        with conn.cursor() as cursor:
            create_schema(cursor)
            summary = []
            for config in WORKBOOKS:
                source_path = config["source_path"]
                sheet_names = workbook_sheet_names(source_path)
                title_text = first_non_empty_cell(source_path)
                anchors = extract_image_anchors(source_path)
                file_hash = sha256_file(source_path)
                source_file_id = upsert_source(
                    cursor, config, file_hash, sheet_names, title_text, len(anchors)
                )
                image_count = save_rule_images(
                    cursor,
                    config,
                    source_file_id,
                    sheet_names[0] if sheet_names else None,
                    anchors,
                )
                summary.append(
                    {
                        "rule_type": config["rule_type"],
                        "source_name": config["source_name"],
                        "title_text": title_text,
                        "images": image_count,
                    }
                )
            conn.commit()
            print(json.dumps(summary, ensure_ascii=False, indent=2))
    finally:
        conn.close()


if __name__ == "__main__":
    main()
