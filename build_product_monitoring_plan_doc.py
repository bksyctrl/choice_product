from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


OUTPUT = "选品与数据监控系统功能规划文档.docx"


sections = [
    {
        "title": "一、跨品类板块",
        "plan": [
            (
                "数据源与抓取",
                "以 FastMoss 平台商品搜索榜为切入点（例如选取女装与女士内衣中的女士上装），抓取每个商品的主图及核心数据。底层的数据表结构与现有的手机壳类目结构保持一致，方便系统统一纳管。"
            ),
            (
                "专区建立与前端展示",
                "新建“跨平台热品榜”栏目，聚合多种类目的商品。在商品列表页顶部增加精细化筛选栏，筛选维度包括但不限于：近7天销量、主图有无 Logo、是否有图案、图案清晰度等。"
            ),
            (
                "图像工具集成",
                "列表中增加一个“一键抠图/提取插画”的快捷交互按钮，提升素材处理效率。"
            ),
        ],
        "questions": [
            (
                "品类选择标准",
                "跨品类的栏目中，到底需要纳入哪些具体的商品排名集进行常态化监控？"
            ),
            (
                "TikTok 热搜对接",
                "TikTok 的爆点热搜具体需要爬取什么内容？关键词该怎么细分？这些热搜数据怎么结合到当前系统中，期望的最终呈现形式是怎么样的？"
            ),
            (
                "深度归因分析",
                "跨品类热搜榜的打分机制是否需要体现差异化？是不是不仅要抓取图片，还要有 AI 介入分析“为什么这么热”（是因为材质、标题，还是因为印花设计）？"
            ),
        ],
    },
    {
        "title": "二、跨平台板块",
        "plan": [
            (
                "多平台数据覆盖",
                "定向爬取 Temu、各类定制化独立站、亚马逊等网站的手机壳销售数据。"
            ),
            (
                "聚合表与差异化兼容",
                "将多源数据规划到选品平台的“聚合表（统一表）”中。设定最低爬取字段标准：评分、评论数、销量、商品标题。同时，底层数据库在设计时，必须保留每个平台专属的数据结构扩展字段。"
            ),
        ],
        "questions": [
            (
                "关键词策略",
                "跨平台的爬取，是直接搜索核心短词（如：手机壳），还是主要搜索长尾词（如：手机壳防水，防摔）来获取更精准的细分市场数据？"
            ),
            (
                "风向标选择",
                "跨平台的数据抓取，是拿 FastMoss 等平台的热搜品还是热销品？热销品意味着已被市场验证过有需求，但需要最终确认在 TikTok 生态中，到底是“搜索”带来的流量大，还是“销量排行”带来的流量大。"
            ),
        ],
    },
    {
        "title": "三、商品列表页与详细页优化",
        "plan": [
            (
                "全局综合打分",
                "在所有的商品列表页（包括跨品类热搜表）中引入多维度打分功能，直观量化商品潜力。"
            ),
            (
                "详情页深度分析",
                "商品详细页内嵌深度数据分析功能，展示趋势图和各项指标表现。"
            ),
            (
                "IP 分析与 UI 呈现",
                "针对 IP 修改意见或提示词优化，采用折叠式阅读体验。初始界面先输出初步分析和“一句话总结”；若用户想看详细原因，点击“详细分析”按钮后才展开全文，提升页面整洁度。"
            ),
        ],
        "questions": [
            (
                "打分算法与权重",
                "打分功能到底用什么样的具体策略？如果依据数据变化规则（如近几天对比近一个月的数据变化率）、TikTok 热搜命中率、热推材质等多角度来评估，各项指标的具体权重比例如何界定？例如：变化率权重 30%，热搜词命中权重 20%，材质属性权重 30% 等，需要制定明确的模型。"
            ),
        ],
    },
    {
        "title": "四、数据监控与预警板块",
        "plan": [
            (
                "全方位精准追踪",
                "提供灵活的数据监控功能，用户只需输入店名、商品名或特定的特征值，即可准确调取并查看该商品在店铺中的历史与实时表现。"
            ),
            (
                "AI 智能报警箱",
                "引入自动化的数据报警机制。设定基础报警规则线（例如：评分跌破 3 分、点击量或转化率急剧下滑等）。由 AI 实时巡检分析，表现差的异常数据自动进入“报警箱”集中处理。"
            ),
        ],
        "questions": [
            (
                "监控策略细化",
                "监控平台的具体监控策略是什么？包含监控频率、数据抓取深度、以及 AI 判断“异常”的综合阈值该如何具体设定？"
            ),
        ],
    },
]


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, width_dxa=9360):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(width_dxa))
    tbl_w.set(qn("w:type"), "dxa")


def style_paragraph(paragraph, font_name="Calibri", size=11, color="000000", bold=False):
    for run in paragraph.runs:
        run.font.name = font_name
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        run.font.size = Pt(size)
        run.font.color.rgb = RGBColor.from_string(color)
        run.font.bold = bold


def add_label_paragraph(doc, label, text, style=None):
    p = doc.add_paragraph(style=style)
    r = p.add_run(f"{label}：")
    r.bold = True
    r.font.color.rgb = RGBColor(31, 77, 120)
    r.font.name = "Calibri"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r.font.size = Pt(11)
    r2 = p.add_run(text)
    r2.font.name = "Calibri"
    r2._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r2.font.size = Pt(11)
    return p


def build_doc():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    for name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 16, 8),
        ("Heading 2", 13, "2E74B5", 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("选品与数据监控系统功能规划文档")
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(22)
    run.font.bold = True
    run.font.color.rgb = RGBColor(11, 37, 69)
    title.paragraph_format.space_after = Pt(10)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = subtitle.add_run("功能方向、数据接入、页面优化与监控预警规划")
    r.font.name = "Calibri"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(102, 112, 133)
    subtitle.paragraph_format.space_after = Pt(18)

    intro_table = doc.add_table(rows=1, cols=1)
    set_table_width(intro_table)
    cell = intro_table.cell(0, 0)
    set_cell_shading(cell, "F4F6F9")
    set_cell_margins(cell, top=140, bottom=140, start=180, end=180)
    p = cell.paragraphs[0]
    r = p.add_run("文档定位：")
    r.bold = True
    r.font.color.rgb = RGBColor(31, 58, 95)
    r.font.name = "Calibri"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r.font.size = Pt(11)
    r2 = p.add_run("围绕跨品类、跨平台、商品页面优化与数据监控预警四个方向，梳理系统建设思路与后续需要确认的问题。")
    r2.font.name = "Calibri"
    r2._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r2.font.size = Pt(11)

    for block in sections:
        doc.add_heading(block["title"], level=1)

        doc.add_heading("1. 规划思路", level=2)
        for label, text in block["plan"]:
            add_label_paragraph(doc, label, text)

        doc.add_heading("2. 待确认问题", level=2)
        for label, text in block["questions"]:
            add_label_paragraph(doc, label, text)

    doc.add_paragraph()
    footer_note = doc.add_paragraph()
    footer_note.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = footer_note.add_run("文档状态：初版规划草案")
    r.font.name = "Calibri"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(102, 112, 133)

    doc.save(OUTPUT)


if __name__ == "__main__":
    build_doc()
